# -*- coding: utf-8 -*-
"""修正版实际价 Excel：结论、逐笔、每日权益、旧版对比。"""
import json

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = '03_结果'
XLSX = f'{OUT}/实际价_修正版.xlsx'

M = json.load(open(f'{OUT}/实际价_修正版_指标.json', encoding='utf-8'))
old = json.load(open(f'{OUT}/最终_指标.json', encoding='utf-8'))['actual']
tr = pd.read_csv(f'{OUT}/交易记录_中文_修正版.csv',
                 parse_dates=['信号日', '买入日', '卖出/期末日'])
eq = pd.read_csv(f'{OUT}/实际价_修正版_每日权益.csv', parse_dates=['date'])
eq['drawdown'] = eq.equity / eq.equity.cummax() - 1
old_tr = pd.read_csv(f'{OUT}/旧版_不可靠/实际价_逐笔_旧版.csv')
old365 = old_tr[(old_tr.thscode == '002365.SZ') & (old_tr.signal_date == '2018-04-23')].iloc[0]
old607 = old_tr[(old_tr.thscode == '002607.SZ') & (old_tr.signal_date == '2018-12-17')].iloc[0]
new365 = tr[(tr['代码'] == '002365.SZ') & (tr['信号日'] == pd.Timestamp('2018-04-23'))].iloc[0]
new607 = tr[(tr['代码'] == '002607.SZ') & (tr['信号日'] == pd.Timestamp('2018-12-17'))].iloc[0]

concl = pd.DataFrame([
    ('旧版实际价（不可用）', '42.58%', '2,362,270', '69.7%', '6.26%', '3.19',
     '送转/成本 bug；002365 假亏损'),
    ('修正版实际价', f"{M['cagr']*100:.2f}%", f"{M['final']:,.0f}",
     f"{M['win']*100:.1f}%", f"{M['avg']*100:.2f}%", '4.08',
     '100股整手；送转/分红/成本基准修正'),
    ('复权主口径', '40.01%', '2,074,584', '71.7%', '6.68%', '4.06',
     '与修正版实际价只差 0.58pp'),
], columns=['口径', '年化', '期末权益', '胜率', '平均单笔', '利润因子', '说明'])

compare = pd.DataFrame([
    ('002365.SZ', '2018-04-23', '2018-04-26', '2018-05-18',
     old365.entry_px, old365.exit_px, old365.shares, old365.pnl, old365.ret,
     new365['股数'], new365['净盈亏(元)'], new365['净收益率'],
     new365['复权收益率'], new365['现金分红(元)']),
    ('002607.SZ', '2018-12-17', '2018-12-20', '2019-01-11',
     old607.entry_px, old607.exit_px, old607.shares, old607.pnl, old607.ret,
     new607['股数'], new607['净盈亏(元)'], new607['净收益率'],
     new607['复权收益率'], new607['现金分红(元)']),
], columns=['代码', '信号日', '买入日', '卖出日',
            '旧买入价', '旧卖出价', '旧股数', '旧净盈亏', '旧净收益率',
            '新股数', '新净盈亏', '新净收益率', '复权收益率', '现金分红'])

e = eq.set_index('date').equity
ye = e.resample('YE').last()
ann = ye.pct_change()
ann.iloc[0] = ye.iloc[0] / 100000 - 1
ann = ann.to_frame('年度收益')
ann.index = [d.year for d in ann.index]

notes = pd.DataFrame([
    ('标题', '实际价逐笔修正版（旧版不可靠）'),
    ('旧版问题1', '002365 十转五：2,500 股被调成 3,868.95（应 3,750）；'
                  '成本用“调整后股数×原始买入价”，把送转股重复计入，'
                  '出现 -24,533 假亏损。'),
    ('旧版问题2', '002607 漏了 4,095 元现金分红：旧版 -129/-0.16%，'
                  '应约 +3,974/+5.28%。'),
    ('旧版问题3', '旧版实际价 42.58% 比复权 40.01% 高 2.57pp，是虚高；'
                  '旧版每日权益、年度/月度实际价表同源，不要直接引用。'),
    ('修正方法', '100股整手；送转用同口径因子（002365 显式 1.5 倍）；'
                 '现金分红计入盈亏；成本用原始买入现金；'
                 '佣金万2.5（最低5元）+印花千0.5（卖）+滑点千1（单边）。'),
    ('验证', f"99笔；修正版 vs 复权平均绝对差 "
             f"{M['mean_abs_diff']*100:.3f}pp；最大差 0.82pp。"),
    ('新文件', '实际价_修正版_逐笔.csv / 交易记录_中文_修正版.csv / '
               '实际价_修正版_每日权益.csv / 实际价_修正版_指标.json / '
               '实际价_修正说明.md'),
    ('不构成', '历史回测不代表未来；不构成投资建议。'),
], columns=['项目', '内容'])

with pd.ExcelWriter(XLSX, engine='openpyxl') as w:
    concl.to_excel(w, sheet_name='修正结论', index=False)
    tr.to_excel(w, sheet_name='修正版逐笔', index=False)
    eq.to_excel(w, sheet_name='修正版每日权益', index=False)
    ann.to_excel(w, sheet_name='修正版年度')
    compare.to_excel(w, sheet_name='旧版对比', index=False)
    notes.to_excel(w, sheet_name='说明', index=False)
    wb = w.book
    wb.move_sheet('说明', offset=-(len(wb.sheetnames) - 1))

wb = load_workbook(XLSX)
HEAD_FILL = PatternFill('solid', fgColor='1F4E79')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
for ws in wb.worksheets:
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal='center', vertical='center',
                                wrap_text=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
for ws in wb.worksheets:
    headers = [c.value for c in ws[1]]
    for j, h in enumerate(headers, 1):
        letter = get_column_letter(j)
        if h in ('净收益率', '复权收益率', '旧净收益率', '新净收益率',
                 'drawdown'):
            for c in ws[letter][1:]:
                c.number_format = '0.00%'
        if h in ('净盈亏(元)', '旧净盈亏', '新净盈亏', '现金分红',
                 'equity', 'cash'):
            for c in ws[letter][1:]:
                c.number_format = '#,##0'
        if h in ('股数', '新股数', '旧股数', 'npos'):
            for c in ws[letter][1:]:
                c.number_format = '#,##0'
        if h in ('买入价', '卖出价', '旧买入价', '旧卖出价'):
            for c in ws[letter][1:]:
                c.number_format = '0.00'
        ws.column_dimensions[letter].width = min(max(10, len(str(h)) * 1.8), 26)
    if ws.title == '说明':
        ws.column_dimensions['A'].width = 16
        ws.column_dimensions['B'].width = 100
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[r].height = 35
        ws.auto_filter.ref = None
    if ws.title == '修正版逐笔':
        ws.freeze_panes = 'D2'
    if ws.title == '修正版年度':
        for c in ws['B'][1:]:
            c.number_format = '0.00%'
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 14
    if ws.title == '修正版每日权益':
        ws.freeze_panes = 'B2'
wb.save(XLSX)
print('excel done', XLSX)
