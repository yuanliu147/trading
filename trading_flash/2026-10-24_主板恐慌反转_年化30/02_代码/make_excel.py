# DEPRECATED 2026-10-25: 本脚本的实际价部分读取旧版文件，不要重跑覆盖；
# 修正版见 actual_engine_fixed.py / make_actual_fixed_excel.py。
import pandas as pd, numpy as np, json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
M=json.load(open('03_结果/最终_指标.json'))
ann=pd.read_csv('03_结果/最终_年度_展示.csv',index_col=0); ann.columns=['年度收益']
mon=pd.read_csv('03_结果/最终_月度.csv',index_col=0); mon.index=mon.index.astype(str); mon.columns=['月度收益']
tr=pd.read_csv('03_结果/实际价_逐笔.csv',parse_dates=['signal_date','entry_date','exit_date'])
tr.columns=['代码','信号日','买入日','卖出日','买入价','卖出价','股数','净盈亏','净收益率','退出原因','复权收益率']
tr['退出原因']=tr['退出原因'].map({'time':'15日到期','target':'+20%止盈','eod':'期末'})
eq=pd.read_csv('03_结果/实际价_每日权益.csv',parse_dates=['date'])
eq['drawdown']=eq.equity/eq.equity.cummax()-1
eq.columns=['日期','权益','现金','持仓数','回撤']
sp=pd.read_csv('03_结果/最终_回撤区间.csv',parse_dates=['start','trough','recovery'])
sp.columns=['开始','谷底','修复','最大回撤','持续天数']
g=pd.read_csv('03_结果/grid_all_periods.csv'); g['min3']=g[['train_cagr','val_cagr','test_cagr']].min(axis=1)
g=g[['dd','rsi','hold','delay','tp','train_cagr','val_cagr','test_cagr','full_cagr','full_dd','full_n','min3']]
g.columns=['回撤阈值','RSI阈值','持有日','延迟入场','止盈','训练年化','验证年化','检验年化','全期年化','全期回撤','全期笔数','三段最低']
ev=pd.read_csv('03_结果/event_scan.csv')
ev=ev[ev.h==10][['event','all_n','all_r','all_ex','test_ex','all_win']]
ev.columns=['事件','全期样本','全期10日收益','全期10日超额','检验期10日超额','全期胜率']
ev['事件']=ev['事件'].map({'rev_panic':'恐慌超跌反转','rev_deep':'深度超跌','pull_ma20_tight':'均线回踩(紧)','macd_lowcross':'MACD零下金叉','macd_highcross':'MACD零上金叉','brk60_vol':'60日突破放量','lu_first':'首板','lu_multi':'连板','lu_broken':'炸板','gap_up':'高开高走','trend_quiet':'安静趋势','engulf_up':'看涨吞没'}).fillna(ev['事件'])
wf=pd.DataFrame([['候选参数（全期选参）',0.339,0.400,0.518,0.400,-0.160,99],['走步选参（仅用过去）',0.468,0.260,0.485,0.407,-0.225,103]],columns=['口径','训练','验证','检验','全期','全期回撤','笔数'])
cur=pd.DataFrame([['2026-10-23','主板等权20日',-0.0519,'开（<-5%）'],['2026-10-23','主板宽度20',0.1900,'低'],['2026-10-23','个股全条件',123,'有信号'],['2026-10-23','前2候选','603778.SH / 605376.SH','按回撤最深'],['2026-10-28','计划入场','2 仓各 50%','一字板/停牌不买']],columns=['日期','项目','数值','判定'])
one=pd.DataFrame([
 ['策略名称','主板恐慌洗盘反转（Panic Shakeout Reversal）'],['股票池','沪深主板（600/601/603/605/000/001/002/003），剔除5%涨跌停样态（ST代理）'],
 ['信号','年内仍强（250日收益>0）+ 自250日高点回撤>45% + RSI14<35 + 收盘位于当日振幅上半部 + 主板等权20日收益<-5%'],
 ['排序','按250日回撤深度升序，取前2只'],['入场','信号后第3个交易日开盘；一字涨停/停牌不买'],['仓位','2 仓，每仓 50% 权益；空仓等待'],
 ['止盈','盘中触及 +20% 即卖（限价/触价），T+1 后生效'],['到期','持有15个交易日收盘卖出'],['止损','无普通止损；灾难止损经测试会扩大回撤，不采用'],
 ['成本','佣金万2.5（最低5元）+ 印花税千0.5（卖）+ 滑点千1（单边）'],['本金','10 万元'],
 ['全期年化（实际价整手）',f"{M['hfq']['cagr']*100:.1f}%"],['全期最大回撤',f"{M['hfq']['dd']*100:.1f}%"],
 ['交易笔数',M['hfq_trades']],['胜率',f"{M['hfq_win']*100:.1f}%"],['平均单笔',f"{M['hfq_avg']*100:.2f}%"],
 ['盈亏比（利润因子）','4.06'],['持仓时间占比','31.1%'],
 ['前5笔盈利占比',f"{M['hfq_top5']*100:.1f}%"],['最大回撤区间','2026-06-08 → 2026-06-30，-16.00%，2026-08-06 修复'],
 ['诚实走步预期','验证 26.0%、2024–2026 检验 48.5%（仅用过去选参）；576 组中 7 组三段≥30%，按 25%–30% 规划'],
 ['当前状态','2026-10-23：市场闸门开，123 只全条件；前2 603778.SH、605376.SH，计划 2026-10-28 开盘'],
],columns=['项目','内容'])
with pd.ExcelWriter('恐慌洗盘反转_数据表.xlsx',engine='openpyxl') as w:
    one.to_excel(w,sheet_name='一页纸',index=False)
    pd.DataFrame([['信号','年内仍强 + 回撤>45% + RSI<35 + 反转收盘 + 市场20日<-5%'],['排序','回撤最深优先'],['入场','信号后第3日开盘'],['仓位','2×50%'],['止盈','+20%触价'],['持有','15日'],['止损','无'],['成本','万2.5+印花千0.5+滑点千1']],columns=['项目','规则']).to_excel(w,sheet_name='规则',index=False)
    ann.to_excel(w,sheet_name='年度'); mon.to_excel(w,sheet_name='月度'); tr.to_excel(w,sheet_name='逐笔',index=False)
    eq.to_excel(w,sheet_name='每日权益',index=False); sp.to_excel(w,sheet_name='回撤区间',index=False)
    g.to_excel(w,sheet_name='参数敏感性',index=False); wf.to_excel(w,sheet_name='走步检验',index=False)
    cur.to_excel(w,sheet_name='当前状态',index=False); ev.to_excel(w,sheet_name='事件扫描',index=False)
    pd.DataFrame([['数据源','同花顺本地 DuckDB：~/Library/Application Support/hithink-finance/data/market.duckdb'],['窗口','2016-10-13 至 2026-10-23；信号从 2017-11 起'],['股票池','沪深主板；仅当前上市，未含退市股（幸存者偏差）'],['复权','后复权算收益；实际价整手复核'],['ST','5%涨跌停样态代理剔除'],['第三方','未使用 baostock 等第三方数据'],['注意','历史回测不代表未来；先纸面50笔']],columns=['项目','说明']).to_excel(w,sheet_name='数据说明',index=False)
wb=load_workbook('恐慌洗盘反转_数据表.xlsx')
hdr=PatternFill('solid',fgColor='1F4E79'); hfont=Font(color='FFFFFF',bold=True,size=10)
thin=Side(style='thin',color='D9D9D9'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
for ws in wb.worksheets:
    ws.freeze_panes='A2'
    for c in ws[1]:
        c.fill=hdr; c.font=hfont; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=border
    for row in ws.iter_rows(min_row=1,max_row=ws.max_row):
        for c in row: c.border=border
    for col in ws.columns:
        mx=max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width=min(max(10,mx*1.6),55)
    if ws.title=='一页纸': ws.column_dimensions['A'].width=24; ws.column_dimensions['B'].width=100
    if ws.title=='逐笔':
        for c in ws['I'][1:]: c.number_format='0.00%'
        for c in ws['H'][1:]: c.number_format='#,##0'
    if ws.title=='每日权益':
        for c in ws['B'][1:]: c.number_format='#,##0'
        for c in ws['E'][1:]: c.number_format='0.00%'
    if ws.title=='年度':
        for c in ws['B'][1:]: c.number_format='0.00%'
    if ws.title=='月度':
        for c in ws['B'][1:]: c.number_format='0.00%'
    if ws.title=='参数敏感性':
        for col in ['F','G','H','I','J','L']:
            for c in ws[col][1:]: c.number_format='0.0%'
    if ws.title=='走步检验':
        for col in ['B','C','D','E','F']:
            for c in ws[col][1:]: c.number_format='0.0%'
wb.save('恐慌洗盘反转_数据表.xlsx')
print('excel done')
