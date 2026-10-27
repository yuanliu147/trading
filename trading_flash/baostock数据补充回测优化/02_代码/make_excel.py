#!/usr/bin/env python3
"""生成中文 Excel 数据表与 CSV 副本。"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
R = ROOT / "03_结果"
E = ROOT / "05_证据"
OUT = ROOT / "baostock补充研究_数据表.xlsx"
CSVDIR = R / "数据表_CSV"
CSVDIR.mkdir(exist_ok=True)

GREEN = "1F6FB2"
LIGHT = "EAF1F8"
ORANGE = "C0392B"
WARN = "FDE9E7"

NAME_MAP = {
    "A_仅当前上市": "仅当前上市",
    "B_含退市股": "含退市股",
    "B_含退市_无ST过滤": "含退市股（不排除风险警示）",
    "C_含退市_排除ST": "含退市股+排除风险警示",
    "C_排除ST": "排除风险警示",
    "C_基础_无止损": "排除风险警示（无止损）",
    "C_止损15": "排除风险警示+止损15%",
    "C_止损20": "排除风险警示+止损20%",
    "C_止损25": "排除风险警示+止损25%",
    "C_止损15_仓位35": "止损15%+单票35%",
    "C_止损25_仓位35": "止损25%+单票35%",
    "C_止损15_仓位25": "止损15%+单票25%",
    "C_止损25_仓位25": "止损25%+单票25%",
    "M1_市场跌超10": "市场20日跌幅超10%才买",
    "M2_涨幅超20": "250日涨幅>20%",
    "M3_涨幅超50": "250日涨幅>50%",
    "M4_波动超45": "20日波动率>45%",
    "M5_波动超60": "20日波动率>60%",
    "M6_回撤不超75": "排除回撤超过75%",
    "M7_换手不低于2": "20日换手率≥2%",
    "M8_组合": "组合过滤（市场+涨幅+波动+回撤）",
    "M9_组合轻": "组合过滤（市场+涨幅）",
}
COL_MAP = {
    "期间": "期间", "口径": "口径", "标签": "方案", "过滤": "过滤条件",
    "年化": "年化收益", "最大回撤": "最大回撤", "卡玛": "卡玛比率",
    "夏普": "夏普比率", "交易数": "交易笔数", "胜率": "胜率",
    "单笔均值": "单笔平均收益", "单笔中位数": "单笔收益中位数",
    "持仓占比": "持仓天数占比", "期末倍数": "期末净值倍数",
    "期末金额": "期末金额（元）", "起始": "起始日", "结束": "结束日",
    "年数": "年数", "信号日": "信号日数", "候选行": "候选行数",
    "候选股票": "候选股票数", "左Top2": "左侧前二数", "右Top2": "右侧前二数",
    "交集": "前二交集", "Jaccard": "前二重合度", "相同日": "相同信号日",
    "对比": "对比", "维度": "维度", "分组": "分组", "样本": "样本数",
    "均值": "平均收益", "中位数": "收益中位数", "平均MFE": "平均最大浮盈",
    "平均MAE": "平均最大浮亏", "止损": "止损", "单票仓位": "单票仓位",
    "代码": "代码", "名称": "名称",
    "cagr": "年化收益", "max_dd": "最大回撤", "calmar": "卡玛比率", "sharpe": "夏普比率",
    "n_trades": "交易笔数", "win": "胜率", "avg": "单笔平均收益", "median": "单笔收益中位数",
    "exposure": "持仓天数占比", "final": "期末净值倍数", "final_yuan": "期末金额（元）",
    "years": "年数", "first_date": "起始日", "last_date": "结束日",
}


def load(path: Path) -> pd.DataFrame:
    return pd.read_csv(path)


def prettify(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for col in ["口径", "标签", "方案"]:
        if col in df.columns:
            df[col] = df[col].map(lambda x: NAME_MAP.get(str(x), str(x)))
    df = df.rename(columns={k: v for k, v in COL_MAP.items() if k in df.columns})
    return df


def pct_cols(df: pd.DataFrame) -> list[str]:
    return [c for c in df.columns
            if any(k in c for k in ["收益", "回撤", "胜率", "占比", "年化", "浮盈", "浮亏",
                                    "触达率", "仓位", "止损", "重合度", "位置"])]


def write_sheet(wb, name: str, df: pd.DataFrame, note: str = "", pct: bool = True, width: int = 18):
    ws = wb.create_sheet(name[:31])
    ws.sheet_view.showGridLines = False
    ws["A1"] = name
    ws["A1"].font = Font(bold=True, size=14, color="1F3B57")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, min(len(df.columns), 8)))
    ws.row_dimensions[1].height = 24
    header_row = 3
    for j, col in enumerate(df.columns, 1):
        cell = ws.cell(row=header_row, column=j, value=col)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="FFFFFF"))
    for i, row in enumerate(df.itertuples(index=False), header_row + 1):
        for j, val in enumerate(row, 1):
            cell = ws.cell(row=i, column=j)
            if isinstance(val, float):
                cell.value = val
                if pct and df.columns[j - 1] in pct_cols(df):
                    cell.number_format = "0.0%"
                elif df.columns[j - 1] == "期末净值倍数":
                    cell.number_format = "0.00"
                elif abs(val) < 10:
                    cell.number_format = "0.000"
                else:
                    cell.number_format = "#,##0"
            else:
                cell.value = val
            cell.font = Font(size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if (i - header_row) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
    for j, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(width, max(10, len(str(col)) * 2 + 4))
    if note:
        row = header_row + len(df) + 2
        ws.cell(row=row, column=1, value=note).font = Font(size=9, italic=True, color="666666")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(1, min(len(df.columns), 8)))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def main():
    wb = __import__("openpyxl").Workbook()
    wb.remove(wb.active)

    # 1 结论摘要
    core = load(R / "核心结论.csv")
    core = prettify(core)
    summary = core[[c for c in ["期间", "口径", "年化收益", "最大回撤", "卡玛比率", "交易笔数", "胜率", "期末净值倍数"] if c in core.columns]]
    write_sheet(wb, "结论摘要", summary,
                "口径：后复权收益、实际价成交；CAGR 从 2017-10-19 或 2005-01-04 起算；"
                "含退市股后 2017-2026 从 43.2% 降到 29.8%，排除风险警示后 35.0%。")

    # 2 口径分解
    dec = load(R / "口径分解.csv")
    dec = prettify(dec)
    keep = [c for c in ["期间", "方案", "年化收益", "最大回撤", "卡玛比率", "交易笔数", "胜率", "期末净值倍数"] if c in dec.columns]
    write_sheet(wb, "口径分解", dec[keep],
                "仅当前上市→排除风险警示→含退市股：用于区分“风险警示过滤”和“退市股/排名变化”各自的贡献。")

    # 3 年度收益
    def annual_table(period_folder: str, label_map: dict[str, str]) -> pd.DataFrame:
        frames = {}
        for label, name in label_map.items():
            p = R / period_folder / f"{label}_年度.csv"
            if not p.exists():
                continue
            s = pd.read_csv(p, index_col=0).iloc[:, 0]
            s.index = s.index.astype(int).astype(str)
            frames[name] = s
        df = pd.DataFrame(frames)
        df.index.name = "年份"
        return df.reset_index()
    a171 = annual_table("2017_2026", {"A_仅当前上市": "仅当前上市", "B_含退市股": "含退市股"})
    c171 = annual_table("指标组合_2017_2026", {"C_排除ST": "含退市+排除风险警示"})
    write_sheet(wb, "年度收益_2017_2026", a171.merge(c171, on="年份", how="outer"),
                "独立 2017-2026 区间的年度收益；不是 2005-2026 全区间中的分段。")
    a052 = annual_table("2005_2026", {"A_仅当前上市": "仅当前上市", "B_含退市股": "含退市股"})
    c052 = annual_table("指标优化/2005-2026", {"C_排除ST": "含退市+排除风险警示"})
    write_sheet(wb, "年度收益_2005_2026", a052.merge(c052, on="年份", how="outer"),
                "2005-2026 连续复利路径；2008 年 C 为 -52.7%，最大回撤 -71.5%。")

    # 4 事件研究
    ev1 = load(R / "事件研究_2005_2016.csv"); ev1["期间"] = "2005-2016"
    ev2 = load(R / "事件研究_2017_2026.csv"); ev2["期间"] = "2017-2026"
    ev = pd.concat([ev1, ev2], ignore_index=True)
    ev = ev[["期间", "维度", "分组", "样本", "胜率", "止盈20率", "均值", "中位数", "平均MFE", "平均MAE"]]
    ev = ev.rename(columns={"止盈20率": "20%止盈触达率", "平均MFE": "平均最大浮盈",
                            "平均MAE": "平均最大浮亏", **COL_MAP})
    ev["维度"] = ev["维度"].replace({"ST": "风险警示", "PB": "市净率"})
    ev["分组"] = ev["分组"].replace({"ST": "是", "非ST": "否", "是": "是", "否": "否"})
    write_sheet(wb, "指标事件研究", ev,
                "所有满足冻结条件的信号（不只前二），T+3 开盘至 T+17 收盘；信号日高度重叠，只作描述统计。")

    # 5 组合测试
    opt = load(R / "指标优化" / "指标优化汇总.csv")
    opt = prettify(opt)
    keep = [c for c in ["期间", "方案", "信号日数", "年化收益", "最大回撤", "卡玛比率", "交易笔数", "胜率", "期末净值倍数"] if c in opt.columns]
    write_sheet(wb, "组合测试", opt[keep],
                "每个过滤单独测试；没有一项能在 2005-2016 与 2017-2026 同时改善收益和回撤。")

    # 6 止损与仓位
    risk = load(R / "风险叠加汇总.csv")
    risk = prettify(risk)
    keep = [c for c in ["期间", "方案", "止损", "单票仓位", "年化收益", "最大回撤", "卡玛比率", "交易笔数", "胜率", "期末净值倍数"] if c in risk.columns]
    write_sheet(wb, "止损与仓位", risk[keep],
                "止损 15%/25% 在两个样本期都同时改善收益和回撤；仓位缩放按比例降低收益与回撤，卡玛基本不变。")

    # 7 退市股交易
    b = load(R / "退市股交易_2017-2026_B_含退市_无ST过滤.csv")
    c = load(R / "退市股交易_2017-2026_C_含退市_排除ST.csv")
    for df, name in [(b, "含退市"), (c, "含退市+排除风险警示")]:
        df["口径"] = name
    delist = pd.concat([b, c], ignore_index=True)
    delist["pnl"] = delist.proceeds - delist.invested
    cols = ["口径", "code", "signal_date", "entry_date", "exit_date", "entry_px", "exit_px",
            "ret", "hfq_ret", "actual_px_ret", "split_factor", "exit_reason", "is_st_signal", "pnl"]
    delist = delist[[c for c in cols if c in delist.columns]].rename(columns={
        "code": "代码", "signal_date": "信号日", "entry_date": "买入日", "exit_date": "卖出日",
        "entry_px": "买入价", "exit_px": "卖出价", "ret": "净收益率", "hfq_ret": "后复权收益率",
        "actual_px_ret": "实际价涨跌幅", "split_factor": "复权因子", "exit_reason": "退出原因",
        "is_st_signal": "信号日风险警示", "pnl": "净盈亏（元）"})
    write_sheet(wb, "退市股交易", delist,
                "“后来退市”指该股在 2026-10-23 之前退市，不代表买入时已知会退市；信号日风险警示为 1 的名字，"
                "在排除风险警示版本中不会买入。")

    # 8 Top2
    ov = load(R / "Top2重叠.csv").rename(columns={"期间": "期间", "对比": "对比",
                                                   "左Top2": "左侧前二数", "右Top2": "右侧前二数",
                                                   "交集": "前二交集", "Jaccard": "前二重合度",
                                                   "相同日": "相同信号日", "信号日": "信号日数"})
    write_sheet(wb, "前二重合度", ov,
                "前二重合度=交集/并集；含退市股后约 14%-20% 的前二配对不同，2022 年 600077 挤掉 600706 是典型例子。")

    # 9 数据校验
    ver = load(E / "校验报告.csv")
    ver["结果"] = ver.ok.map({True: "通过", False: "失败"})
    ver = ver[["name", "结果", "detail"]].rename(columns={"name": "检查项", "detail": "细节"})
    write_sheet(wb, "数据校验", ver, "107 项独立校验，全部通过；详见 05_证据/校验报告.md。", pct=False, width=60)

    # 10 复权修复
    rep = load(E / "复权修复事件.csv")
    rep = rep.rename(columns={"code": "代码", "date": "日期", "close": "收盘价", "preclose": "前收盘",
                              "prev_close": "上一收盘", "factor": "复权因子",
                              "prev_factor": "上一复权因子", "observed_jump": "观察跳变",
                              "expected_jump": "应有跳变", "mult": "修正乘数"})
    write_sheet(wb, "复权修复", rep,
                "16 行“仅因子跳变”已修复；判断标准：前收盘未变但复权因子跳变>2%。", width=22)

    # 11 002684
    x = load(E / "002684_信号日特征.csv")
    x = x[(x.date >= "2022-04-15") & (x.date <= "2022-06-10")]
    x = x.rename(columns={"date": "日期", "close": "收盘价", "drawdown250": "250日回撤",
                          "return_250d": "250日涨幅", "rsi14": "RSI14",
                          "close_location": "收盘位置", "amount20": "20日成交额",
                          "is_st": "风险警示", "market_return20_all": "市场20日收益"})
    x["风险警示"] = x["风险警示"].map({1: "是", 0: "否"})
    write_sheet(wb, "002684个案", x,
                "信号日 2022-04-21：风险警示=1，回撤 -61.4%，RSI 27.8；4/26 开盘 2.69 买入；"
                "5/5 起停牌；6/6 退市整理首日开盘 0.46，-82.9%。")

    # 中文交易记录
    ctr = load(OUT.parent / "03_结果" / "交易记录" / "交易记录_基准C_2005-2026.csv")
    write_sheet(wb, "交易记录_基准C", ctr,
                "2005-2026 基准C（含退市+排除风险警示）逐笔；名称为当前名称，复权因子>1 表示持有期间有送转/分红。",
                pct=True, width=16)
    rtr = load(OUT.parent / "03_结果" / "交易记录" / "交易记录_推荐配置_2005-2026.csv")
    write_sheet(wb, "交易记录_推荐配置", rtr,
                "2005-2026 推荐配置（2×35% + 止损25%）逐笔；净收益率已含费用。",
                pct=True, width=16)

    wb.save(OUT)
    print("saved", OUT)

    # CSV 副本
    summary.to_csv(CSVDIR / "结论摘要.csv", index=False, encoding="utf-8-sig")
    dec[[c for c in ["期间", "方案", "年化收益", "最大回撤", "卡玛比率", "交易笔数", "胜率", "期末净值倍数"]
         if c in dec.columns]].to_csv(CSVDIR / "口径分解.csv", index=False, encoding="utf-8-sig")
    ev.to_csv(CSVDIR / "指标事件研究.csv", index=False, encoding="utf-8-sig")
    opt[[c for c in ["期间", "方案", "信号日数", "年化收益", "最大回撤", "卡玛比率", "交易笔数", "胜率", "期末净值倍数"]
         if c in opt.columns]].to_csv(CSVDIR / "组合测试.csv", index=False, encoding="utf-8-sig")
    risk[[c for c in ["期间", "方案", "止损", "单票仓位", "年化收益", "最大回撤", "卡玛比率", "交易笔数", "胜率", "期末净值倍数"]
          if c in risk.columns]].to_csv(CSVDIR / "止损与仓位.csv", index=False, encoding="utf-8-sig")
    delist.to_csv(CSVDIR / "退市股交易.csv", index=False, encoding="utf-8-sig")
    ov.to_csv(CSVDIR / "前二重合度.csv", index=False, encoding="utf-8-sig")


if __name__ == "__main__":
    main()
