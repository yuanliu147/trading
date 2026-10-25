# -*- coding: utf-8 -*-
"""重复信号与买入时点：交付校验。"""
from pathlib import Path
import json

import numpy as np
import pandas as pd
import openpyxl

HERE = Path(__file__).resolve()
OUT = HERE.parents[1]
RES = OUT / '03_结果'
FIG = OUT / '04_图表'
EV = OUT / '05_证据'
EV.mkdir(parents=True, exist_ok=True)

rows = []


def check(name, cond, got=''):
    rows.append((bool(cond), name, str(got)))


P = pd.read_csv(RES / '位置统计.csv')
PAIR = pd.read_csv(RES / '首笔vs第二笔.csv')
PAIRD = pd.read_csv(RES / '首笔vs第二笔_明细.csv')
BEST = pd.read_csv(RES / '最佳位置.csv')
DELAY = pd.read_csv(RES / '入场延迟.csv')
PF = pd.read_csv(RES / '组合入场延迟对照.csv')
F = pd.read_csv(RES / '特征分组.csv')
CUR = pd.read_csv(RES / '当前信号_250日回撤.csv')
CASE = pd.read_csv(RES / '603778案例.csv')
A = pd.read_csv(RES / '全部信号重复序列_完整.csv')
T = pd.read_csv(RES / '前二重复序列.csv')
M = json.loads((RES / '指标.json').read_text(encoding='utf-8'))

check('信号池4195', M['样本']['信号池'] == 4195, M['样本']['信号池'])
check('全部信号4195', M['样本']['全部信号'] == 4195, M['样本']['全部信号'])
check('信号日251', M['样本']['信号日'] == 251, M['样本']['信号日'])
check('股票1246', M['样本']['股票'] == 1246, M['样本']['股票'])
check('全部完整序列4019', len(A) == 4019, len(A))
check('全部完整4019', int(A['complete'].sum()) == 4019, int(A['complete'].sum()))
check('全部段2478', M['样本']['全部段'] == 2478, M['样本']['全部段'])
check('全部完整段2311', M['样本']['全部完整段'] == 2311, M['样本']['全部完整段'])
check('前二完整450', len(T) == 450, len(T))
check('前二段321', T.eid.nunique() == 321, T.eid.nunique())
check('配对数78', len(PAIRD) == 78, len(PAIRD))

p0 = P.set_index(['样本', '位置'])
check('前二第1次321/+4.74%', int(p0.loc[('前二', '1'), '笔数']) == 321 and abs(p0.loc[('前二', '1'), '平均'] - .047366) < 1e-5,
      p0.loc[('前二', '1'), '平均'])
check('前二第2次78/+6.79%', int(p0.loc[('前二', '2'), '笔数']) == 78 and abs(p0.loc[('前二', '2'), '平均'] - .067857) < 1e-5,
      p0.loc[('前二', '2'), '平均'])
check('前二第3次30/+6.96%', int(p0.loc[('前二', '3'), '笔数']) == 30 and abs(p0.loc[('前二', '3'), '平均'] - .069639) < 1e-5,
      p0.loc[('前二', '3'), '平均'])
check('前二第4+次21/+5.10%', int(p0.loc[('前二', '4+'), '笔数']) == 21 and abs(p0.loc[('前二', '4+'), '平均'] - .050983) < 1e-5,
      p0.loc[('前二', '4+'), '平均'])
check('全部第4+次330/+2.75%', int(p0.loc[('全部信号', '4+'), '笔数']) == 330 and abs(p0.loc[('全部信号', '4+'), '平均'] - .027455) < 1e-5,
      p0.loc[('全部信号', '4+'), '平均'])

pp = PAIR.set_index('样本')
check('配对全部78', int(pp.loc['全部', '配对数']) == 78, pp.loc['全部', '配对数'])
check('首笔+6.97%', abs(pp.loc['全部', '首笔均值'] - .069743) < 1e-5, pp.loc['全部', '首笔均值'])
check('第二笔+6.79%', abs(pp.loc['全部', '第二笔均值'] - .067857) < 1e-5, pp.loc['全部', '第二笔均值'])
check('差值-0.19pp', abs(pp.loc['全部', '差值均值'] + .001886) < 1e-5, pp.loc['全部', '差值均值'])
check('第二笔更好37.18%', abs(pp.loc['全部', '第二笔更好占比'] - .371795) < 1e-5, pp.loc['全部', '第二笔更好占比'])
check('相关0.7577', abs(pp.loc['全部', '相关性'] - .757655) < 1e-5, pp.loc['全部', '相关性'])

bb = BEST.set_index('best_pos')
check('最佳位置1=43/55.1%', int(bb.loc[1, '次数']) == 43 and abs(bb.loc[1, '占比'] - .551282) < 1e-5,
      (bb.loc[1, '次数'], bb.loc[1, '占比']))
check('最佳位置2=24/30.8%', int(bb.loc[2, '次数']) == 24 and abs(bb.loc[2, '占比'] - .307692) < 1e-5)
check('最佳位置3=7/9.0%', int(bb.loc[3, '次数']) == 7 and abs(bb.loc[3, '占比'] - .089744) < 1e-5)
check('最佳位置4=4/5.1%', int(bb.loc[4, '次数']) == 4 and abs(bb.loc[4, '占比'] - .051282) < 1e-5)

dd = DELAY[(DELAY['样本'] == '前二全部')].set_index('N')
for n, val in [(1, .049014), (2, .050455), (3, .051307), (4, .053831), (5, .050958)]:
    check(f'前二N={n} 平均{val*100:.1f}%', abs(dd.loc[n, '平均'] - val) < 1e-4, dd.loc[n, '平均'])
pf = PF.set_index('N')
check('组合N=3 40.0%', abs(pf.loc[3, '组合年化'] - .4001) < 1e-5, pf.loc[3, '组合年化'])
check('组合N=4 37.7%', abs(pf.loc[4, '组合年化'] - .377478) < 1e-5, pf.loc[4, '组合年化'])
check('组合N=1 23.7%', abs(pf.loc[1, '组合年化'] - .236816) < 1e-5, pf.loc[1, '组合年化'])

ft = F[F['样本'] == '前二'].set_index(['特征', '分组'])
check('前一日上涨+5.58%', abs(ft.loc[('买入日前一天涨跌', '前一日上涨'), '平均'] - .055805) < 1e-5,
      ft.loc[('买入日前一天涨跌', '前一日上涨'), '平均'])
check('前一日下跌+4.89%', abs(ft.loc[('买入日前一天涨跌', '前一日下跌'), '平均'] - .048877) < 1e-4,
      ft.loc[('买入日前一天涨跌', '前一日下跌'), '平均'])
check('MACD柱为负+5.91%', abs(ft.loc[('买入日前一天MACD柱', 'MACD柱为负'), '平均'] - .059125) < 1e-5)
check('MACD金叉+3.37%', abs(ft.loc[('买入日前一天MACD金叉', '前一日金叉'), '平均'] - .033674) < 1e-4)
check('MA20下方+5.40%', abs(ft.loc[('前一日收盘 vs MA20', 'MA20下方'), '平均'] - .053959) < 1e-4)
check('同向涨涨+5.97%', abs(ft.loc[('个股涨跌 × 市场涨跌', '个股涨+市场涨'), '平均'] - .059700) < 1e-4)
check('背离涨/市跌+3.74%', abs(ft.loc[('个股涨跌 × 市场涨跌', '个股涨+市场跌'), '平均'] - .037439) < 1e-4)

cc = CUR.set_index('代码')
check('603778回撤-78.61%', abs(cc.loc['603778.SH', '回撤_hfq'] + .786110) < 1e-5,
      cc.loc['603778.SH', '回撤_hfq'])
check('605376回撤-71.30%', abs(cc.loc['605376.SH', '回撤_hfq'] + .713027) < 1e-5,
      cc.loc['605376.SH', '回撤_hfq'])
check('002969回撤-70.68%', abs(cc.loc['002969.SZ', '回撤_hfq'] + .706788) < 1e-5,
      cc.loc['002969.SZ', '回撤_hfq'])
check('605117回撤-45.32%', abs(cc.loc['605117.SH', '回撤_hfq'] + .453227) < 1e-5,
      cc.loc['605117.SH', '回撤_hfq'])
check('605117实际回撤-61.07%', abs(cc.loc['605117.SH', '回撤_实际'] + .610678) < 1e-5,
      cc.loc['605117.SH', '回撤_实际'])
check('603778深于阈值33.61pp', abs(cc.loc['603778.SH', '深于45阈值_pp'] - 33.611) < 1e-3,
      cc.loc['603778.SH', '深于45阈值_pp'])

case_top = CASE[(CASE['是否前二'] == True) & CASE['收益'].notna()]
check('603778前二8笔', len(case_top) == 8, len(case_top))
check('603778前二均值+6.73%', abs(case_top['收益'].mean() - .067343) < 1e-5, case_top['收益'].mean())
check('603778前二胜率62.5%', abs((case_top['收益'] > 0).mean() - .625) < 1e-9)

wb = openpyxl.load_workbook(OUT / '重复信号与买入时点.xlsx', read_only=True)
check('Excel 18个sheet', len(wb.sheetnames) == 18, len(wb.sheetnames))
check('Excel含当前信号', '当前信号' in wb.sheetnames)
check('Excel含前二重复序列', '前二重复序列' in wb.sheetnames)
check('严格连续3天前二11', M['样本']['严格连续3天前二'] == 11, M['样本']['严格连续3天前二'])
check('严格连续3天全部181', M['样本']['严格连续3天全部'] == 181, M['样本']['严格连续3天全部'])
check('图表6张', len(list(FIG.glob('*.png'))) == 6, len(list(FIG.glob('*.png'))))

ok = sum(1 for r in rows if r[0])
lines = [f'重复信号与买入时点 · 交付校验  {ok}/{len(rows)} PASS', '']
for passed, name, got in rows:
    lines.append(f"{'PASS' if passed else 'FAIL'}  {name}" + (f"  [{got}]" if got else ''))
lines += ['', '数据源：同花顺本地 DuckDB 派生 research_panel.parquet；沪深主板等权 market_daily.parquet。',
          '口径：T+3 开盘、+20% 止盈、15 日到期、0.175%/边；逐只独立。']
text = '\n'.join(lines)
(EV / '校验.txt').write_text(text, encoding='utf-8')
print(text)
