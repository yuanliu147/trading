# -*- coding: utf-8 -*-
"""重复信号与买入时点：中文 Excel。"""
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve()
OUT = HERE.parents[1]
RES = OUT / '03_结果'

CUR = pd.read_csv(RES / '当前信号_250日回撤.csv', parse_dates=['信号日', '最高日_hfq', '最高日_实际'])
PLAN = pd.read_csv(RES / '当前执行方案.csv')
HOLD = pd.read_csv(RES / '当前持仓.csv', parse_dates=['signal_date', 'entry_date'])
P = pd.read_csv(RES / '位置统计.csv')
PAIR = pd.read_csv(RES / '首笔vs第二笔.csv')
PAIRD = pd.read_csv(RES / '首笔vs第二笔_明细.csv', parse_dates=['first_date', 'second_date'])
PAIRD = PAIRD.rename(columns={'eid': '配对编号', 'code': '代码', 'first_date': '首笔信号日',
                              'second_date': '第二笔信号日', 'first': '首笔收益',
                              'second': '第二笔收益', 'first_dd': '首笔回撤',
                              'second_dd': '第二笔回撤', 'y': '年份'})
DELAY = pd.read_csv(RES / '入场延迟.csv')
PF = pd.read_csv(RES / '组合入场延迟对照.csv')
F = pd.read_csv(RES / '特征分组.csv')
CASE = pd.read_csv(RES / '603778案例.csv', parse_dates=['信号日', '入场日', '退出日'])
MULTI = pd.read_csv(RES / '多发信号案例.csv', parse_dates=['信号日', '入场日', '退出日'])
T = pd.read_csv(RES / '前二重复序列.csv', parse_dates=['signal_date', 'entry_date', 'exit_date'])
A = pd.read_csv(RES / '全部信号重复序列_完整.csv', parse_dates=['signal_date', 'entry_date', 'exit_date'])
STRICT = pd.read_csv(RES / '严格连续3天.csv')
BEST = pd.read_csv(RES / '最佳位置.csv').rename(columns={'best_pos': '最佳位置', '次数': '次数',
                                                          '占比': '占比', 'ep_len': '段长',
                                                          'code': '代码', 'eid': '配对编号',
                                                          'best_ret': '最佳收益'})

cur = CUR.rename(columns={'信号日': '信号日', '排名': '排名', '代码': '代码',
                          '当日收盘_hfq': '当日收盘(后复权)', '最高250_hfq': '250日最高(后复权)',
                          '最高日_hfq': '250日最高日(后复权)', '回撤_hfq': '回撤(后复权)',
                          '当日收盘_实际': '当日收盘(实际价)', '最高250_实际': '250日最高(实际价)',
                          '跌价_hfq': '跌价(后复权)', '跌价_实际': '跌价(实际价)',
                          '最高日_实际': '250日最高日(实际价)', '回撤_实际': '回撤(实际价)',
                          '深于45阈值_pp': '深于-45%(pp)', '当日RSI': 'RSI14',
                          '收盘位置': '收盘位置', '量比': '量比(20日)', '涨幅250': '250日涨幅',
                          '当日信号数': '当日信号数'})

hold = HOLD.rename(columns={'thscode': '代码', 'signal_date': '信号日', 'entry_date': '入场日',
                            'entry_px': '入场价(后复权)', 'last_px': '期末价(后复权)',
                            'ret': '浮动收益', 'exit_reason': '状态'})

T2 = T.rename(columns={
    'signal_date': '信号日', 'code': '代码', 'pos': '段内位置', 'ep_len': '段长',
    'gap': '间隔交易日', 'rank': '排名', 'rk': '排名', 'dd250': '回撤250', 'entry_date': '入场日',
    'exit_date': '退出日', 'exit_reason': '退出原因', 'ret': '收益',
    'prev_up': '前一日上涨', 'amount_ratio1_20_prev': '前一日量比',
    'macd_hist_pct_prev': '前一日MACD柱', 'close_ma20_prev': '前一日距MA20',
    'market_prev_up': '市场前一日上涨', '阶段': '市场阶段'})
T2 = T2[['信号日', '代码', '段内位置', '段长', '间隔交易日', '排名', '回撤250',
         '入场日', '退出日', '退出原因', '收益', '前一日上涨', '前一日量比',
         '前一日MACD柱', '前一日距MA20', '市场前一日上涨', '市场阶段']]

A2 = A.rename(columns={
    'signal_date': '信号日', 'code': '代码', 'pos': '段内位置', 'ep_len': '段长',
    'gap': '间隔交易日', 'rk': '排名', 'top2': '是否前二', 'dd250': '回撤250',
    'entry_date': '入场日', 'exit_date': '退出日', 'exit_reason': '退出原因',
    'ret': '收益', 'status': '状态', 'complete': '完整',
    'prev_up': '前一日上涨', 'amount_ratio1_20_prev': '前一日量比',
    'macd_hist_pct_prev': '前一日MACD柱', 'close_ma20_prev': '前一日距MA20',
    'market_prev_up': '市场前一日上涨', '阶段': '市场阶段'})
A2 = A2[['信号日', '代码', '段内位置', '段长', '间隔交易日', '排名', '是否前二',
         '回撤250', '状态', '完整', '入场日', '退出日', '退出原因', '收益',
         '前一日上涨', '前一日量比', '前一日MACD柱', '前一日距MA20',
         '市场前一日上涨', '市场阶段']]

case = CASE.rename(columns={'信号日': '信号日', '排名': '排名', '是否前二': '是否前二',
                            '段内位置': '段内位置', '回撤250': '回撤250', 'RSI': 'RSI14',
                            '收盘位置': '收盘位置', '入场日': '入场日', '退出日': '退出日',
                            '收益': '收益', '状态': '状态'})

summary = pd.DataFrame([
    ('配对', '前二 78 对可比：首笔均值 +6.97%，第二笔 +6.79%，差值 -0.19pp；第二次更好仅 37.2%，相关 0.76。'),
    ('最佳位置', '多日信号中，事后最好的是第一次 55.1%，第二次 30.8%，第三次 9.0%，第四次 5.1%。'),
    ('全部信号', '第一次 +5.60%、第二次 +6.20%、第三次 +5.10%、第四次及以后 +2.75%；等太久变差。'),
    ('前一日涨跌', '前二：上涨 +5.58% vs 下跌 +4.89%；剔除 2026 +6.75% vs +5.42%；2026 反向（+0.38% vs +2.64%），不稳定。'),
    ('量能', '前二量比五分位 U 形：Q1 +6.68%、Q5 +6.43%，中间 Q2/Q4 更低；没有单调优势。'),
    ('MACD', '前二 MACD 柱为正 +3.97% vs 为负 +5.91%；金叉 +3.37% vs 非金叉 +5.45%。金叉是“晚了”，不是买点。'),
    ('MA20', '前二站上 MA20 +2.68% vs 下方 +5.40%（仅 23 笔在上面）；全部信号 +0.02% vs +5.59%。别追已修复的。'),
    ('最稳交互', '个股与市场同向：涨涨 +5.97%、跌跌 +5.72%；背离：涨/市跌 +3.74%、跌/市涨 +3.51%。剔除 2026 仍成立。'),
    ('入场延迟', '前二单笔 N=1..5：+4.90/+5.05/+5.13/+5.38/+5.10%；组合年化 N=3 最高 40.0%，N=4 37.7% 但回撤更大。'),
    ('当前', '603778 -78.61%、605376 -71.30%、002969 -70.68%、605117 -45.32%（后复权）；603778 已持有，不重复买。'),
], columns=['项目', '结论'])

notes = pd.DataFrame([
    ('数据源', '同花顺本地 DuckDB 派生 research_panel.parquet；沪深主板等权 market_daily.parquet。'),
    ('样本', '4195 条信号池；4019 条完整；2478 段股票信号序列；前二 450 条完整、321 段。'),
    ('重复段定义', '同一股票两次信号间隔 <=3 个交易日，算同一段“持续信号”；段内位置=第几次信号。'),
    ('前二', '每个信号日按个股 250 日回撤最深取前二；只有前二才是策略实际买入候选。'),
    ('收益口径', 'T+3 开盘买、+20% 止盈、15 日到期、0.175%/边；逐只独立，不模拟 2 仓占位。'),
    ('特征口径', '买入日前一天（E-1）的 K 线/量能/MACD/均线，以及 E-1 市场涨跌；E-1 收盘在 E 日开盘前已知。'),
    ('当前信号', '2026-10-22/23；后复权回撤是信号口径；实际价回撤仅用于核对，605117 因 2025/2026 送转，实际价回撤失真。'),
    ('重要警告', '信号日高度重叠；配对样本只有 78 对；阶段/特征差异是描述性的，不是可实盘的择时规则。'),
    ('幸存者偏差', '股票池为当前上市主板，未含退市股；熊市年份可能高估。'),
    ('历史不代表未来', '先纸面 50 笔；任何新过滤都要独立样本外验证。'),
], columns=['项目', '说明'])

params = pd.DataFrame([
    ('信号条件', '250 日回撤<-45%；250 日涨幅>0；RSI14<35；收盘位置>0.45；市场 20 日<-5%'),
    ('排序', '按个股 250 日回撤最深取前二'),
    ('仓位', '2 仓，每仓 50% 当前权益；无信号空仓'),
    ('入场', '信号后第 3 个交易日开盘（N=3）'),
    ('退出', '+20% 盘中止盈；15 个交易日收盘到期；无普通止损'),
    ('成本', '0.175%/边'),
    ('数据截止', '2026-10-23'),
], columns=['参数', '设定'])

path = OUT / '重复信号与买入时点.xlsx'
with pd.ExcelWriter(path, engine='openpyxl') as writer:
    notes.to_excel(writer, sheet_name='说明', index=False)
    summary.to_excel(writer, sheet_name='结论摘要', index=False)
    cur.to_excel(writer, sheet_name='当前信号', index=False)
    PLAN.to_excel(writer, sheet_name='执行方案', index=False)
    hold.to_excel(writer, sheet_name='当前持仓', index=False)
    P.to_excel(writer, sheet_name='位置统计', index=False)
    BEST.to_excel(writer, sheet_name='最佳位置', index=False)
    STRICT.to_excel(writer, sheet_name='严格连续3天', index=False)
    PAIR.to_excel(writer, sheet_name='首笔vs第二笔', index=False)
    PAIRD.to_excel(writer, sheet_name='首笔vs第二笔明细', index=False)
    DELAY.to_excel(writer, sheet_name='入场延迟', index=False)
    PF.to_excel(writer, sheet_name='组合延迟对照', index=False)
    F.to_excel(writer, sheet_name='特征分组', index=False)
    case.to_excel(writer, sheet_name='603778案例', index=False)
    MULTI.to_excel(writer, sheet_name='多发信号案例', index=False)
    T2.to_excel(writer, sheet_name='前二重复序列', index=False)
    A2.to_excel(writer, sheet_name='全部重复序列', index=False)
    params.to_excel(writer, sheet_name='口径与参数', index=False)

    wb = writer.book
    head_fill = PatternFill('solid', fgColor='1F4E79')
    head_font = Font(color='FFFFFF', bold=True, size=10)
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    date_headers = ('信号日', '入场日', '退出日', '250日最高日(后复权)', '250日最高日(实际价)',
                    '首笔信号日', '第二笔信号日')
    pct_tokens = ('收益', '胜率', '平均', '中位', '回撤', '占比', '涨跌', '涨幅', 'MACD柱',
                  '距MA20', '浮动', '均值', '年化')
    int_headers = ('笔数', '段长', '间隔交易日', '排名', '当日信号数', '起配对数', '配对数',
                   '次数', 'N', '交易笔数', '起点数')
    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            header = str(ws.cell(1, col).value)
            width = max(10, min(34, max(len(str(ws.cell(r, col).value or '')) for r in range(1, min(ws.max_row, 300) + 1)) + 2))
            ws.column_dimensions[letter].width = width
            fmt = None
            if header in date_headers:
                fmt = 'yyyy-mm-dd'
            elif header == '深于-45%(pp)':
                fmt = '0.00"pp"'
            elif header == '当日收盘(后复权)' or header == '250日最高(后复权)' or \
                    header == '当日收盘(实际价)' or header == '250日最高(实际价)' or \
                    header == '入场价(后复权)' or header == '期末价(后复权)' or \
                    header == '跌价(后复权)' or header == '跌价(实际价)':
                fmt = '0.000'
            elif header in int_headers:
                fmt = '0'
            elif any(t in header for t in pct_tokens):
                fmt = '0.00%'
            elif header == '量比(20日)' or header == '前一日量比':
                fmt = '0.00'
            elif header == 'RSI14':
                fmt = '0.0'
            elif header == '收盘位置' or header == '前一日距MA20':
                fmt = '0.00%' if header == '前一日距MA20' else '0.000'
            elif header == '相关性':
                fmt = '0.000'
            if fmt:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, col).number_format = fmt
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical='center')

    for name in ('说明', '结论摘要', '执行方案', '口径与参数'):
        ws = wb[name]
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 100
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)

    wb['当前信号'].freeze_panes = 'B2'
    wb['前二重复序列'].freeze_panes = 'B2'

print('written', path)
