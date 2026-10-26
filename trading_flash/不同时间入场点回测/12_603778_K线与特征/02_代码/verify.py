# -*- coding: utf-8 -*-
"""603778 K 线与特征：交付校验。"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve()
OUT = HERE.parents[1]
RES = OUT / '03_结果'
FIG = OUT / '04_图表'
REPORT = OUT / '05_证据' / '校验.txt'

checks: list[tuple[bool, str]] = []


def check(cond, msg):
    checks.append((bool(cond), msg))


def approx(a, b, tol=1e-4):
    return abs(float(a) - float(b)) <= tol


# 数据
panel = '/Users/admin/WorkSpace/trading/trading_flash/2026-10-24_主板恐慌反转_年化30/01_数据/research_panel.parquet'
cols = ['thscode', 'date', 'open', 'high', 'low', 'close', 'hfq_close', 'drawdown250',
        'rsi14', 'close_location', 'macd_hist_pct', 'amount_ratio1_20', 'close_ma20',
        'return_250d', 'return_60d', 'return_20d']
f = pd.read_parquet(panel, columns=cols)
f['date'] = pd.to_datetime(f['date'])
g = f[f['thscode'] == '603778.SH'].sort_values('date').reset_index(drop=True)
cur = g[g['date'] == '2026-10-23'].iloc[0]
check(len(g) == 2151, f'603778 日频 2,151 行（实际 {len(g)}）')
check(approx(cur['close'], 7.33, 5e-4), '10/23 实际收盘 7.33')
check(approx(cur['hfq_close'], 18.4647, 5e-4), '10/23 后复权收盘 18.4647')
check(approx(cur['drawdown250'], -0.78611, 5e-4), '250 日回撤 -78.6%')
check(approx(cur['rsi14'], 32.2411, 5e-3), 'RSI14 32.24')
check(approx(cur['close_location'], 0.924242, 5e-4), '收盘位置 0.9242')
check(approx(cur['macd_hist_pct'], -0.005654, 5e-4), 'MACD 柱 -0.00565')
check(approx(cur['amount_ratio1_20'], 0.807953, 5e-4), '量比1/20 0.808')
check(approx(cur['close_ma20'], -0.092878, 5e-4), '距MA20 -9.29%')
check(approx(cur['return_250d'], 1.106322, 5e-4), '250日涨幅 +110.6%')
# BOR
g['prior_low'] = g['low'].shift(1)
g['prior_close'] = g['close'].shift(1)
g['BOR'] = (g['low'] < g['prior_low']) & (g['close'] > g['prior_close']) & (g['close'] > g['open'])
check(bool(g.loc[g['date'] == '2026-10-23', 'BOR'].iloc[0]), '10/23 BOR=True')
check(bool(g.loc[g['date'] == '2026-07-21', 'BOR'].iloc[0]), '7/21 BOR=True')
check(bool(g.loc[g['date'] == '2026-07-02', 'BOR'].iloc[0]), '7/2 BOR=True（浅回撤输家）')
check(not bool(g.loc[g['date'] == '2026-07-08', 'BOR'].iloc[0]), '7/8 BOR=False')

# 10/26 partial
s26 = pd.read_csv('/Users/admin/WorkSpace/trading/trading_flash/2026-10-24_主板恐慌反转_年化30/03_结果/当前信号_2026-10-26.csv')
r26 = s26[s26['thscode'] == '603778.SH'].iloc[0]
check(int(r26['rank']) == 1, '10/26 603778 排名 1')
check(approx(r26['actual_close'], 7.29, 5e-4), '10/26 收盘 7.29')
check(approx(r26['close_location'], 0.454545, 5e-4), '10/26 收盘位置 0.4545')
check(approx(r26['rsi14'], 31.7976, 5e-3), '10/26 RSI 31.8')
check(approx(r26['drawdown250'], -0.787278, 5e-4), '10/26 250日回撤 -78.7%')

# 信号表
S = pd.read_csv(RES / '信号日特征.csv', parse_dates=['信号日'])
check(len(S) == 9, f'603778 信号表 9 行（实际 {len(S)}）')
check(S['排名'].eq(1).sum() == 8, '8 次榜一')
check(S['排名'].eq(2).sum() == 1, '1 次第 2（07-23）')
for d, ret in [('2026-07-02', -0.140516), ('2026-07-08', -0.188433), ('2026-07-21', 0.195804),
               ('2026-10-15', -0.026084)]:
    v = S.loc[S['信号日'] == pd.Timestamp(d), '15日收益'].iloc[0]
    check(approx(v, ret, 5e-4), f'{d} 15日收益与 06 一致')
check(bool(S.loc[S['信号日'] == pd.Timestamp('2026-10-23'), 'BOR'].iloc[0]), '信号表 10/23 BOR=True')

# 条件
C = pd.read_csv(RES / '条件统计.csv')
b17 = C[C['条件'] == '深回撤 + BOR + 高收盘'].iloc[0]
check(int(b17['样本']) == 17, '深回撤+BOR+高收盘 17 笔')
check(approx(b17['均值'], 0.144846, 5e-4), '17 笔均值 +14.5%')
check(approx(b17['中位'], 0.195804, 5e-4), '17 笔中位 +19.6%')
check(approx(b17['胜率'], 1.0, 5e-6), '17 笔胜率 100%')
check(approx(b17['最差'], 0.014179, 5e-4), '17 笔最差 +1.4%')
b26 = C[C['条件'] == '深回撤 + BOR'].iloc[0]
check(int(b26['样本']) == 26 and approx(b26['胜率'], 0.923077, 5e-4), '深回撤+BOR 26 笔/92.3%')
check(int(C[C['条件'] == '深回撤 + 收盘位置≤0.8']['样本'].iloc[0]) == 73, '深回撤+低收盘 73 笔')

# 关键价位
K = pd.read_csv(RES / '关键价位.csv')
for name, price in [('历史高点', 34.27), ('7月低点', 7.56), ('10月低点', 6.72),
                    ('回撤 88.6%', 6.11), ('+20%目标（按10/23收盘）', 8.80),
                    ('回撤 78.6%', 9.29), ('MA60', 9.31), ('回撤 100%', 2.49)]:
    v = K.loc[K['类型'] == name, '价格']
    check(len(v) and approx(v.iloc[0], price, 5e-3), f'关键价位 {name} ≈ {price}')

# 相对强弱
mkt = pd.read_parquet('/Users/admin/WorkSpace/trading/trading_hithink/初始研究/data/work/market_daily.parquet',
                      columns=['date', 'equal_weight_index'])
mkt['date'] = pd.to_datetime(mkt['date'])
mg = g.set_index('date')['hfq_close']
mm = mkt.set_index('date')['equal_weight_index']
for a, b, sret, mret in [('2026-04-09', '2026-10-23', -0.7768, -0.1124),
                          ('2026-07-21', '2026-10-23', -0.1447, 0.0377),
                          ('2026-09-04', '2026-10-23', -0.3207, -0.0688)]:
    s = mg.loc[b] / mg.loc[a] - 1
    m = mm.loc[b] / mm.loc[a] - 1
    check(approx(s, sret, 5e-4), f'{a}→{b} 个股 {sret:.1%}')
    check(approx(m, mret, 5e-4), f'{a}→{b} 市场 {mret:.1%}')

# 文件
for name in ['01_两年K线.png', '02_近月K线.png', '03_信号日散点.png', '04_条件收益.png', '05_相对强弱.png']:
    p = FIG / name
    check(p.exists() and p.stat().st_size > 50_000, f'图表 {name} 存在且 >50KB')
html = OUT / '603778_K线与特征.html'
check(html.exists() and html.stat().st_size > 800_000, 'HTML 存在且 >800KB')
h = html.read_text(encoding='utf-8')
check(h.count('data:image/png;base64,') == 5, 'HTML 内嵌 5 张图')
for s in ['17 笔全为正', '0.924', '0.4545', 'BOR']:
    check(s in h, f'HTML 含 {s}')
xlsx = OUT / '603778_K线与特征.xlsx'
check(xlsx.exists(), 'Excel 存在')
wb = load_workbook(xlsx)
check(wb.sheetnames == ['说明', '信号日特征', '条件统计', '关键价位', '当前快照', '相对强弱'],
      'Excel 6 个工作表')
for name in ['00_一页纸_603778_K线与特征.md', '01_完整报告_603778_K线与特征.md', '03_数据与复现.md']:
    check((OUT / name).exists(), f'{name} 存在')

passed = sum(1 for ok, _ in checks if ok)
lines = ['603778 K 线与特征 · 交付校验', f'日期：2026-10-23', '', f'通过：{passed}/{len(checks)}', '']
for ok, msg in checks:
    lines.append(('✓' if ok else '✗') + ' ' + msg)
lines += ['', '数据：research_panel.parquet、06 逐日首笔、07 全部信号、market_daily。',
          '注意：BOR 17 笔为样本内、事后定义；10/26 OHLC 未入库。']
REPORT.write_text('\n'.join(lines), encoding='utf-8')
print('\n'.join(lines))
if passed != len(checks):
    raise SystemExit(1)
