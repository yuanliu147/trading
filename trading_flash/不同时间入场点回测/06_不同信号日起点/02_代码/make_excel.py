# -*- coding: utf-8 -*-
"""不同信号日起点：中文 Excel。"""
from pathlib import Path
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve()
OUT = HERE.parents[1]
RES = OUT / '03_结果'
XLSX = OUT / '不同信号日起点回测.xlsx'

C = pd.read_csv(RES / '起点汇总.csv',
                parse_dates=['signal_date', 'scheduled_entry', 'first_entry', 'last_exit'])
B = pd.read_csv(RES / '逐日首笔.csv',
                parse_dates=['signal_date', 'scheduled_entry', 'first_entry', 'last_exit'])
DET = pd.read_csv(RES / '逐日首笔_逐只.csv',
                  parse_dates=['signal_date', 'entry_date', 'exit_date'])
P = pd.read_csv(RES / '相邻信号日.csv', parse_dates=['date_a', 'date_b'])
TR = pd.read_csv(RES / '全部起点_逐笔.csv',
                 parse_dates=['起点', 'signal_date', 'entry_date', 'exit_date'])
IDX = pd.read_csv(RES / '每信号日交易记录' / '00_索引.csv')
D = pd.read_csv(RES / '起点分布.csv')
Y = pd.read_csv(RES / '起点年度.csv')
M = json.load(open(RES / '指标.json', encoding='utf-8'))

# 补充 P 的买入只数（兼容旧结果）
if 'n_buy_a' not in P.columns:
    nb = C[['signal_date', 'n_buy']]
    P = P.merge(nb.rename(columns={'signal_date': 'date_a', 'n_buy': 'n_buy_a'}),
                on='date_a', how='left')
    P = P.merge(nb.rename(columns={'signal_date': 'date_b', 'n_buy': 'n_buy_b'}),
                on='date_b', how='left')
P['both_buy2'] = P['n_buy_a'].eq(2) & P['n_buy_b'].eq(2)

# ---------------------------------------------------------------- 中文表头
C = C.rename(columns={
    'signal_date': '起点信号日', 'n_signals': '当日信号数',
    'first_code': '第一名代码', 'first_dd250': '第一名250日回撤(%)',
    'second_code': '第二名代码', 'second_dd250': '第二名250日回撤(%)',
    'scheduled_entry': '计划入场日', 'first_entry': '首笔入场日',
    'last_exit': '首笔最后退出日', 'n_buy': '首笔买入只数',
    'batch_complete': '首笔是否完整', 'batch_ret': '首笔组合收益(%)',
    'batch_float': '首笔浮动收益(%)', 'batch_avg': '首笔平均单笔(%)',
    'final': '期末倍数', 'total_ret': '总收益(%)', 'cagr': '全期年化(%)',
    'dd': '最大回撤(%)', 'n': '笔数', 'win': '胜率(%)',
    'avg': '平均单笔(%)', 'median': '中位单笔(%)', 'best': '最好单笔(%)',
    'worst': '最差单笔(%)', 'pf': '利润因子', 'expo': '持仓时间占比(%)',
    'n_target': '止盈笔数', 'n_time': '到期笔数', 'n_eod': '期末未平笔数',
    'profit_100k': '总利润(元)', 'ret_1年': '1年收益(%)',
    'dd_1年': '1年最大回撤(%)', 'ret_3年': '3年收益(%)',
    'dd_3年': '3年最大回撤(%)', 'ret_5年': '5年收益(%)',
    'dd_5年': '5年最大回撤(%)', '样本': '样本',
})
C['首笔是否完整'] = C['首笔是否完整'].map({True: '是', False: '否'})

B = B.rename(columns={
    'signal_date': '信号日', 'n_signals': '当日信号数',
    'first_code': '第一名代码', 'first_dd250': '第一名250日回撤(%)',
    'second_code': '第二名代码', 'second_dd250': '第二名250日回撤(%)',
    'scheduled_entry': '计划入场日', 'first_entry': '首笔入场日',
    'last_exit': '首笔最后退出日', 'n_buy': '首笔买入只数',
    'n_complete': '首笔完整只数', 'batch_complete': '首笔是否完整',
    'batch_ret': '首笔组合收益(%)', 'batch_float': '首笔浮动收益(%)',
    'batch_avg': '首笔平均单笔(%)', 'codes': '首笔代码',
    'rets': '首笔单笔收益', 'status': '状态',
})
B['首笔是否完整'] = B['首笔是否完整'].map({True: '是', False: '否'})
B['状态'] = (B['状态'].replace({'ok': '已完成', '数据结束': '尚未入场'})
             .str.replace('持有未满', '持有未满', regex=False))

DET = DET.rename(columns={
    'signal_date': '信号日', 'n_signals': '当日信号数', 'rank': '排名',
    'code': '代码', 'dd250': '250日回撤(%)', 'rsi14': '相对强弱14',
    'close_loc': '收盘位置', 'status': '状态', 'complete': '是否完整',
    'entry_date': '入场日', 'exit_date': '退出日', 'exit_reason': '退出原因',
    'ret': '单笔收益(%)', 'pre_ret': '入场前涨幅(%)',
})
DET['是否完整'] = DET['是否完整'].map({True: '是', False: '否'})
DET['状态'] = (DET['状态'].replace({'ok': '已完成', '数据结束': '尚未入场'})
               .str.replace('持有未满', '持有未满', regex=False))
DET['退出原因'] = DET['退出原因'].map(
    {'target': '+20%止盈', 'time': '15日到期', 'eod': '期末未平'}
).fillna(DET['退出原因'])

P = P.rename(columns={
    'date_a': '信号日A', 'date_b': '信号日B',
    'gap_td': '间隔交易日', 'gap_cd': '间隔自然日',
    'n_a': 'A当日信号数', 'n_b': 'B当日信号数',
    'n_buy_a': 'A首笔买入只数', 'n_buy_b': 'B首笔买入只数',
    'top_a': 'A前二', 'top_b': 'B前二',
    'overlap': '前二重合只数', 'top_jac': '前二重合度(%)',
    'same_top': '前二完全相同',
    'batch_a': 'A首笔收益(%)', 'batch_b': 'B首笔收益(%)',
    'batch_diff': '首笔差(百分点)',
    'cagr_a': 'A全期年化(%)', 'cagr_b': 'B全期年化(%)',
    'cagr_diff': '年化差(百分点)',
    'y1_a': 'A首年收益(%)', 'y1_b': 'B首年收益(%)', 'y1_diff': '首年差(百分点)',
    'y3_a': 'A三年收益(%)', 'y3_b': 'B三年收益(%)', 'y3_diff': '三年差(百分点)',
    'y5_a': 'A五年收益(%)', 'y5_b': 'B五年收益(%)', 'y5_diff': '五年差(百分点)',
    'trade_overlap': '交易重合数', 'trade_jac': '交易重合度(%)',
    'both_ge2': '两日信号都≥2只', 'both_buy2': '两日各买满2只',
})
for c in ('前二完全相同', '两日信号都≥2只', '两日各买满2只'):
    P[c] = P[c].map({True: '是', False: '否'})
P['配对类型'] = np.where(P['间隔交易日'] == 1,
                     '次日', '隔' + P['间隔交易日'].astype(str) + '个交易日')
for c in ('首笔差(百分点)', '年化差(百分点)', '首年差(百分点)',
          '三年差(百分点)', '五年差(百分点)'):
    P[c] = P[c] * 100.0
P = P[['信号日A', '信号日B', '配对类型', '间隔交易日', '间隔自然日',
       'A当日信号数', 'B当日信号数', 'A首笔买入只数', 'B首笔买入只数',
       'A前二', 'B前二', '前二重合只数', '前二重合度(%)', '前二完全相同',
       'A首笔收益(%)', 'B首笔收益(%)', '首笔差(百分点)',
       'A全期年化(%)', 'B全期年化(%)', '年化差(百分点)',
       'A首年收益(%)', 'B首年收益(%)', '首年差(百分点)',
       'A三年收益(%)', 'B三年收益(%)', '三年差(百分点)',
       'A五年收益(%)', 'B五年收益(%)', '五年差(百分点)',
       '交易重合数', '交易重合度(%)', '两日信号都≥2只', '两日各买满2只']]

D = D.rename(columns={
    '样本': '样本', '指标': '指标', 'n': '样本数', 'mean': '均值',
    'p5': '第5百分位', 'p25': '第25百分位', 'median': '中位数',
    'p75': '第75百分位', 'p95': '第95百分位', 'min': '最小值',
    'max': '最大值', 'neg': '负收益占比(%)',
    'ge30': '≥30%占比(%)', 'ge40': '≥40%占比(%)',
})
for c in ('≥30%占比(%)', '≥40%占比(%)'):
    D.loc[D['指标'].str.contains('最大回撤'), c] = np.nan

Y = Y.rename(columns={
    '起点年份': '起点年份', '起点数': '起点数',
    '满1年数': '满1年数', '满3年数': '满3年数', '满5年数': '满5年数',
    '年化中位数': '全期年化中位数(%)', '年化均值': '全期年化均值(%)',
    '年化负收益占比': '全期年化负收益占比(%)',
    '全期总收益中位数': '全期总收益中位数(%)',
    '1年收益中位数': '1年收益中位数(%)',
    '3年收益中位数': '3年收益中位数(%)',
    '5年收益中位数': '5年收益中位数(%)',
    '首笔批次收益中位数': '首笔批次收益中位数(%)',
})
Y['备注'] = Y['起点年份'].map({
    2017: '满5年（2017为部分年度）', 2018: '满5年', 2019: '满5年',
    2020: '满5年', 2021: '满5年（仅5个满5年）', 2022: '满3年',
    2023: '5个起点，仅2个满3年', 2024: '满1年，全期约2.3年',
    2025: '满1年，全期约1.3年', 2026: '不足1年，年化不可比',
}).fillna('')

# 最近案例
R = C.tail(10).copy()
R['首笔状态'] = np.where(R['首笔入场日'].isna(), '尚未入场',
                     np.where(R['首笔是否完整'].eq('是'), '已完成', '持有未满'))
R['说明'] = ''
for i in R.index:
    d = R.loc[i, '起点信号日']
    if d == pd.Timestamp('2026-10-23'):
        R.loc[i, '说明'] = ('123只候选；前二 603778.SH、605376.SH；'
                            '计划 10-28 开盘；尚未入场')
    elif d == pd.Timestamp('2026-10-22'):
        R.loc[i, '说明'] = ('仅1只候选 605117.SH；计划 10-27 开盘；'
                            '与 10-23 是两批不同信号')
    elif d == pd.Timestamp('2026-10-15'):
        R.loc[i, '说明'] = '22只候选；前二 603778.SH、603986.SH；持有未满'
    elif d == pd.Timestamp('2026-10-16'):
        R.loc[i, '说明'] = '13只候选；前二 000426.SZ、600879.SH；持有未满'
R = R[['起点信号日', '当日信号数', '第一名代码', '第一名250日回撤(%)',
       '第二名代码', '第二名250日回撤(%)', '计划入场日', '首笔入场日',
       '首笔买入只数', '首笔浮动收益(%)', '首笔状态', '期末倍数',
       '全期年化(%)', '1年收益(%)', '3年收益(%)', '5年收益(%)', '说明']]

# 全部起点逐笔
TR['退出价(后复权)'] = TR['exit_fill'].fillna(TR['last_px'])
TR['退出原因'] = TR['exit_reason'].map(
    {'target': '+20%止盈', 'time': '15日到期', 'eod': '期末未平'}
).fillna(TR['exit_reason'])
TR['分配额(初始本金%)'] = TR['alloc'] * 100
TR['仓位占当时权益(%)'] = TR['alloc_frac'] * 100
TR['单笔收益(%)'] = TR['ret'] * 100
TR['盈亏贡献(初始本金%)'] = TR['alloc'] * TR['ret'] * 100
TR['250日回撤(%)'] = TR['drawdown250'] * 100
TR['市场20日(%)'] = TR['market_return20'] * 100
TR = TR.rename(columns={
    '起点': '起点信号日', 'signal_date': '信号日', 'thscode': '代码',
    'entry_date': '入场日', 'exit_date': '退出日', 'days': '持有交易日',
    'entry_px': '入场价(后复权)', 'rk': '信号排名', 'rsi14': '相对强弱14',
    'close_location': '收盘位置',
})
TR = TR[['起点信号日', '信号日', '代码', '入场日', '退出日', '退出原因',
         '持有交易日', '入场价(后复权)', '退出价(后复权)',
         '分配额(初始本金%)', '仓位占当时权益(%)', '单笔收益(%)',
         '盈亏贡献(初始本金%)', '信号排名', '250日回撤(%)',
         '相对强弱14', '收盘位置', '市场20日(%)']]
TR = TR.sort_values(['起点信号日', '入场日', '代码']).reset_index(drop=True)

# ---------------------------------------------------------------- 结论摘要
def getv(sample, metric, field):
    q = D[(D['样本'] == sample) & (D['指标'] == metric)]
    return float(q.iloc[0][field])


m1_cagr = getv('满1年', '全期年化', '中位数')
m1_y1 = getv('满1年', '1年收益', '中位数')
m1_y1min = getv('满1年', '1年收益', '最小值')
m1_y1max = getv('满1年', '1年收益', '最大值')
m3_cagr = getv('满3年', '全期年化', '中位数')
m3_min = getv('满3年', '全期年化', '最小值')
m3_max = getv('满3年', '全期年化', '最大值')
m5_cagr = getv('满5年', '全期年化', '中位数')
m5_min = getv('满5年', '全期年化', '最小值')
m5_max = getv('满5年', '全期年化', '最大值')
m5_mult_min = 1 + getv('满5年', '5年收益', '最小值')
m5_mult_max = 1 + getv('满5年', '5年收益', '最大值')
b_med = getv('全部信号日', '首笔批次收益', '中位数')
b_mean = getv('全部信号日', '首笔批次收益', '均值')
b_p5 = getv('全部信号日', '首笔批次收益', '第5百分位')
b_p95 = getv('全部信号日', '首笔批次收益', '第95百分位')
b_neg = getv('全部信号日', '首笔批次收益', '负收益占比(%)')

summary_rows = [
    ('样本', '4,195 条信号；251 个信号日；208 个信号日 ≥2 只；'
             '250 对相邻信号日；158 对“次日”；114 对“次日且两日都 ≥2 只”'
             '（其中 113 对各买满 2 只）'),
    ('最早起点', '2017-11-20：99 笔，期末 20.75 倍，年化 40.47%；'
                 '母报告从 2017-10-19 起算是 40.0%，同一批 99 笔、同一终值'),
    ('满1年（202个起点）',
     f'全期年化中位 {m1_cagr*100:.1f}%；首年收益中位 {m1_y1*100:.1f}%，'
     f'范围 {m1_y1min*100:.1f}%–{m1_y1max*100:.1f}%'),
    ('满3年（146个起点）',
     f'全期年化中位 {m3_cagr*100:.1f}%，范围 {m3_min*100:.1f}%–{m3_max*100:.1f}%'),
    ('满5年（101个起点）',
     f'全期年化中位 {m5_cagr*100:.1f}%，范围 {m5_min*100:.1f}%–{m5_max*100:.1f}%；'
     f'5年期末 {m5_mult_min:.2f}–{m5_mult_max:.2f} 倍'),
    ('首笔批次（246个完整批次）',
     f'中位 {b_med*100:.1f}%，均值 {b_mean*100:.1f}%，'
     f'5%–95% 分位 {b_p5*100:.1f}% 到 {b_p95*100:.1f}%；'
     f'负收益占 {b_neg*100:.1f}%'),
    ('次日配对：前二',
     '114 对两日都 ≥2 只：前二零重合 63 对、重合1只 38 对、完全相同 13 对；'
     '“第二天有新的两只”是常态，不是例外'),
    ('次日配对：首年差',
     '81 对可算首年：B−A 中位 -0.6 个百分点，5%–95% 分位 -23.1 到 +19.5；'
     '最大负差 -29.4，最大正差 +26.3'),
    ('次日配对：5年差',
     '36 对可算5年：B−A 中位 -9.1 个百分点，5%–95% 分位 -73.6 到 +70.8；'
     '绝对值中位 22.2 个百分点'),
    ('次日配对：交易重合',
     '114 对全期交易集合重合度中位 89.7%，5%分位 16.2%；'
     '第一笔不同，但后面的交易大多会重新汇合'),
    ('怎么读', '40% 不是某个起点的偶然；但起点决定你前1–5年看到什么。'
               '长期年化落在 31%–43%，不等于每年、每个起点都 40%。'),
    ('当前案例', '2026-10-22 只有 605117.SH，计划 10-27 开盘；'
                 '2026-10-23 有 123 只，前二 603778.SH、605376.SH，计划 10-28 开盘。'
                 '两批不同，不能互相替换。'),
    ('重要限制', '251个起点高度重叠，不是251次独立实验；后复权；不含退市股；'
                 '无盘中分时；未对251个起点逐一做实际价整手复核；'
                 '2024为异常大年；事后挑最好起点=过拟合。'),
]
summary = pd.DataFrame(summary_rows, columns=['项目', '内容'])

# ---------------------------------------------------------------- 起点最好最差
c3 = C[C['3年收益(%)'].notna()].copy()
c1 = C[C['1年收益(%)'].notna()].copy()
bestworst = []
for group, df, col, asc in [
        ('满3年·年化最低10', c3, '全期年化(%)', True),
        ('满3年·年化最高10', c3, '全期年化(%)', False),
        ('满1年·首年最低10', c1, '1年收益(%)', True),
        ('满1年·首年最高10', c1, '1年收益(%)', False)]:
    q = df.sort_values(col, ascending=asc).head(10).copy()
    q.insert(0, '分组', group)
    bestworst.append(q)
bestworst = pd.concat(bestworst, ignore_index=True)
bestworst = bestworst[['分组', '起点信号日', '当日信号数', '第一名代码', '第二名代码',
                       '计划入场日', '首笔组合收益(%)', '期末倍数', '全期年化(%)',
                       '最大回撤(%)', '1年收益(%)', '3年收益(%)', '5年收益(%)']]
bestworst.insert(1, '提示', '事后统计，不可用于挑起点')

# ---------------------------------------------------------------- 口径
params = pd.DataFrame([
    ('策略', '主板恐慌洗盘反转（母策略：trading_flash/2026-10-24_主板恐慌反转_年化30）'),
    ('本报告问题', 'A 在 m 日看到信号、B 在 m+1 日看到新信号，各买各的前二；'
                   '不是“同一个信号日延后 N 天买”'),
    ('起点定义', '每个“有信号的交易日”作为一个起点；当天开始、10万本金；'
                 '若当天只有1只，就只买1只（半仓），不递补'),
    ('入场', '信号后第 3 个交易日开盘（N=3，与母策略基准一致）'),
    ('仓位', '2 仓；每仓 50% 当前权益；无信号空仓'),
    ('持有', '15 个交易日（含入场日）'),
    ('止盈', '+20%（T+1 起，盘中触价；开盘跳空高于止盈按开盘）'),
    ('止损', '无普通止损'),
    ('成本', '0.175%/边（复权模拟）'),
    ('一字涨停', '入场实际价 ≥ 前收×1.098 不买；被跳过不递补'),
    ('停牌', '入场日无行情不买'),
    ('排序', '同日按 250 日回撤升序取前 2'),
    ('复权', '收益用后复权；涨跌停/停牌判断用不复权实际价'),
    ('数据', f"{M['start']} 至 {M['end']}；{M['n_days']:,} 个交易日；"
             f"{M['n_syms']:,} 只；{M['n_signals']:,} 条信号；"
             f"{M['n_signal_days']} 个信号日"),
    ('首年/3年/5年', '起点后 252/756/1260 个交易日的收盘市值，未强制平仓；'
                     '不是年化，是区间总收益'),
    ('期末', '2026-10-23 收盘；未平仓按最后收盘价计'),
    ('当前案例', '2026-10-22 → 计划 2026-10-27 开盘；'
                 '2026-10-23 → 计划 2026-10-28 开盘（工作日推算）'),
    ('CSV 单位', '收益/回撤/胜率为小数（0.40=40%）；Excel 已格式化为百分比；'
                 '差值在 CSV 为小数，Excel ×100 后标“个百分点”'),
    ('逐笔表字段', '起点信号日=你开始的那天；信号日=这笔交易对应的信号日'
                   '（可能更晚）；分配额(初始本金%)=该笔投入/10万，复利后可能超过100%；'
                   '仓位占当时权益(%) 才是不超过50%的当前权益比例'),
    ('未平仓差异', '单笔收益/盈亏贡献按双边成本；期末未平仓的账户市值只扣买入成本、'
                   '未扣卖出成本，所以未平仓部分的贡献之和会比账户总收益低约0.17%'),
    ('复现', 'cd 不同时间入场点回测 && python3 06_不同信号日起点/02_代码/run_cohorts.py '
             '&& python3 06_不同信号日起点/02_代码/make_charts.py '
             '&& python3 06_不同信号日起点/02_代码/make_excel.py '
             '&& python3 06_不同信号日起点/02_代码/verify_cohorts.py'),
    ('不构成', '历史回测不是投资建议；先纸面 50 笔，再小资金'),
], columns=['参数/项目', '值'])

# ---------------------------------------------------------------- 写
with pd.ExcelWriter(XLSX, engine='openpyxl') as writer:
    summary.to_excel(writer, sheet_name='结论摘要', index=False)
    C.to_excel(writer, sheet_name='起点汇总', index=False)
    D.to_excel(writer, sheet_name='起点分布', index=False)
    Y.to_excel(writer, sheet_name='起点年度', index=False)
    bestworst.to_excel(writer, sheet_name='起点最好最差', index=False)
    P.to_excel(writer, sheet_name='相邻信号日', index=False)
    # 宽表：相邻汇总
    rows = []
    for label, q in [
            ('全部相邻信号日', P),
            ('次日', P[P['配对类型'] == '次日']),
            ('次日且两日都≥2只', P[(P['配对类型'] == '次日') & (P['两日信号都≥2只'] == '是')]),
            ('次日且两日各买满2只', P[(P['配对类型'] == '次日') & (P['两日各买满2只'] == '是')])]:
        y1 = q['首年差(百分点)'].dropna()
        y3 = q['三年差(百分点)'].dropna()
        y5 = q['五年差(百分点)'].dropna()
        bd = q['首笔差(百分点)'].dropna()
        rows.append({
            '样本': label, '对数': len(q),
            '首笔可算对数': len(bd), '首年可算对数': len(y1),
            '三年可算对数': len(y3), '五年可算对数': len(y5),
            '前二零重合占比(%)': (q['前二重合只数'] == 0).mean(),
            '前二完全相同占比(%)': (q['前二完全相同'] == '是').mean(),
            '首笔差中位(百分点)': bd.median(),
            '首年差中位(百分点)': y1.median(),
            '首年差5%分位(百分点)': y1.quantile(.05),
            '首年差95%分位(百分点)': y1.quantile(.95),
            '首年差绝对值中位(百分点)': y1.abs().median(),
            '三年差中位(百分点)': y3.median(),
            '五年差中位(百分点)': y5.median(),
            '五年差绝对值中位(百分点)': y5.abs().median(),
            '交易重合度中位(%)': q['交易重合度(%)'].median(),
        })
    PS = pd.DataFrame(rows)
    PS.to_excel(writer, sheet_name='相邻信号日汇总', index=False)
    B.to_excel(writer, sheet_name='逐日首笔', index=False)
    DET.to_excel(writer, sheet_name='逐只首笔', index=False)
    TR.to_excel(writer, sheet_name='全部起点逐笔', index=False)
    IDX.to_excel(writer, sheet_name='每信号日索引', index=False)
    R.to_excel(writer, sheet_name='最近信号日', index=False)
    params.to_excel(writer, sheet_name='口径与参数', index=False)
    # 说明放最前
    notes = pd.DataFrame([
        ('标题', '不同信号日起点：A 在 m 日、B 在 m+1 日，各买各的前二'),
        ('一句话', '固定“信号后第3个交易日开盘买”，只换第一个信号日：'
                   '满3年的146个起点，年化 30.8%–51.3%（中位38.8%）；'
                   '满5年的101个，30.8%–42.7%（中位37.9%）。'
                   '但首年收益 1.3%–138.6%，相邻两天起点可差20–30个百分点。'),
        ('先分清', '上一份《不同入场时点》固定信号日、改变 N；'
                   '本报告固定 N=3、改变第一个信号日。'),
        ('看哪张表', '起点汇总=251个起点逐个结果；相邻信号日=250对A/B；'
                      '逐日首笔=每个信号日的第一笔；全部起点逐笔=每个起点的所有交易；'
                      '最近信号日=2026-10案例。'),
        ('逐日交易记录', '每信号日交易记录/YYYY-MM-DD.csv：251 个信号日各一份；'
                          '00_索引.csv 可按日期找文件；同目录有 zip。'),
        ('逐笔表字段', '全部起点逐笔：起点信号日=你开始的那天；信号日=这笔交易对应的信号日'
                       '（可能更晚）；分配额(初始本金%)=投入/10万，复利后可能>100%。'),
        ('分布表口径', '全部信号日：全期年化246（5个短样本/未入场不计算）；'
                       '全期总收益251（2个未入场按0计）；首笔批次246（完整批次）。'),
        ('数据', f"{M['start']} 至 {M['end']}；本地 DuckDB 派生面板；10万本金；"
                 f"4,195条信号；251个信号日"),
        ('重要限制', '251个起点高度重叠，不是251次独立实验；后复权；不含退市股；'
                     '无盘中分时；未对251个起点逐一做实际价整手复核；'
                     '2024为异常大年；事后挑最好起点=过拟合。'),
        ('不构成', '历史回测不是投资建议；先纸面50笔，再小资金。'),
    ], columns=['项目', '内容'])
    notes.to_excel(writer, sheet_name='说明', index=False)
    # 把说明移到第一个 sheet
    wb = writer.book
    wb.move_sheet('说明', offset=-(len(wb.sheetnames) - 1))

# ---------------------------------------------------------------- 格式
wb = load_workbook(XLSX)
HEAD_FILL = PatternFill('solid', fgColor='1F4E79')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
PCT = '0.0%'
PCT2 = '0.00%'
DEC2 = '0.00'
MULT = '0.00"倍"'
YUAN = '#,##0'
PP = '0.0"个百分点"'
DATE = 'yyyy-mm-dd'


def col_index(ws, name):
    for cell in ws[1]:
        if cell.value == name:
            return cell.column
    raise KeyError(name)


def style_sheet(ws, pct=(), pct2=(), mult=(), yuan=(), pp=(), date=(),
                num=(), dec2=(), width=None, freeze='A2', wrap=(), title_rows=()):
    if ws.max_row < 1:
        return
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center',
                                   wrap_text=True)
    if freeze:
        ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    headers = [c.value for c in ws[1]]
    for j, h in enumerate(headers, 1):
        if h is None:
            continue
        letter = get_column_letter(j)
        fmt = None
        if h in pct:
            fmt = PCT
        elif h in pct2:
            fmt = PCT2
        elif h in mult:
            fmt = MULT
        elif h in dec2:
            fmt = DEC2
        elif h in yuan:
            fmt = YUAN
        elif h in pp:
            fmt = PP
        elif h in date:
            fmt = DATE
        elif h in num:
            fmt = '0'
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                for c in row:
                    c.number_format = fmt
        if h in wrap:
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical='top')
        if width and h in width:
            ws.column_dimensions[letter].width = width[h]
    if ws.max_row > 1:
        for r in range(2, ws.max_row + 1):
            ws.row_dimensions[r].height = None


# 说明
ws = wb['说明']
for r in range(1, ws.max_row + 1):
    ws.row_dimensions[r].height = 34
for r in range(2, ws.max_row + 1):
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
style_sheet(ws, width={'项目': 16, '内容': 120}, freeze='A2')
ws.auto_filter.ref = None

# 结论摘要
ws = wb['结论摘要']
for r in range(2, ws.max_row + 1):
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[r].height = 34
style_sheet(ws, width={'项目': 22, '内容': 120}, freeze='A2')
ws.auto_filter.ref = None

# 起点汇总
style_sheet(
    wb['起点汇总'],
    pct=('第一名250日回撤(%)', '第二名250日回撤(%)', '首笔组合收益(%)',
         '首笔浮动收益(%)', '首笔平均单笔(%)', '总收益(%)', '全期年化(%)',
         '最大回撤(%)', '胜率(%)', '平均单笔(%)', '中位单笔(%)',
         '最好单笔(%)', '最差单笔(%)', '持仓时间占比(%)',
         '1年收益(%)', '1年最大回撤(%)', '3年收益(%)', '3年最大回撤(%)',
         '5年收益(%)', '5年最大回撤(%)'),
    mult=('期末倍数',), yuan=('总利润(元)',),
    date=('起点信号日', '计划入场日', '首笔入场日', '首笔最后退出日'),
    num=('当日信号数', '首笔买入只数', '笔数', '止盈笔数', '到期笔数', '期末未平笔数'),
    width={'起点信号日': 12, '第一名代码': 11, '第二名代码': 11,
           '计划入场日': 12, '首笔入场日': 12, '首笔最后退出日': 12,
           '首笔代码': 22, '首笔单笔收益': 22, '样本': 9,
           '总利润(元)': 13, '期末倍数': 10},
    freeze='C2',
)
# 样本配色
ws = wb['起点汇总']
c_complete = col_index(ws, '首笔是否完整')
for r in range(2, ws.max_row + 1):
    if ws.cell(r, c_complete).value == '否':
        for c in range(1, c_complete + 1):
            ws.cell(r, c).fill = PatternFill('solid', fgColor='FDEBD0')

# 起点分布
style_sheet(
    wb['起点分布'],
    pct=('均值', '第5百分位', '第25百分位', '中位数', '第75百分位',
         '第95百分位', '最小值', '最大值', '负收益占比(%)',
         '≥30%占比(%)', '≥40%占比(%)'),
    num=('样本数',),
    width={'样本': 20, '指标': 18}, freeze='C2',
)
for r in range(2, wb['起点分布'].max_row + 1):
    if '最大回撤' in str(wb['起点分布'].cell(r, 2).value):
        for c in (12, 13, 14):
            wb['起点分布'].cell(r, c).value = None

# 起点年度
style_sheet(
    wb['起点年度'],
    pct=('全期年化中位数(%)', '全期年化均值(%)', '全期年化负收益占比(%)',
         '全期总收益中位数(%)', '1年收益中位数(%)', '3年收益中位数(%)',
         '5年收益中位数(%)', '首笔批次收益中位数(%)'),
    num=('起点年份', '起点数', '满1年数', '满3年数', '满5年数'),
    width={'起点年份': 10, '备注': 30}, freeze='B2', wrap=('备注',),
)

# 起点最好最差
style_sheet(
    wb['起点最好最差'],
    pct=('首笔组合收益(%)', '全期年化(%)', '最大回撤(%)',
         '1年收益(%)', '3年收益(%)', '5年收益(%)'),
    mult=('期末倍数',), date=('起点信号日', '计划入场日'),
    num=('当日信号数',),
    width={'分组': 18, '提示': 22, '起点信号日': 12, '第一名代码': 11,
           '第二名代码': 11, '计划入场日': 12},
    freeze='C2',
)

# 相邻信号日
style_sheet(
    wb['相邻信号日'],
    pct=('前二重合度(%)', 'A首笔收益(%)', 'B首笔收益(%)', 'A全期年化(%)',
         'B全期年化(%)', 'A首年收益(%)', 'B首年收益(%)',
         'A三年收益(%)', 'B三年收益(%)', 'A五年收益(%)', 'B五年收益(%)',
         '交易重合度(%)'),
    pp=('首笔差(百分点)', '年化差(百分点)', '首年差(百分点)',
        '三年差(百分点)', '五年差(百分点)'),
    num=('间隔交易日', '间隔自然日', 'A当日信号数', 'B当日信号数',
         'A首笔买入只数', 'B首笔买入只数', '前二重合只数', '交易重合数'),
    date=('信号日A', '信号日B'),
    width={'信号日A': 11, '信号日B': 11, '配对类型': 14,
           'A前二': 22, 'B前二': 22},
    freeze='C2',
)
ws = wb['相邻信号日']
c_same = col_index(ws, '前二完全相同')
c_both = col_index(ws, '两日信号都≥2只')
for r in range(2, ws.max_row + 1):
    if ws.cell(r, c_same).value == '是':
        for c in range(1, c_same + 1):
            ws.cell(r, c).fill = PatternFill('solid', fgColor='D6EAF8')
    if ws.cell(r, c_both).value == '否':
        for c in range(1, c_both + 1):
            ws.cell(r, c).fill = PatternFill('solid', fgColor='F2F3F4')

# 相邻汇总
style_sheet(
    wb['相邻信号日汇总'],
    pct=('前二零重合占比(%)', '前二完全相同占比(%)', '交易重合度中位(%)'),
    pp=('首笔差中位(百分点)', '首年差中位(百分点)', '首年差5%分位(百分点)',
        '首年差95%分位(百分点)', '首年差绝对值中位(百分点)',
        '三年差中位(百分点)', '五年差中位(百分点)', '五年差绝对值中位(百分点)'),
    num=('对数', '首笔可算对数', '首年可算对数', '三年可算对数', '五年可算对数'),
    width={'样本': 24}, freeze='B2',
)

# 逐日首笔
style_sheet(
    wb['逐日首笔'],
    pct=('第一名250日回撤(%)', '第二名250日回撤(%)', '首笔组合收益(%)',
         '首笔浮动收益(%)', '首笔平均单笔(%)'),
    date=('信号日', '计划入场日', '首笔入场日', '首笔最后退出日'),
    num=('当日信号数', '首笔买入只数', '首笔完整只数'),
    width={'信号日': 12, '第一名代码': 11, '第二名代码': 11,
           '计划入场日': 12, '首笔入场日': 12, '首笔最后退出日': 12,
           '首笔代码': 24, '首笔单笔收益': 24, '状态': 12},
    freeze='C2',
)
ws = wb['逐日首笔']
c_complete = col_index(ws, '首笔是否完整')
for r in range(2, ws.max_row + 1):
    if ws.cell(r, c_complete).value == '否':
        for c in range(1, c_complete + 1):
            ws.cell(r, c).fill = PatternFill('solid', fgColor='FDEBD0')

# 逐只首笔
style_sheet(
    wb['逐只首笔'],
    pct=('250日回撤(%)', '单笔收益(%)', '入场前涨幅(%)'),
    num=('当日信号数', '排名'),
    date=('信号日', '入场日', '退出日'),
    width={'信号日': 12, '代码': 11, '状态': 10, '退出原因': 12},
    freeze='C2',
)
ws = wb['逐只首笔']
c_complete = col_index(ws, '是否完整')
for r in range(2, ws.max_row + 1):
    if ws.cell(r, c_complete).value == '否':
        for c in range(1, c_complete + 1):
            ws.cell(r, c).fill = PatternFill('solid', fgColor='FDEBD0')

# 全部起点逐笔
style_sheet(
    wb['全部起点逐笔'],
    pct=('单笔收益(%)', '分配额(初始本金%)', '仓位占当时权益(%)',
         '盈亏贡献(初始本金%)', '250日回撤(%)', '市场20日(%)'),
    dec2=('入场价(后复权)', '退出价(后复权)'),
    num=('持有交易日', '信号排名'),
    date=('起点信号日', '信号日', '入场日', '退出日'),
    width={'起点信号日': 12, '信号日': 12, '代码': 11, '入场日': 12,
           '退出日': 12, '退出原因': 12, '入场价(后复权)': 13,
           '退出价(后复权)': 13},
    freeze='D2',
)
for r in range(2, wb['全部起点逐笔'].max_row + 1):
    if wb['全部起点逐笔'].cell(r, 6).value == '期末未平':
        for c in range(1, 7):
            wb['全部起点逐笔'].cell(r, c).fill = PatternFill('solid', fgColor='FDEBD0')

# 每信号日索引
style_sheet(
    wb['每信号日索引'],
    mult=('期末倍数',),
    num=('当日信号数', '笔数'),
    date=('起点信号日', '首笔入场日', '末笔退出日', '计划入场日'),
    width={'起点信号日': 12, '首笔入场日': 12, '末笔退出日': 12,
           '计划入场日': 12, '文件名': 16, '样本': 9},
    freeze='B2',
)
for r in range(2, wb['每信号日索引'].max_row + 1):
    for c in (8, 9, 10, 11, 12):
        wb['每信号日索引'].cell(r, c).number_format = '0.0'

# 最近案例
style_sheet(
    wb['最近信号日'],
    pct=('第一名250日回撤(%)', '第二名250日回撤(%)', '首笔浮动收益(%)',
         '全期年化(%)', '1年收益(%)', '3年收益(%)', '5年收益(%)'),
    date=('起点信号日', '计划入场日', '首笔入场日'),
    num=('当日信号数', '首笔买入只数'),
    width={'起点信号日': 12, '第一名代码': 11, '第二名代码': 11,
           '计划入场日': 12, '首笔入场日': 12, '首笔状态': 10, '说明': 55},
    wrap=('说明',), freeze='C2',
)

# 口径
ws = wb['口径与参数']
for r in range(2, ws.max_row + 1):
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[r].height = 30
style_sheet(ws, width={'参数/项目': 20, '值': 120}, freeze='A2')
ws.auto_filter.ref = None

wb.save(XLSX)
print('excel done', XLSX, len(wb.sheetnames), 'sheets')
