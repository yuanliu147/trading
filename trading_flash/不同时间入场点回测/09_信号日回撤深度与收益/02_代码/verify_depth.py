# -*- coding: utf-8 -*-
"""信号日回撤深度与收益：独立复核。

原则：核心数字从 06/08 的原始逐笔表重算，再与 09 的输出表 / Excel / 图表对照；
不拿 09 的输出反推 09 的结论。
"""
from pathlib import Path
import json

import numpy as np
import pandas as pd
import openpyxl
from scipy import stats

HERE = Path(__file__).resolve()
OUT = HERE.parents[1]
BASE = HERE.parents[2]
RES = OUT / '03_结果'
FIG = OUT / '04_图表'
EV = OUT / '05_证据'
EV.mkdir(parents=True, exist_ok=True)

lines = []
fails = []


def chk(name, cond, detail=''):
    ok = bool(cond)
    lines.append(('PASS' if ok else 'FAIL') + ' | ' + name + (' | ' + detail if detail else ''))
    if not ok:
        fails.append(name)


def near(name, got, want, tol=1e-6):
    chk(name, abs(float(got) - float(want)) <= tol, f'got={got:.6f} want={want:.6f} tol={tol}')


# ---------- 原始数据 ----------
T0 = pd.read_csv(BASE / '06_不同信号日起点/03_结果/逐日首笔_逐只.csv')
T = T0[T0['complete'] == True].copy()  # noqa: E712
T['y'] = pd.to_datetime(T['signal_date']).dt.year
T['m'] = pd.to_datetime(T['signal_date']).dt.strftime('%Y-%m')
T['days'] = (pd.to_datetime(T['exit_date']) - pd.to_datetime(T['entry_date'])).dt.days

A0 = pd.read_csv(BASE / '08_重复信号与买入时点/03_结果/全部信号重复序列_完整.csv')
A = A0[A0['complete'] == True].copy()  # noqa: E712

chk('前二原始 459 行', len(T0) == 459, f'{len(T0)}')
chk('前二完整 450 行', len(T) == 450, f'{len(T)}')
Araw = pd.read_csv(BASE / '08_重复信号与买入时点/03_结果/全部信号池_原始.csv')
chk('全部信号原始 4195 行', len(Araw) == 4195, f'{len(Araw)}')
chk('全部信号完整 4019 行', len(A) == 4019, f'{len(A)}')
chk('前二完整信号日 246 天', T['signal_date'].nunique() == 246, f"{T['signal_date'].nunique()}")
chk('全部信号信号日 246 天', A['signal_date'].nunique() == 246, f"{A['signal_date'].nunique()}")

# ---------- 核心分档 ----------
deep = T[T['dd250'] <= -0.65]
sh = T[T['dd250'] > -0.65]
chk('前二深组 109 笔', len(deep) == 109, f'{len(deep)}')
chk('前二浅组 341 笔', len(sh) == 341, f'{len(sh)}')
near('深组平均 9.21%', deep['ret'].mean() * 100, 9.2058, .01)
near('浅组平均 3.99%', sh['ret'].mean() * 100, 3.9950, .01)
near('深浅差 5.21pp', (deep['ret'].mean() - sh['ret'].mean()) * 100, 5.2109, .01)
near('深组胜率 77.06%', (deep['ret'] > 0).mean() * 100, 77.0642, .01)
near('浅组胜率 61.88%', (sh['ret'] > 0).mean() * 100, 61.8768, .01)
near('深组止盈率 43.12%', (deep['exit_reason'] == 'target').mean() * 100, 43.1193, .01)
near('浅组止盈率 23.46%', (sh['exit_reason'] == 'target').mean() * 100, 23.4604, .01)
near('深组平均持有 16.25 日', deep['days'].mean(), 16.2477, .01)
near('浅组平均持有 19.20 日', sh['days'].mean(), 19.2023, .01)
near('深组最差 -22.46%', deep['ret'].min() * 100, -22.4561, .01)
near('浅组最差 -33.03%', sh['ret'].min() * 100, -33.0342, .01)
p_mw = stats.mannwhitneyu(deep['ret'], sh['ret'], alternative='two-sided').pvalue
chk('秩和 p = 5.58e-05', abs(p_mw - 5.5804e-05) < 1e-8, f'{p_mw:.3e}')

deep60 = T[T['dd250'] <= -0.60]
sh60 = T[T['dd250'] > -0.60]
chk('前二 -60 深组 217 笔', len(deep60) == 217, f'{len(deep60)}')
near('-60 深浅差 2.31pp', (deep60['ret'].mean() - sh60['ret'].mean()) * 100, 2.3115, .01)
p60 = stats.mannwhitneyu(deep60['ret'], sh60['ret'], alternative='two-sided').pvalue
near('-60 秩和 p 0.0478', p60, 0.047822, 1e-5)

deep_a = A[A['dd250'] <= -0.65]
sh_a = A[A['dd250'] > -0.65]
chk('全部信号深组 181 笔', len(deep_a) == 181, f'{len(deep_a)}')
near('全部信号深组 11.08%', deep_a['ret'].mean() * 100, 11.0803, .01)
near('全部信号浅组 5.20%', sh_a['ret'].mean() * 100, 5.1963, .01)

# 分档计数与均值
b = pd.read_csv(RES / '前二深度分档.csv')
chk('前二分档 7 档', len(b) == 7, f'{len(b)}')
chk('前二分档合计 450', int(b['笔数'].sum()) == 450, f"{b['笔数'].sum()}")
near('≤-75 档 19.70%', b.loc[b['深度档'] == '≤-75%', '平均收益'].iloc[0], 19.70, .01)
near('-75~-70 档 9.68%', b.loc[b['深度档'] == '-75~-70%', '平均收益'].iloc[0], 9.68, .01)
near('-55~-50 档 1.65%', b.loc[b['深度档'] == '-55~-50%', '平均收益'].iloc[0], 1.65, .01)
ba = pd.read_csv(RES / '全部信号深度分档.csv')
chk('全部信号分档合计 4019', int(ba['笔数'].sum()) == 4019, f"{ba['笔数'].sum()}")

# ---------- 同日排名 ----------
ret = T.pivot_table(index='signal_date', columns='rank', values='ret', aggfunc='first').dropna()
dd = T.pivot_table(index='signal_date', columns='rank', values='dd250', aggfunc='first').dropna()
chk('同日可配对 204 天', len(ret) == 204, f'{len(ret)}')
near('排名1 平均 6.29%', ret[1].mean() * 100, 6.2906, .01)
near('排名2 平均 4.25%', ret[2].mean() * 100, 4.2462, .01)
near('排名1-2 差 2.04pp', (ret[1].mean() - ret[2].mean()) * 100, 2.0444, .01)
near('配对 t p 0.0262', stats.ttest_rel(ret[1], ret[2]).pvalue, 0.026228, 1e-5)
chk('排名1 回撤更深的天数占多数', (dd[1] <= dd[2]).mean() > .99, f"{(dd[1] <= dd[2]).mean():.3f}")
pr = pd.read_csv(RES / '同日排名对比.csv')
near('输出表 同日差 2.04', pr.loc[pr['样本'] == '全部', '平均差(百分点)'].iloc[0], 2.04, .01)
pyr = pd.read_csv(RES / '同日排名_年度.csv')
chk('同日年度 10 年', len(pyr) == 10, f'{len(pyr)}')
chk('同日年度 10 年差值全为正', (pyr['平均差(百分点)'] > 0).all(), f"{pyr['平均差(百分点)'].tolist()}")

# ---------- 回归 ----------
T['depth_pp'] = (-T['dd250'] - 0.45) * 100
A['depth_pp'] = (-A['dd250'] - 0.45) * 100
meta = A.drop_duplicates('signal_date').set_index('signal_date')
T['mkt_dd250'] = T['signal_date'].map(meta['mkt_dd250'])


def ols(D, xcols):
    D = D.dropna(subset=['ret', 'y'] + xcols).copy()
    X = np.column_stack([np.ones(len(D))] + [D[c].astype(float).values for c in xcols]
                        + [pd.get_dummies(D['y'], drop_first=True, dtype=float).values])
    y = D['ret'].astype(float).values
    bb = np.linalg.lstsq(X, y, rcond=None)[0]
    r = y - X @ bb
    s2 = float(r @ r / (len(D) - X.shape[1]))
    se = np.sqrt(np.diag(s2 * np.linalg.pinv(X.T @ X)))
    return bb, se


bb, se = ols(T, ['depth_pp'])
near('前二斜率 3.95pp/10pp', bb[1] * 1000, 3.9467, .01)
near('前二斜率 t=5.04', bb[1] / se[1], 5.0354, .02)
bb2, se2 = ols(T, ['depth_pp', 'mkt_dd250'])
near('前二控市场后 3.20pp', bb2[1] * 1000, 3.2029, .01)
near('前二控市场 t=3.97', bb2[1] / se2[1], 3.9657, .02)
ba2, sea2 = ols(A, ['depth_pp'])
near('全部信号斜率 2.79pp', ba2[1] * 1000, 2.7871, .01)
ba3, sea3 = ols(A, ['depth_pp', 'mkt_dd250'])
near('全部信号控市场后 1.40pp', ba3[1] * 1000, 1.4029, .01)
reg = pd.read_csv(RES / '深度回归.csv')
near('输出 前二斜率 3.95', reg.loc[(reg['样本'] == '前二') & (reg['模型'] == '年度固定效应'),
                                  '每深10个百分点收益(百分点)'].iloc[0], 3.95, .01)

# ---------- 日级 ----------
day = T.groupby('signal_date').agg(dd=('dd250', 'mean'), ret=('ret', 'mean'))
day_ex = T[T['y'] != 2026].groupby('signal_date').agg(dd=('dd250', 'mean'), ret=('ret', 'mean'))
near('日级 Pearson -0.160', stats.pearsonr(day['dd'], day['ret']).statistic, -0.15993, .001)
near('日级 p 0.012', stats.pearsonr(day['dd'], day['ret']).pvalue, 0.012011, 1e-5)
near('日级 ex2026 -0.192', stats.pearsonr(day_ex['dd'], day_ex['ret']).statistic, -0.19154, .001)
chk('日级 246 天', len(day) == 246, f'{len(day)}')
crowd = T.groupby('signal_date').agg(dd=('dd250', 'mean'), ret=('ret', 'mean'), pool=('n_signals', 'first'))
crowd['grp'] = pd.qcut(crowd['pool'], 3, labels=['小池子', '中池子', '大池子'])
crowd_out = pd.read_csv(RES / '信号拥挤度分组.csv')
chk('拥挤度 3 组', len(crowd_out) == 3, f'{len(crowd_out)}')
for lab, want_r, want_p in [('小池子', -0.062, .5601), ('中池子', -0.112, .341), ('大池子', -0.228, .042)]:
    gg = crowd[crowd['grp'] == lab]
    rr = stats.spearmanr(gg['dd'], gg['ret'])
    near(f'{lab} 相关 {want_r}', rr.statistic, want_r, .001)
    near(f'{lab} p {want_p}', rr.pvalue, want_p, .001)

# ---------- 年度 / 阶段 ----------
yt = pd.read_csv(RES / '年度深度对比.csv')
top = yt[yt['样本'] == '前二']
chk('年度表 20 行', len(yt) == 20, f'{len(yt)}')
both = top.dropna(subset=['深组平均', '浅组平均'])
chk('前二 8 个可比年份', len(both) == 8, f'{len(both)}')
chk('前二 8 个可比年份深组全胜', (both['深组平均'] > both['浅组平均']).all(),
    f"{both[['年份', '平均差(百分点)']].to_dict('records')}")
near('2026 深组 8.92', top.loc[top['年份'] == 2026, '深组平均'].iloc[0], 8.92, .01)
near('2026 浅组 -3.37', top.loc[top['年份'] == 2026, '浅组平均'].iloc[0], -3.37, .01)
ph = pd.read_csv(RES / '阶段深度对比.csv')
chk('阶段表 8 行', len(ph) == 8, f'{len(ph)}')
near('D 阶段深组 9.81', ph.loc[(ph['样本'] == '前二') & (ph['阶段'].str.startswith('D')),
                              '深组平均'].iloc[0], 9.81, .01)

# ---------- 2026 ----------
june = T[(T['y'] == 2026) & (T['m'] == '2026-06')]
july = T[(T['y'] == 2026) & (T['m'] == '2026-07')]
chk('2026-06 深组 9 笔', len(june[june['dd250'] <= -0.65]) == 9)
near('2026-06 深组 0.29', june[june['dd250'] <= -0.65]['ret'].mean() * 100, 0.2923, .01)
near('2026-06 浅组 -9.30', june[june['dd250'] > -0.65]['ret'].mean() * 100, -9.3013, .01)
chk('2026-07 深组 24 笔', len(july[july['dd250'] <= -0.65]) == 24)
near('2026-07 深组 12.15', july[july['dd250'] <= -0.65]['ret'].mean() * 100, 12.1548, .01)
near('2026-07 浅组 6.11', july[july['dd250'] > -0.65]['ret'].mean() * 100, 6.1081, .01)
m26 = pd.read_csv(RES / '2026深度分档.csv')
chk('2026 月度 4 行', len(m26) == 4, f'{len(m26)}')

# ---------- 机制 / 极值 ----------
mech = pd.read_csv(RES / '深度机制.csv')
near('深组止盈占比 43.12', mech.loc[mech['深度组'] == '≤-65%', '止盈占比'].iloc[0], 43.12, .01)
near('浅组止盈占比 23.46', mech.loc[mech['深度组'] == '>-65%', '止盈占比'].iloc[0], 23.46, .01)
ext = T[T['dd250'] <= -0.70]
chk('≤-70 极值 37 笔', len(ext) == 37, f'{len(ext)}')
near('≤-70 平均 11.85', ext['ret'].mean() * 100, 11.8496, .01)
near('≤-70 胜率 89.19', (ext['ret'] > 0).mean() * 100, 89.1892, .01)
ext_out = pd.read_csv(RES / '深度极值_70以上.csv')
chk('极值输出 37 行', len(ext_out) == 37, f'{len(ext_out)}')

# ---------- 重复信号 ----------
P = pd.read_csv(BASE / '08_重复信号与买入时点/03_结果/首笔vs第二笔_明细.csv')
P['dd_chg'] = (P['second_dd'] - P['first_dd']) * 100
P['ret_chg'] = (P['second'] - P['first']) * 100
P['deeper'] = P['dd_chg'] < 0
chk('重复配对 78 对', len(P) == 78, f'{len(P)}')
chk('第二笔更深 39 对', int(P['deeper'].sum()) == 39, f"{P['deeper'].sum()}")
near('第二笔更深 收益差 -0.28', P[P['deeper']]['ret_chg'].mean(), -0.2837, .01)
near('第二笔更浅 收益差 -0.09', P[~P['deeper']]['ret_chg'].mean(), -0.0935, .01)
c = stats.pearsonr(P['dd_chg'], P['ret_chg'])
near('深度变化相关 0.336', c.statistic, 0.33633, .001)
near('相关 p 0.0026', c.pvalue, 0.002607, 1e-5)

C = pd.read_csv(BASE / '08_重复信号与买入时点/03_结果/603778案例.csv')
Cok = C[(C['状态'] == 'ok') & (C['是否前二'] == True)]  # noqa: E712
chk('603778 前二完整 7 笔', len(Cok) == 7, f'{len(Cok)}')
near('603778 深度-收益相关 -0.8525', stats.pearsonr(Cok['回撤250'], Cok['收益']).statistic, -0.8525, .001)

# ---------- 当前信号 ----------
CUR = pd.read_csv(RES / '当前信号_深度对照.csv', parse_dates=['信号日'])
chk('当前信号 6 行', len(CUR) == 6, f'{len(CUR)}')
v = lambda code, col: CUR.loc[CUR['代码'] == code, col].iloc[0]
near('605117 信号日回撤 -45.32', v('605117.SH', '信号日回撤_后复权'), -45.32, .01)
near('605117 实际价回撤 -61.07', v('605117.SH', '实际价回撤'), -61.07, .01)
near('603778 信号日回撤 -78.61', v('603778.SH', '信号日回撤_后复权'), -78.61, .01)
near('605376 信号日回撤 -71.30', v('605376.SH', '信号日回撤_后复权'), -71.30, .01)
near('002969 信号日回撤 -70.68', v('002969.SZ', '信号日回撤_后复权'), -70.68, .01)
near('603778 同档历史 19.70', v('603778.SH', '历史同档平均'), 19.70, .01)
near('605376 同档历史 9.68', v('605376.SH', '历史同档平均'), 9.68, .01)
near('605117 同档历史 4.65', v('605117.SH', '历史同档平均'), 4.65, .01)

# ---------- 输出完整性 ----------
trade = pd.read_csv(RES / '前二逐笔_深度.csv')
chk('前二逐笔 450 行', len(trade) == 450, f'{len(trade)}')
chk('前二逐笔列无英文 hfq', not any('hfq' in c for c in trade.columns))
chk('当前表列无 hfq', not any('hfq' in c for c in CUR.columns))
mj = json.loads((RES / '指标.json').read_text(encoding='utf-8'))
near('指标 json 深组平均', mj['深度分档']['前二≤-65平均'] * 100, 9.2058, .01)
near('指标 json 同日差', mj['同日排名']['平均差(百分点)'], 2.0444, .01)

wb = openpyxl.load_workbook(OUT / '信号日回撤深度与收益.xlsx')
chk('Excel 23 个 sheet', len(wb.sheetnames) == 23, f'{len(wb.sheetnames)}')
chk('Excel 结论摘要 14 行', wb['结论摘要'].max_row == 14, f"{wb['结论摘要'].max_row}")
chk('Excel 当前信号 7 行', wb['当前信号'].max_row == 7, f"{wb['当前信号'].max_row}")
chk('Excel 前二逐笔 451 行', wb['前二逐笔'].max_row == 451, f"{wb['前二逐笔'].max_row}")
chk('Excel 深度分档 15 行', wb['深度分档'].max_row == 15, f"{wb['深度分档'].max_row}")
chk('Excel 当前回撤单元格 -45.32', abs(wb['当前信号'].cell(2, 6).value + 45.32) < .01,
    f"{wb['当前信号'].cell(2, 6).value}")

charts = sorted(FIG.glob('*.png'))
chk('图表 5 张', len(charts) == 5, f'{len(charts)}')
chk('图表均大于 10KB', all(p.stat().st_size > 10000 for p in charts),
    f"{[(p.name, p.stat().st_size) for p in charts]}")

# ---------- 汇总 ----------
n_pass = sum(1 for x in lines if x.startswith('PASS'))
n_fail = len(fails)
lines.append('')
lines.append(f'合计：{n_pass} PASS / {n_fail} FAIL')
(EV / '校验.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8')
print('\n'.join(lines))
raise SystemExit(1 if n_fail else 0)
