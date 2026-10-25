# -*- coding: utf-8 -*-
"""市场阶段适配：交付校验。"""
from pathlib import Path
import json

import numpy as np
import pandas as pd
import openpyxl

HERE = Path(__file__).resolve()
OUT = HERE.parents[1]
RES = OUT / '03_结果'
FIG = OUT / '04_图表'
EV = OUT / '05_证据'
EV.mkdir(parents=True, exist_ok=True)

rows = []


def check(name, cond, got=''):
    rows.append((bool(cond), name, str(got)))


A = pd.read_csv(RES / '年度对比.csv')
P = pd.read_csv(RES / '阶段统计.csv')
V = pd.read_csv(RES / '波动分组.csv')
F = pd.read_csv(RES / '起点固定持有期.csv', parse_dates=['signal_date'])
C = pd.read_csv(RES / '2026起点.csv', parse_dates=['signal_date'])
D = pd.read_csv(RES / '全部信号日逐只_带阶段.csv', parse_dates=['signal_date'])
S26 = pd.read_csv(RES / '2026信号逐只.csv', parse_dates=['signal_date'])
DD = pd.read_csv(RES / '回撤事实.csv')
M = pd.read_csv(RES / '指标.json') if False else json.loads((RES / '指标.json').read_text(encoding='utf-8'))
CO = pd.read_csv(OUT.parent / '06_不同信号日起点' / '03_结果' / '起点汇总.csv', parse_dates=['signal_date'])

check('信号日总数=251', len(CO) == 251, len(CO))
check('逐只总数=459', len(D) == 459, len(D))
check('完整逐只=450', int(D['complete'].sum()) == 450, int(D['complete'].sum()))
check('未完成=9', int((~D['complete']).sum()) == 9, int((~D['complete']).sum()))
check('2026 信号日=49', len(CO[CO.signal_date.dt.year == 2026]) == 49, len(CO[CO.signal_date.dt.year == 2026]))
check('2026 逐只=92', len(S26) == 92, len(S26))
check('2026 完整=84', int(S26['complete'].sum()) == 84, int(S26['complete'].sum()))
check('2026 未完成=8', int((~S26['complete']).sum()) == 8, int((~S26['complete']).sum()))

jun = S26[(S26.signal_date.dt.strftime('%Y-%m') == '2026-06') & S26.complete]
jul = S26[(S26.signal_date.dt.strftime('%Y-%m') == '2026-07') & S26.complete]
check('6月38笔', len(jun) == 38, len(jun))
check('6月均值-7.03%', abs(jun.ret.mean() + 0.07029148) < 1e-6, round(jun.ret.mean(), 8))
check('6月胜率18.42%', abs((jun.ret > 0).mean() - 0.184210526) < 1e-6, round((jun.ret > 0).mean(), 8))
check('7月32笔', len(jul) == 32, len(jul))
check('7月均值+10.64%', abs(jul.ret.mean() - 0.10643125) < 1e-6, round(jul.ret.mean(), 8))
check('7月胜率81.25%', abs((jul.ret > 0).mean() - 0.8125) < 1e-6, round((jul.ret > 0).mean(), 8))

check('2026起点49', len(C) == 49, len(C))
check('2026起点13负', int((C.total_ret < 0).sum()) == 13, int((C.total_ret < 0).sum()))
check('2026起点34正', int((C.total_ret > 0).sum()) == 34, int((C.total_ret > 0).sum()))
check('2026起点2个未入场', int((C.total_ret == 0).sum()) == 2, int((C.total_ret == 0).sum()))
check('2026起点中位+8.93%', abs(C.total_ret.median() - 0.0892576) < 1e-6, round(C.total_ret.median(), 8))
worst = C.loc[C.total_ret.idxmin()]
best = C.loc[C.total_ret.idxmax()]
check('最差起点6/4 -28.07%', str(worst.signal_date.date()) == '2026-06-04' and abs(worst.total_ret + 0.2807146) < 1e-6,
      f'{worst.signal_date.date()} {worst.total_ret:.6f}')
check('最好起点7/9 +48.92%', str(best.signal_date.date()) == '2026-07-09' and abs(best.total_ret - 0.4891522) < 1e-6,
      f'{best.signal_date.date()} {best.total_ret:.6f}')

hist60 = F.loc[F.signal_date.dt.year < 2026, 'ret_60'].dropna()
cur60 = F.loc[F.signal_date.dt.year == 2026, 'ret_60'].dropna()
check('历史60日202个', len(hist60) == 202, len(hist60))
check('2026 60日40个', len(cur60) == 40, len(cur60))
check('历史60日中位+10.85%', abs(hist60.median() - 0.1085) < 1e-4, round(hist60.median(), 6))
check('2026 60日中位+5.29%', abs(cur60.median() - 0.0529) < 1e-4, round(cur60.median(), 6))
check('历史60日亏损15.8%', abs((hist60 < 0).mean() - 0.1584) < 1e-3, round((hist60 < 0).mean(), 6))
check('2026 60日亏损40.0%', abs((cur60 < 0).mean() - 0.40) < 1e-9, round((cur60 < 0).mean(), 6))
check('2026 60日最差-27.34%', abs(cur60.min() + 0.2734125) < 1e-6, round(cur60.min(), 6))
dd60 = F.loc[F.signal_date.dt.year == 2026, 'dd_60'].min()
check('2026 60日最差回撤-38.29%', abs(dd60 + 0.3829465) < 1e-6, round(dd60, 6))
check('历史60日最差回撤-30.05%', abs(F.loc[F.signal_date.dt.year < 2026, 'dd_60'].min() + 0.3005) < 1e-3,
      round(F.loc[F.signal_date.dt.year < 2026, 'dd_60'].min(), 6))

aa = A.set_index('年份')
check('2018策略+47.41%', abs(aa.loc[2018, '策略收益'] - 0.474056) < 1e-5, aa.loc[2018, '策略收益'])
check('2018市场-30.42%', abs(aa.loc[2018, '市场等权'] + 0.304208) < 1e-5, aa.loc[2018, '市场等权'])
check('2021策略+1.34%', abs(aa.loc[2021, '策略收益'] - 0.013413) < 1e-5, aa.loc[2021, '策略收益'])
check('2021市场+27.71%', abs(aa.loc[2021, '市场等权'] - 0.277085) < 1e-5, aa.loc[2021, '市场等权'])
check('2022策略+49.58%', abs(aa.loc[2022, '策略收益'] - 0.495808) < 1e-5, aa.loc[2022, '策略收益'])
check('2026策略+19.75%', abs(aa.loc[2026, '策略收益'] - 0.197469) < 1e-5, aa.loc[2026, '策略收益'])
check('2026市场-8.47%', abs(aa.loc[2026, '市场等权'] + 0.084733) < 1e-5, aa.loc[2026, '市场等权'])

check('阶段A=132', int(P.set_index('分类').loc['A 牛市上升段急跌', '笔数']) == 132,
      int(P.set_index('分类').loc['A 牛市上升段急跌', '笔数']))
check('阶段B=94', int(P.set_index('分类').loc['B 高位转弱/跌破年线', '笔数']) == 94,
      int(P.set_index('分类').loc['B 高位转弱/跌破年线', '笔数']))
check('阶段C=22', int(P.set_index('分类').loc['C 熊市反弹', '笔数']) == 22,
      int(P.set_index('分类').loc['C 熊市反弹', '笔数']))
check('阶段D=202', int(P.set_index('分类').loc['D 熊市下跌', '笔数']) == 202,
      int(P.set_index('分类').loc['D 熊市下跌', '笔数']))
check('四阶段合计450', int(P.set_index('分类').loc[['A 牛市上升段急跌', 'B 高位转弱/跌破年线', 'C 熊市反弹', 'D 熊市下跌'], '笔数'].sum()) == 450)
check('回归beta=1.2796', abs(M['回归']['beta'] - 1.2795747) < 1e-6, M['回归']['beta'])
check('回归corr=0.5235', abs(M['回归']['corr'] - 0.5235447) < 1e-6, M['回归']['corr'])
check('回归R2=0.2741', abs(M['回归']['r2'] - 0.2740990) < 1e-6, M['回归']['r2'])

check('2026账户+19.75%', abs(M['二零二六']['账户收益'] - 0.1974694) < 1e-6, M['二零二六']['账户收益'])
check('2026最大回撤-16.00%', abs(M['二零二六']['账户最大回撤'] + 0.1600028) < 1e-6, M['二零二六']['账户最大回撤'])
check('2026回撤峰4/16', M['二零二六']['回撤峰'] == '2026-04-16', M['二零二六']['回撤峰'])
check('2026回撤谷6/30', M['二零二六']['回撤谷'] == '2026-06-30', M['二零二六']['回撤谷'])
check('2026修复8/6', M['二零二六']['修复'] == '2026-08-06', M['二零二六']['修复'])
check('全期最大回撤-16.00%', abs(M['二零二六']['全期最大回撤'] + 0.1600028) < 1e-6, M['二零二六']['全期最大回撤'])

check('2026 6月笔数38', M['二零二六']['六月笔数'] == 38, M['二零二六']['六月笔数'])
# 直接用表算 6 月环境
jun_all = S26[S26.signal_date.dt.strftime('%Y-%m') == '2026-06']
check('2026 6月波动中位0.1785', abs(jun_all.market_volatility20.median() - 0.1785) < 1e-4,
      round(jun_all.market_volatility20.median(), 6))
check('2026 6月信号数中位34', abs(jun_all.n_signals.median() - 34) < 1e-9, jun_all.n_signals.median())
check('2026 6月250日涨幅中位19.66%', abs(jun_all.mkt_ret250.median() - 0.1966) < 1e-3,
      round(jun_all.mkt_ret250.median(), 6))

wb = openpyxl.load_workbook(OUT / '市场阶段适配.xlsx', read_only=True)
check('Excel 17个sheet', len(wb.sheetnames) == 17, len(wb.sheetnames))
check('Excel含2026起点', '2026起点' in wb.sheetnames)
check('Excel含结论摘要', '结论摘要' in wb.sheetnames)
check('图表7张', len(list(FIG.glob('*.png'))) == 7, len(list(FIG.glob('*.png'))))

# 输出
ok = sum(1 for r in rows if r[0])
lines = [f'市场阶段适配 · 交付校验  {ok}/{len(rows)} PASS', '']
for passed, name, got in rows:
    lines.append(f"{'PASS' if passed else 'FAIL'}  {name}" + (f"  [{got}]" if got else ''))
lines.append('')
lines.append('数据源：同花顺本地 DuckDB 派生 research_panel.parquet；沪深主板等权 market_daily.parquet。')
lines.append('口径：T+3 开盘、2 仓 50%、+20% 止盈、15 日到期、0.175%/边。')
text = '\n'.join(lines)
(EV / '校验.txt').write_text(text, encoding='utf-8')
print(text)
