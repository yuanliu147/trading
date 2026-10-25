# -*- coding: utf-8 -*-
"""把市场阶段适配分析写成中文 Excel。"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve()
OUT = HERE.parents[1]
RES = OUT / '03_结果'

A = pd.read_csv(RES / '年度对比.csv')
P = pd.read_csv(RES / '阶段统计.csv')
V = pd.read_csv(RES / '波动分组.csv')
V2 = pd.read_csv(RES / '年线波动分组.csv')
Q = pd.read_csv(RES / '持有期市场分位.csv')
C = pd.read_csv(RES / '2026起点.csv', parse_dates=['signal_date'])
F = pd.read_csv(RES / '固定持有期对比.csv')
L = pd.read_csv(RES / '长期对比.csv')
DD = pd.read_csv(RES / '回撤事实.csv')
W = pd.read_csv(RES / '最差交易.csv', parse_dates=['signal_date'])
B = pd.read_csv(RES / '最好交易.csv', parse_dates=['signal_date'])
S26 = pd.read_csv(RES / '2026信号逐只.csv', parse_dates=['signal_date'])
S = pd.read_csv(RES / '全部信号日逐只_带阶段.csv', parse_dates=['signal_date'])
M = pd.read_csv(RES / '市场阶段日度.csv', parse_dates=['date'])

# 中文列名
S26 = S26.drop(columns=['complete'], errors='ignore')
S26 = S26.rename(columns={
    'signal_date': '信号日', 'n_signals': '当日信号数', 'rank': '排名', 'code': '代码',
    'dd250': '个股250日回撤', 'rsi14': 'RSI14', 'close_loc': '收盘位置',
    'status': '状态', 'complete': '是否完成', 'entry_date': '入场日',
    'exit_date': '退出日', 'exit_reason': '退出原因', 'ret': '收益',
    '阶段': '市场阶段', 'mkt_ret250': '市场250日涨幅', 'mkt_dd250': '市场250日回撤',
    'market_volatility20': '市场20日波动率', 'breadth20': '市场宽度20',
    '信号后15日市场': '信号后15日市场', '持有期市场': '持有期市场',
    '是否完整': '是否完整'})
S = S.drop(columns=['complete'], errors='ignore')
S = S.rename(columns={
    'signal_date': '信号日', 'n_signals': '当日信号数', 'rank': '排名', 'code': '代码',
    'dd250': '个股250日回撤', 'rsi14': 'RSI14', 'close_loc': '收盘位置',
    'status': '状态', 'complete': '是否完成', 'entry_date': '入场日',
    'exit_date': '退出日', 'exit_reason': '退出原因', 'ret': '收益',
    '阶段': '市场阶段', 'mkt_ret250': '市场250日涨幅', 'mkt_dd250': '市场250日回撤',
    'market_volatility20': '市场20日波动率', 'breadth20': '市场宽度20',
    '信号后15日市场': '信号后15日市场', '持有期市场': '持有期市场',
    '是否完整': '是否完整'})

C = C.rename(columns={
    'signal_date': '信号日', 'n_signals': '当日信号数', 'n_buy': '买入只数',
    'batch_ret': '首笔批次收益', 'final': '期末倍数', 'total_ret': '期末总收益',
    'n': '交易笔数', 'win': '胜率', 'avg': '单笔均值', 'ret_60': '60日收益',
    'dd_60': '60日最大回撤', 'ret_120': '120日收益', '是否亏损': '是否亏损'})
C = C[['信号日', '当日信号数', '买入只数', '首笔批次收益', '期末总收益', '期末倍数',
       '交易笔数', '胜率', '单笔均值', '60日收益', '60日最大回撤', '120日收益', '是否亏损']]

W = W.rename(columns={'signal_date': '信号日', 'code': '代码', 'entry_date': '入场日',
                      'exit_date': '退出日', 'ret': '收益', 'exit_reason': '退出原因',
                      '阶段': '市场阶段', 'mkt_ret250': '市场250日涨幅',
                      'mkt_dd250': '市场250日回撤', '持有期市场': '持有期市场'})
B = B.rename(columns={'signal_date': '信号日', 'code': '代码', 'entry_date': '入场日',
                      'exit_date': '退出日', 'ret': '收益', 'exit_reason': '退出原因',
                      '阶段': '市场阶段', 'mkt_ret250': '市场250日涨幅',
                      'mkt_dd250': '市场250日回撤', '持有期市场': '持有期市场'})

M = M.rename(columns={'date': '日期', 'equal_weight_index': '等权指数', 'ma250': '250日均线',
                      'mkt_ret250': '市场250日涨幅', 'mkt_ret120': '市场120日涨幅',
                      'mkt_dd250': '市场250日回撤', 'mkt_above250': '年线上方',
                      'market_return20': '市场20日收益', 'market_return60': '市场60日收益',
                      'breadth20': '市场宽度20', 'market_volatility20': '市场20日波动率',
                      '阶段': '市场阶段', '是否信号日': '是否信号日'})

# 结论摘要
summary = pd.DataFrame([
    ('一句话', '不是“牛市策略”，也不是“熊市策略”；它赚的是急跌后 15 个交易日内有没有反弹的钱。'),
    ('年度证据', '2018 等权 -30.4%，策略 +47.4%；2022 等权 -3.5%，策略 +49.6%；2021 等权 +27.7%，策略仅 +1.3%。'),
    ('长期证据', '1 年起点 73.3% 跑赢等权，3 年 95.2%，5 年 100%；5 年策略中位 +386.9%，市场 +71.9%。'),
    ('单笔驱动', '持有期市场上涨 291 笔均值 +9.6%、胜率 82%；下跌 159 笔均值 -2.6%、胜率 36%。Beta=1.28。'),
    ('最好的环境', '高波动、深度回撤、恐慌充分释放后的快速反弹；不一定非要是牛市。'),
    ('最差的环境', '阴跌/滚动下跌，20 日跌 5% 反复触发但 15 天内不反弹；2026 年 6 月是典型。'),
    ('2026 账户', '至 10-23 +19.75%，不是亏损；但 4/16→6/30 回撤 -16.0%，是 9 年最大回撤，8/6 修复。'),
    ('2026 信号', '3-7 月 84 笔完整：均值 +1.5%、胜率 48.8%；6 月 38 笔 -7.0%/18.4%，7 月 32 笔 +10.6%/81.3%。'),
    ('2026 起点', '49 个起点：34 正、13 负、2 个未入场，中位 +8.9%；最差 6/4 -28.1%，最好 7/9 +48.9%。'),
    ('2026 同口径', '60 日：2026 中位 +5.3%、亏损 40%；历史 +10.9%、亏损 15.8%。确实低于常态。'),
    ('为什么 6/4 亏', '首批 605255/002097 亏 15%；晚一天换仓，买到 002342(-33.0%)/603878(-13.9%)；2 仓 15 天路径依赖。'),
    ('风险提示', '母路径 99 笔最差单笔 -17.2%，但全部 450 笔最差 -33.0%；2026-06-04 起点 60 日回撤 -38.3%。'),
    ('不要做的事', '不要因为 2026 就加“只在牛市做”的过滤：年线上方 +4.0%，年线下方 +5.9%。'),
    ('如果要做开关', '波动率比牛熊方向更有信息量（高波动 +7.8% vs 低波动 +2.7%），但剔除 2026 后差距收窄，且会砍掉一半信号；先另开研究。'),
], columns=['项目', '内容'])

notes = pd.DataFrame([
    ('数据源', '同花顺本地 DuckDB；沪深主板等权指数来自 trading_hithink/data/work/market_daily.parquet。'),
    ('样本', '251 个信号日；459 条信号日逐只；450 条完整；9 条未完成（1 条一字涨停、8 条持有未满/数据结束）。'),
    ('2026 样本', '49 个信号日；92 条逐只；84 条完整（3-7 月）；10 月 8 条未完成，不计入完整统计。'),
    ('口径', '信号日收盘确认，T+3 开盘买；2 仓、每仓 50% 当前权益；+20% 盘中止盈；15 个交易日到期卖；成本 0.175%/边；一字涨停/停牌不买、不递补。'),
    ('收益口径', '逐笔为后复权、含成本；年度账户为连续账户（10 万本金、2 仓 50%）后复权。'),
    ('2026 年度', '2025-12-31 至 2026-10-23；市场等权 -8.5%，策略 +19.8%，超额 +28.2pp。'),
    ('2017 年度', '策略从 2017-11-23 开始，2017 为部分年度。'),
    ('阶段定义', 'A 年线上方且市场 250 日涨幅>0；B 年线下方且 250 日涨幅>0；C 年线上方且 250 日涨幅<=0；D 年线下方且 250 日涨幅<=0。'),
    ('波动定义', '信号日市场 20 日波动率，以 450 条完整样本的中位数分高/低。'),
    ('重要警告', '信号日高度重叠（2026 年 49 天只有 2 个长簇），不能把 450 笔当成 450 个独立样本；阶段差异是描述性的，不是可实盘的择时规则。'),
    ('幸存者偏差', '股票池为当前上市的主板股票，未含退市股；2018 等熊市年份的收益可能被高估。'),
    ('历史不代表未来', '先纸面 50 笔；任何环境开关都要独立样本外验证。'),
], columns=['项目', '说明'])

params = pd.DataFrame([
    ('信号条件', '250 日回撤<-45%；250 日涨幅>0；RSI14<35；收盘位置>0.45；市场 20 日<-5%'),
    ('排序', '按个股 250 日回撤最深取前二'),
    ('仓位', '2 仓，每仓 50% 当前权益；无信号空仓'),
    ('入场', '信号后第 3 个交易日开盘'),
    ('退出', '+20% 盘中止盈；15 个交易日收盘到期；无普通止损'),
    ('成本', '0.175%/边（佣金、印花税、滑点）'),
    ('不可买', '一字涨停、停牌不买、不递补'),
    ('数据截止', '2026-10-23'),
], columns=['参数', '设定'])

path = OUT / '市场阶段适配.xlsx'
with pd.ExcelWriter(path, engine='openpyxl') as writer:
    notes.to_excel(writer, sheet_name='说明', index=False)
    summary.to_excel(writer, sheet_name='结论摘要', index=False)
    A.to_excel(writer, sheet_name='年度对比', index=False)
    P.to_excel(writer, sheet_name='阶段统计', index=False)
    V.to_excel(writer, sheet_name='波动分组', index=False)
    V2.to_excel(writer, sheet_name='年线波动', index=False)
    Q.to_excel(writer, sheet_name='持有期分位', index=False)
    L.to_excel(writer, sheet_name='长期对比', index=False)
    C.to_excel(writer, sheet_name='2026起点', index=False)
    F.to_excel(writer, sheet_name='固定持有期', index=False)
    DD.to_excel(writer, sheet_name='回撤事实', index=False)
    W.to_excel(writer, sheet_name='最差交易', index=False)
    B.to_excel(writer, sheet_name='最好交易', index=False)
    S26.to_excel(writer, sheet_name='2026信号逐只', index=False)
    S.to_excel(writer, sheet_name='全部信号日', index=False)
    M.to_excel(writer, sheet_name='市场日度', index=False)
    params.to_excel(writer, sheet_name='口径与参数', index=False)

    wb = writer.book
    head_fill = PatternFill('solid', fgColor='1F4E79')
    head_font = Font(color='FFFFFF', bold=True, size=10)
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    pct_tokens = ('收益', '胜率', '回撤', '涨幅', '波动率', '占比', '亏损',
                  '平均', '中位', '最差', '最好', 'p5', 'p95', 'p10', 'p90',
                  '首笔', '期末总', '60日', '120日', '持有期市场', '市场等权',
                  '均值', '宽度')
    count_tokens = ('笔数', '起点数', '交易笔数', '买入只数', '当日信号数',
                    '信号日数', '完整单笔', '年份', '排名', '是否完整', '是否完成',
                    '年线上方', '是否信号日')

    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            header = str(ws.cell(1, col).value)
            width = max(10, min(34, max(len(str(ws.cell(r, col).value or '')) for r in range(1, min(ws.max_row, 300) + 1)) + 2))
            ws.column_dimensions[letter].width = width
            fmt = None
            if header == '超额百分点':
                fmt = '0.00"pp"'
            elif any(t in header for t in pct_tokens):
                fmt = '0.00%'
            elif any(t in header for t in count_tokens):
                fmt = '0'
            elif header in ('期末倍数', '等权指数', '250日均线'):
                fmt = '0.000'
            elif header in ('信号日', '入场日', '退出日', '日期'):
                fmt = 'yyyy-mm-dd'
            if fmt:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, col).number_format = fmt
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical='center')

    # 说明/结论摘要 左对齐、宽列
    for name in ('说明', '结论摘要', '口径与参数'):
        ws = wb[name]
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 95
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)

    # 回撤事实：数值列是收益率
    ws = wb['回撤事实']
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).number_format = '0.00%'
    # 2026 起点：冻结首列
    ws = wb['2026起点']
    ws.freeze_panes = 'B2'

print('written', path)
