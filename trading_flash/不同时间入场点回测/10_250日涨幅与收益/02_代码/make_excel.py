# -*- coding: utf-8 -*-
"""250 日涨幅专题 Excel。"""
from pathlib import Path
import json

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parents[2]
RES = BASE / "10_250日涨幅与收益" / "03_结果"
OUT = BASE / "10_250日涨幅与收益" / "250日涨幅与收益.xlsx"
M = json.loads((RES / "指标.json").read_text(encoding="utf-8"))

DARK = "17231F"
RED = "B54430"
CREAM = "F7F1E6"
LINE = "E1D7C7"
WHITE = "FFFFFF"

thin = Side(style="thin", color=LINE)
border = Border(bottom=thin)


def pct(x, d=1):
    return f"{x*100:+.{d}f}%"


def pct0(x, d=1):
    return f"{x*100:.{d}f}%"


def add_table(wb, title, df, note="", formats=None, widths=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial Unicode MS", size=15, bold=True, color=DARK)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(df.columns)))
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(name="Arial Unicode MS", size=9, color="6F766F")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(1, len(df.columns)))
    start = 3
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(start, j, col)
        c.fill = PatternFill("solid", fgColor=DARK)
        c.font = Font(name="Arial Unicode MS", size=10, bold=True, color=WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    for i, row in df.iterrows():
        for j, col in enumerate(df.columns, 1):
            v = row[col]
            if pd.isna(v):
                v = None
            c = ws.cell(start + 1 + i, j, v)
            c.font = Font(name="Arial Unicode MS", size=9, color="1C2521")
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=(col in ("模型", "样本", "备注", "说明")))
            if formats and col in formats:
                c.number_format = formats[col]
    ws.freeze_panes = ws.cell(start + 1, 1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(df.columns))}{start + len(df)}"
    for j, col in enumerate(df.columns, 1):
        if widths and col in widths:
            w = widths[col]
        else:
            w = min(26, max(9, max(len(str(col)), *(len(str(x)) for x in df[col].head(40))) + 2))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 26
    if note:
        ws.row_dimensions[2].height = 38
    return ws


wb = Workbook()
wb.remove(wb.active)

# 说明
ws = wb.create_sheet("说明")
ws.sheet_view.showGridLines = False
rows = [
    ("250 日涨幅与单笔收益：历史记录里“最近一年涨幅”有用吗？", 16, True, DARK),
    ("", 9, False, "1C2521"),
    ("问题：信号日当天，个股从 250 个交易日前到信号日的涨幅（return_250d），与这笔交易的盈利有没有关系？", 10, False, "1C2521"),
    ("", 9, False, "1C2521"),
    ("口径：", 10, True, RED),
    ("r250 = 信号日后复权收盘 / 250 个交易日前复权收盘 - 1；250 个交易日约 12 个月，不是 365 个自然日。", 10, False, "1C2521"),
    ("策略本身要求 r250 > 0；本文只在“已经通过 >0 门槛”的样本里，看涨幅大小与收益的关系。", 10, False, "1C2521"),
    ("样本：前二 450 笔（2017-11-20 至 2026-07-31）、母路径 99 笔（含 10/23 两笔持仓按市值计）、全部信号 4,019 笔。", 10, False, "1C2521"),
    ("收益：后复权、含成本、T+3 开盘买、+20% 止盈、15 个交易日到期；前二 450 笔与母路径 99 笔不是同一样本。", 10, False, "1C2521"),
    ("", 9, False, "1C2521"),
    ("数据源：本地同花顺 DuckDB 派生 research_panel.parquet；没有使用 baostock/Tushare 等第三方数据。", 10, False, "6F766F"),
    ("这不是投资建议；历史关系是描述性的，不是因果，也不能直接当实盘过滤。", 10, False, "6F766F"),
]
for i, (text, size, bold, color) in enumerate(rows, 1):
    c = ws.cell(i, 1, text)
    c.font = Font(name="Arial Unicode MS", size=size, bold=bold, color=color)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[i].height = 32 if size == 10 else 24
ws.column_dimensions["A"].width = 120

# 结论摘要
summary = [
    ("一句话结论", "有“止盈概率”上的弱关系，但没有稳定的“平均收益”关系；不能加“涨得越多越买”的过滤。"),
    ("前二 450 笔", f"Pearson {M['相关性']['前二全部']['pearson']:.3f}（p={M['相关性']['前二全部']['p']:.3f}）；Spearman {M['相关性']['前二全部']['spearman']:.3f}（p={M['相关性']['前二全部']['spearman_p']:.3f}）。"),
    ("六档均值", "0–10% +3.4% / 10–20% +4.8% / 20–50% +6.1% / 50–100% +5.4% / 100–200% +6.6% / >200% +6.8%；不单调，>200% 只有 11 笔。"),
    ("止盈率", "0–10% 18.7% → 100–200% 45.9% → >200% 63.6%；每 +100pp r250，止盈 OR 1.84（95%CI 1.31–2.98），控制深度/波动后 1.69（1.19–3.14）。"),
    ("平均收益", f"每 +100pp r250，单笔均值 +{M['回归']['前二_全部_均值_年度FE']['估计']:.2f}pp（95%CI {M['回归']['前二_全部_均值_年度FE']['CI低']:.2f}~{M['回归']['前二_全部_均值_年度FE']['CI高']:.2f}）；控制后 +1.09pp（-1.72~2.86），区间跨 0。"),
    ("剔除 2026", f"均值相关 {M['相关性']['前二_剔除2026']['pearson']:.3f}（p={M['相关性']['前二_剔除2026']['p']:.3f}），几乎为零；止盈 OR 仍为 {M['回归']['前二_剔除2026_止盈_年度FE']['估计']:.2f}（1.09–3.22）。"),
    ("2026", f"均值相关 {M['相关性']['前二_2026']['pearson']:.3f}（p={M['相关性']['前二_2026']['p']:.3f}），但 6 月 +0.124（p=0.459）、7 月 -0.061（p=0.740）；是一年样本，不稳定。"),
    ("为什么均值不涨", "高 r250 的股票波动更大：0–10% 档个股波动均值 0.47，>200% 档 0.75；止盈率上升，但 p10 从 -11.2% 走到 -18.8%，左尾更深，均值互相抵消。"),
    ("更强的变量", "个股波动率五分位：Q1→Q5 单笔均值 +0.9%→+8.3%，止盈率 14.4%→44.4%；比 r250 的区分度更强（但也是样本内观察）。"),
    ("当前", "603778 +110.6%（100–200% 档）、605376 +38.2%（20–50% 档）、002969 +176.3%（100–200% 档）、605117 +31.6%（20–50% 档）、002796 +9.7%（0–10% 档）。前二都不在 0–10% 弱档，但不构成额外保证。"),
]
ws = wb.create_sheet("结论摘要")
ws.sheet_view.showGridLines = False
ws["A1"] = "结论摘要"
ws["A1"].font = Font(name="Arial Unicode MS", size=15, bold=True, color=DARK)
for i, (k, v) in enumerate(summary, 3):
    ws.cell(i, 1, k).font = Font(name="Arial Unicode MS", size=10, bold=True, color=RED)
    ws.cell(i, 2, v).font = Font(name="Arial Unicode MS", size=10, color="1C2521")
    ws.cell(i, 1).alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(i, 2).alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[i].height = 34
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 105

# 数据表
sample_df = pd.read_csv(RES / "样本说明.csv")
add_table(wb, "样本与定义", sample_df,
          "r250 是信号日前的 250 个交易日涨幅；前二 450 笔与母路径 99 笔、全部信号 4,019 笔不是同一样本。",
          formats={"250日涨幅均值": "0.0%", "250日涨幅中位": "0.0%", "单笔均值": "0.0%", "单笔中位": "0.0%", "胜率": "0.0%", "止盈率": "0.0%"})

bin_top2 = pd.read_csv(RES / "六档成绩_前二.csv")
pct_cols = ["单笔均值", "单笔中位", "胜率", "止盈率", "p10", "p90", "最差", "最好",
            "平均回撤", "平均个股波动", "剔除2026均值", "剔除2026胜率", "剔除2026止盈率",
            "2026均值", "2026胜率", "2026止盈率"]
add_table(wb, "六档成绩_前二", bin_top2,
          "前二 450 笔。涨幅档按 r250 切；均值为单笔收益。>200% 只有 11 笔，不能当稳定规律。",
          formats={c: "0.0%" for c in pct_cols})

q_df = pd.read_csv(RES / "五分位_前二.csv")
add_table(wb, "五分位_前二", q_df,
          "按 r250 五等分，每档 90 笔。Q1 最弱，Q2–Q5 均值差距不大；止盈率从 Q1 到 Q5 上升。",
          formats={c: "0.0%" for c in ["最小", "最大"] + pct_cols if c in q_df.columns})

corr = pd.read_csv(RES / "相关性.csv")
add_table(wb, "相关性", corr,
          "Pearson 看线性关系；Spearman 看秩相关。p 值是单变量检验，未做多重比较校正。",
          formats={"Pearson相关": "0.000", "Spearman相关": "0.000", "Pearson p": "0.000", "Spearman p": "0.000"})

reg = pd.read_csv(RES / "回归.csv")
add_table(wb, "回归", reg,
          "均值模型为 OLS；止盈模型为 Logistic，估计为每 +100pp r250 的 OR。区间为按交易 bootstrap 的 95%CI；年度固定效应见模型名。",
          formats={"估计": "0.000", "CI低": "0.000", "CI高": "0.000", "拟合": "0.000"})

yr = pd.read_csv(RES / "年度相关性_前二.csv")
add_table(wb, "年度相关性_前二", yr,
          "年度 Pearson 正负翻转；2023 只有 9 笔，2021 只有 20 笔，单年相关不能外推。",
          formats={"250日涨幅中位": "0.0%", "单笔均值": "0.0%", "单笔中位": "0.0%", "胜率": "0.0%", "止盈率": "0.0%", "Pearson相关": "0.000", "p值": "0.000"})

phase = pd.read_csv(RES / "阶段拆分_前二.csv")
add_table(wb, "阶段拆分_前二", phase,
          "按 r250 中位数（27.4%）分高低。C 只有 22 笔，阶段内差异只能当描述。",
          formats={c: "0.0%" for c in pct_cols if c in phase.columns})

depth = pd.read_csv(RES / "深度拆分_前二.csv")
add_table(wb, "深度拆分_前二", depth,
          "深度≤-65% 与 >-65%，各再按 r250 高低分。深度仍是主变量；高 r250 只在深回撤组里略有增益。",
          formats={c: "0.0%" for c in pct_cols if c in depth.columns})

vol = pd.read_csv(RES / "波动五分位_前二.csv")
add_table(wb, "波动机制_前二", vol,
          "个股 volatility20 五分位。波动率与 r250 正相关（Pearson 0.219），但波动率对收益/止盈的区分度更强。",
          formats={c: "0.0%" for c in ["平均个股波动", "250日涨幅均值"] + pct_cols if c in vol.columns})

bin_all = pd.read_csv(RES / "六档成绩_全部信号.csv")
add_table(wb, "全部信号_六档", bin_all,
          "全部信号 4,019 笔（不只前二）。均值仍不单调；止盈率随 r250 上升。",
          formats={c: "0.0%" for c in pct_cols if c in bin_all.columns})

bin_mother = pd.read_csv(RES / "六档成绩_母路径.csv")
add_table(wb, "母路径_六档", bin_mother,
          "母路径 99 笔（含 10/23 两笔持仓按市值计）。每档样本小，>200% 只有 2 笔；不要据此加过滤。",
          formats={c: "0.0%" for c in pct_cols if c in bin_mother.columns})

cur = pd.read_csv(RES / "当前信号_250日涨幅.csv")
add_table(wb, "当前信号", cur,
          "当前 10/22–10/23 信号。250日涨幅是后复权；同档统计来自前二 450 笔。",
          formats={"250日回撤(%)": '0.00"%"', "250日涨幅": "0.0%", "250日涨幅(%)": '0.00"%"', "同档均值": "0.0%", "同档均值(%)": '0.00"%"',
                   "同档胜率": "0.0%", "同档胜率(%)": '0.00"%"', "同档止盈率": "0.0%", "同档止盈率(%)": '0.00"%"'})

tr = pd.read_csv(RES / "逐笔_前二_带250日涨幅.csv")
add_table(wb, "逐笔明细", tr,
          "前二 450 笔逐笔；信号日、涨幅、收益、退出原因都在。可用于抽查。",
          formats={"250日涨幅": "0.0%", "单笔收益": "0.0%", "250日回撤": "0.0%", "个股波动率": "0.0%", "持有期市场": "0.0%"})

# 口径
ws = wb.create_sheet("口径")
ws.sheet_view.showGridLines = False
lines = [
    ("口径与复现", 15, True, DARK),
    ("", 9, False, "1C2521"),
    ("r250：后复权收盘 / 250 个交易日前复权收盘 - 1；research_panel.return_250d 与逐笔重算最大差 0.0。", 10, False, "1C2521"),
    ("样本：前二 450 笔完整；母路径 99 笔（含 2 笔 10/23 持仓按市值计）；全部信号 4,019 笔完整。", 10, False, "1C2521"),
    ("收益：T+3 开盘买，+20% 盘中止盈，15 个交易日到期，无普通止损，成本 0.175%/边。", 10, False, "1C2521"),
    ("市场：沪深主板等权指数；本文不使用上证指数，也不把“全A”当基准。", 10, False, "1C2521"),
    ("统计：Pearson/Spearman 为单变量；回归为 OLS/Logistic，年度固定效应，bootstrap 95%CI。", 10, False, "1C2521"),
    ("边界：单年/小样本/重复信号；2026 只有 84 笔前二、1645 笔全部信号；分档是描述性的。", 10, False, "1C2521"),
    ("不要直接加“r250 高就买/低就排除”的过滤；如要试，先做独立样本外验证。", 10, True, RED),
    ("数据源：本地同花顺 DuckDB 派生 research_panel.parquet；无第三方数据。", 10, False, "6F766F"),
]
for i, (text, size, bold, color) in enumerate(lines, 1):
    c = ws.cell(i, 1, text)
    c.font = Font(name="Arial Unicode MS", size=size, bold=bold, color=color)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[i].height = 34
ws.column_dimensions["A"].width = 120

wb.save(OUT)
print("written", OUT)
print("sheets", len(wb.sheetnames), wb.sheetnames)
