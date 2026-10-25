# -*- coding: utf-8 -*-
"""校验 250 日涨幅专题。"""
from pathlib import Path
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[2]
ROOT = BASE / "10_250日涨幅与收益"
RES = ROOT / "03_结果"
FIG = ROOT / "04_图表"
OUT = ROOT / "05_证据" / "校验.txt"
NOTES = []


def ok(name, value=None, want=None, tol=None):
    if tol is None:
        passed = True
    else:
        passed = abs(float(value) - float(want)) <= tol
    NOTES.append(f"{'PASS' if passed else 'FAIL'}  {name}" + (f"  got={value} want={want} tol={tol}" if tol is not None else ""))
    return passed


def pct(x):
    return f"{x*100:+.2f}%"


# 文件
for rel in ["00_一页纸_250日涨幅与收益.md", "01_完整报告_250日涨幅与收益.md", "03_数据与复现.md",
            "250日涨幅与收益.xlsx", "README.md"]:
    ok(f"存在 {rel}", (ROOT / rel).exists())
for rel in ["analyze_250d.py", "make_charts.py", "make_excel.py", "verify_250d.py"]:
    ok(f"存在脚本 {rel}", (ROOT / "02_代码" / rel).exists())

# 指标
M = json.loads((RES / "指标.json").read_text(encoding="utf-8"))
ok("前二 450 笔", M["样本"]["前二"], 450, 0)
ok("母路径 99 笔", M["样本"]["母路径"], 99, 0)
ok("全部信号 4019 笔", M["样本"]["全部信号"], 4019, 0)
ok("r250 定义最大差 0.0", M["样本"]["定义校验最大差"], 0.0, 1e-12)

# 相关
for name, want_p, want_s in [
    ("前二全部", 0.05816819996188067, 0.08983691883163403),
    ("前二_剔除2026", -0.002, 0.048),
    ("前二_2026", 0.227, 0.221),
    ("母路径99笔", 0.154, 0.082),
    ("全部信号", 0.011, 0.020),
]:
    got = M["相关性"][name]
    ok(f"{name} Pearson", got["pearson"], want_p, 0.002)
    ok(f"{name} Spearman", got["spearman"], want_s, 0.003)

# 六档
b = pd.read_csv(RES / "六档成绩_前二.csv")
ok("六档 6 行", len(b), 6, 0)
ok("六档笔数合计 450", int(b["笔数"].sum()), 450, 0)
for label, n in [("0–10%", 91), ("10–20%", 75), ("20–50%", 151),
                 ("50–100%", 85), ("100–200%", 37), (">200%", 11)]:
    row = b[b["涨幅档"] == label].iloc[0]
    ok(f"{label} 笔数", int(row["笔数"]), n, 0)
r0 = b[b["涨幅档"] == "0–10%"].iloc[0]
r1 = b[b["涨幅档"] == "100–200%"].iloc[0]
r2 = b[b["涨幅档"] == ">200%"].iloc[0]
ok("0–10% 均值 +3.37%", float(r0["单笔均值"]), 0.0336807, 0.0002)
ok("0–10% 止盈率 18.68%", float(r0["止盈率"]), 0.186813, 0.001)
ok("100–200% 均值 +6.56%", float(r1["单笔均值"]), 0.0656484, 0.0002)
ok("100–200% 止盈率 45.95%", float(r1["止盈率"]), 0.459459, 0.001)
ok(">200% 均值 +6.83%", float(r2["单笔均值"]), 0.068343, 0.001)
ok(">200% 止盈率 63.64%", float(r2["止盈率"]), 0.636364, 0.001)

# 回归
reg = pd.read_csv(RES / "回归.csv")
ok("回归 12 行", len(reg), 12, 0)
rm = reg[reg["模型"] == "前二_全部_均值_年度FE"].iloc[0]
ok("前二均值斜率 +1.69pp", float(rm["估计"]), 1.6869, 0.01)
rt = reg[reg["模型"] == "前二_全部_止盈_年度FE"].iloc[0]
ok("前二止盈 OR 1.84", float(rt["估计"]), 1.8367, 0.01)
re = reg[reg["模型"] == "前二_剔除2026_均值_年度FE"].iloc[0]
ok("剔除2026均值斜率 -0.47pp", float(re["估计"]), -0.4690, 0.02)

# 当前
c = pd.read_csv(RES / "当前信号_250日涨幅.csv")
for code, want in [("603778.SH", 110.63), ("605376.SH", 38.21), ("002969.SZ", 176.31),
                   ("605117.SH", 31.61), ("002796.SZ", 9.66)]:
    got = float(c[c["代码"] == code]["250日涨幅(%)"].iloc[0])
    ok(f"当前 {code} r250 {want}%", got, want, 0.02)
ok("当前 6 行", len(c), 6, 0)

# Excel
wb = load_workbook(ROOT / "250日涨幅与收益.xlsx", read_only=True)
ok("Excel 16 个 sheet", len(wb.sheetnames), 16, 0)
ok("Excel 含结论摘要", "结论摘要" in wb.sheetnames)
ok("Excel 含逐笔明细", "逐笔明细" in wb.sheetnames)
ok("Excel 含口径", "口径" in wb.sheetnames)

# 图表
charts = sorted(FIG.glob("*.png"))
ok("图表 6 张", len(charts), 6, 0)
ok("图表均 > 10KB", all(p.stat().st_size > 10240 for p in charts))

# 文案
t = (ROOT / "00_一页纸_250日涨幅与收益.md").read_text(encoding="utf-8")
ok("一页纸含止盈概率结论", "止盈概率" in t and "平均收益" in t)
ok("一页纸含不能加过滤", "不要加" in t and "r250" in t)

# 写文件
OUT.write_text("\n".join(NOTES) + f"\n\n合计：{sum(x.startswith('PASS') for x in NOTES)} PASS / {sum(x.startswith('FAIL') for x in NOTES)} FAIL\n", encoding="utf-8")
print("\n".join(NOTES[-12:]))
print(f"\n合计：{sum(x.startswith('PASS') for x in NOTES)} PASS / {sum(x.startswith('FAIL') for x in NOTES)} FAIL")
