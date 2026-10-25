# -*- coding: utf-8 -*-
"""交付校验：文件、数字、逐笔、Excel、图表、基准复现。"""
import hashlib
import json
import os
import sys

import numpy as np
import pandas as pd

OUT = '05_证据'
os.makedirs(OUT, exist_ok=True)
results = []


def check(name, cond, detail=''):
    results.append((name, bool(cond), detail))


def close(a, b, tol=1e-6):
    return abs(float(a) - float(b)) <= tol


# ---------------- files
base_files = [
    'README.md', '00_一页纸_不同入场时点.md', '01_完整报告_不同入场时点.md',
    '03_数据与复现.md', '不同入场时点回测.xlsx',
    '03_结果/组合汇总_全部时点.csv', '03_结果/年度收益_全部时点.csv',
    '03_结果/逐笔_组合_全部时点.csv', '03_结果/权益曲线_全部时点.csv',
    '03_结果/基准99笔_不同时点收益明细.csv', '03_结果/信号级汇总_全部信号.csv',
    '03_结果/信号级明细_每日前二.csv', '03_结果/信号集合变化_按档位.csv',
    '03_结果/执行差异_对比基准.csv', '03_结果/跳过统计_全部时点.csv',
    '03_结果/当前信号_2026-10-23.csv', '03_结果/当前信号_入场日期.csv',
    '03_结果/指标.json',
]
for f in base_files:
    check(f'文件存在：{f}', os.path.exists(f))

charts = [f'04_图表/0{i}_' for i in range(1, 8)]
for f in charts:
    hits = [x for x in os.listdir('04_图表') if x.startswith(f.split('/')[-1])]
    check(f'图表存在：{f}*', len(hits) == 1, ','.join(hits))

# ---------------- summary
S = pd.read_csv('03_结果/组合汇总_全部时点.csv')
check('组合汇总 51 行', len(S) == 51, str(len(S)))
check('组合汇总 编号唯一', S.case_id.is_unique)
base = S[S.case_id == 'O03'].iloc[0]
check('O03 年化 40.0%', close(base.cagr, 0.4001, 1e-3), f'{base.cagr:.6f}')
check('O03 回撤 -16.0%', close(base.dd, -0.16, 1e-4), f'{base.dd:.6f}')
check('O03 99 笔', int(base.n) == 99, str(base.n))
check('O03 期末 20.75x', close(base.final, 20.7458, 1e-3), f'{base.final:.4f}')
check('O03 胜率 71.7%', close(base.win, 0.7172, 1e-3), f'{base.win:.4f}')
check('O03 利润因子 4.06', close(base.pf, 4.056, 1e-2), f'{base.pf:.4f}')
check('主档位数 30', (S.类型 == '主').sum() == 30, str((S.类型 == '主').sum()))
check('主档位 ≥30% 共 4 个', (S[S.类型 == '主'].cagr >= 0.30).sum() == 4,
      str((S[S.类型 == '主'].cagr >= 0.30).sum()))
check('主档位年化最小值 8.0%', close(S[S.类型 == '主'].cagr.min(), 0.0802, 1e-3))
check('主档位年化最大值 40.0%', close(S[S.类型 == '主'].cagr.max(), 0.4001, 1e-3))
for c in ['cagr', 'dd', 'n', 'win', 'avg', 'pf', 'final']:
    check(f'组合汇总无缺失 {c}', S[c].notna().all())
check('组合汇总 final>0', (S.final > 0).all())

# ---------------- trades
T = pd.read_csv('03_结果/逐笔_组合_全部时点.csv', parse_dates=['信号日', '入场日', '退出日'])
check('逐笔 4,865 行', len(T) == 4865, str(len(T)))
check('逐笔 51 档', T.case_id.nunique() == 51, str(T.case_id.nunique()))
check('逐笔无重复键', not T.duplicated(['case_id', '信号日', '代码']).any())
check('逐笔入场价/退出价无空值', T['入场价_后复权'].notna().all() and T['退出价_后复权'].notna().all())
check('逐笔价格>0', (T['入场价_后复权'] > 0).all() and (T['退出价_后复权'] > 0).all())
check('退出日>=入场日', (T.退出日 >= T.入场日).all())
check('信号日<=入场日', (T.信号日 <= T.入场日).all())
check('持有 1–15 日', T['持有交易日'].between(1, 15).all(),
      f"{T['持有交易日'].min()}–{T['持有交易日'].max()}")
c00 = T[T.case_id == 'C00']
check('C00 全部同日收盘', (c00.信号日 == c00.入场日).all())
check('逐笔收益范围合理', T['单笔收益'].between(-100, 100).all(),
      f"{T['单笔收益'].min():.2f}–{T['单笔收益'].max():.2f}")
# 每档笔数与汇总一致
counts = T.groupby('case_id').size().to_dict()
bad = [(cid, counts.get(cid, 0), int(S[S.case_id == cid].n.iloc[0]))
       for cid in S.case_id if counts.get(cid, 0) != int(S[S.case_id == cid].n.iloc[0])]
check('每档笔数与汇总一致', len(bad) == 0, str(bad[:3]))

# ---------------- annual
A = pd.read_csv('03_结果/年度收益_全部时点.csv')
check('年度表 51 行', len(A) == 51, str(len(A)))
year_cols = [str(y) for y in range(2017, 2027)]
check('年度表 2017–2026 列', all(y in A.columns for y in year_cols))
a03 = A[A.case_id == 'O03'].iloc[0]
check('O03 2021 +1.3%', close(a03['2021'], 0.01341, 1e-3), f"{a03['2021']:.4f}")
check('O03 2024 +93.0%', close(a03['2024'], 0.92984, 1e-3), f"{a03['2024']:.4f}")

# ---------------- base 99 counterfactual
W = pd.read_csv('03_结果/基准99笔_不同时点收益明细.csv')
check('基准99笔 99 行', len(W) == 99, str(len(W)))
check('基准99笔 O03 列', 'O03' in W.columns)
check('基准99笔 O03=基准收益', np.allclose(W['O03'], W['基准组合收益'], atol=1e-6))

# ---------------- signal level
G = pd.read_csv('03_结果/信号级汇总_全部信号.csv')
check('信号级 31 档', len(G) == 31, str(len(G)))
g03 = G[G.case_id == 'O03'].iloc[0]
check('信号级 4195 条', int(g03['信号总数']) == 4195, str(g03['信号总数']))
check('信号级 O03 基准99平均 6.68%', close(g03['基准99_平均'], 0.066793, 1e-4),
      f"{g03['基准99_平均']:.6f}")
check('信号级 O03 前二平均 5.20%', close(g03['前二_平均'], 0.051972, 1e-4),
      f"{g03['前二_平均']:.6f}")
D = pd.read_csv('03_结果/信号级明细_每日前二.csv')
check('信号级明细 > 10,000 行', len(D) > 10000, str(len(D)))
check('信号级明细状态列', '状态' in D.columns)

# ---------------- overlap
O = pd.read_csv('03_结果/信号集合变化_按档位.csv')
check('信号集合变化 15 行', len(O) == 15, str(len(O)))
o3 = O[O['第N个交易日'] == 3].iloc[0]
check('N=3 Jaccard 6.9%', close(o3['平均Jaccard'], 0.068762, 1e-4), f"{o3['平均Jaccard']:.6f}")
check('N=3 前二存活率 13.5%', close(o3['前二存活率'], 0.134538, 1e-4), f"{o3['前二存活率']:.6f}")
check('N=3 前二完全重合 2.0%', close(o3['前二完全重合比例'], 0.02008, 1e-4), f"{o3['前二完全重合比例']:.6f}")

# ---------------- diff
X = pd.read_csv('03_结果/执行差异_对比基准.csv')
check('执行差异 51 行', len(X) == 51, str(len(X)))
x03 = X[X.case_id == 'O03'].iloc[0]
check('O03 共同 99 / Jaccard 1', int(x03['共同信号数']) == 99 and close(x03['交易集合Jaccard'], 1.0))

# ---------------- current
C = pd.read_csv('03_结果/当前信号_2026-10-23.csv')
check('当前信号 2 只', len(C) == 2, str(len(C)))
check('当前信号代码', set(C['代码']) == {'603778.SH', '605376.SH'}, ','.join(C['代码']))
J = json.load(open('03_结果/指标.json'))
check('当前信号候选 123 只', int(J['signal_events_2026_10_23']) == 123,
      str(J['signal_events_2026_10_23']))

# ---------------- excel
from openpyxl import load_workbook
wb = load_workbook('不同入场时点回测.xlsx', data_only=True)
want = ['说明', '结论摘要', '组合汇总', '年度收益', '逐笔交易', '基准99笔对照',
        '信号级汇总', '信号级明细_前二', '信号集合变化', '执行差异',
        '跳过统计', '当前信号', '口径与参数']
check('Excel 13 个表', wb.sheetnames == want, str(wb.sheetnames))
check('Excel 组合汇总 52 行', wb['组合汇总'].max_row == 52, str(wb['组合汇总'].max_row))
check('Excel 逐笔交易 4866 行', wb['逐笔交易'].max_row == 4866, str(wb['逐笔交易'].max_row))
check('Excel 基准99笔 100 行', wb['基准99笔对照'].max_row == 100, str(wb['基准99笔对照'].max_row))
check('Excel 年化百分比格式', wb['组合汇总']['F2'].number_format == '0.0%',
      wb['组合汇总']['F2'].number_format)

# ---------------- reproduce base
sys.path.insert(0, '02_代码')
import timing_lib as TL
D = TL.load()
TL.add_signal(D)
eq, tr, sk = TL.simulate_timing(D, ev='custom', nth=3, fill='open', hold=15,
                                slots=2, tp=0.20, cost=0.00175,
                                start='2017-01-01', end='2026-10-23')
m = TL.metrics(eq, tr)
check('运行复现 O03 99 笔', len(tr) == 99, str(len(tr)))
check('运行复现 O03 40.0%', close(m['cagr'], 0.4001, 1e-3), f"{m['cagr']:.6f}")
check('运行复现 O03 -16.0%', close(m['dd'], -0.16, 1e-4), f"{m['dd']:.6f}")

# compare to母策略
old_path = '../2026-10-24_主板恐慌反转_年化30/03_结果/最终_逐笔.csv'
if os.path.exists(old_path):
    old = pd.read_csv(old_path, parse_dates=['entry_date', 'signal_date', 'exit_date'])
    old = old.sort_values(['signal_date', 'thscode']).reset_index(drop=True)
    new = tr.sort_values(['signal_date', 'thscode']).reset_index(drop=True)
    same_keys = set(zip(old.signal_date, old.thscode)) == set(zip(new.signal_date, new.thscode))
    same_dates = (old.entry_date.values == new.entry_date.values).all() and \
                 (old.exit_date.values == new.exit_date.values).all()
    maxdiff = np.abs(old.ret.values - new.ret.values).max()
    check('母策略 99 笔键一致', same_keys)
    check('母策略 99 笔日期一致', same_dates)
    check('母策略 99 笔收益差异<1e-6', maxdiff < 1e-6, f'{maxdiff:.2e}')

# ---------------- write
lines = ['不同入场时点回测 — 交付校验', f'时间：{pd.Timestamp.now():%Y-%m-%d %H:%M:%S}', '']
for name, ok, detail in results:
    lines.append(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" | {detail}" if detail else ''))
n_fail = sum(1 for _, ok, _ in results if not ok)
lines.append('')
lines.append(f'合计 {len(results)} 项，PASS {len(results)-n_fail}，FAIL {n_fail}')
open(f'{OUT}/交付校验.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('\n'.join(lines[-8:]))

# hashes
targets = ['README.md', '00_一页纸_不同入场时点.md', '01_完整报告_不同入场时点.md',
           '03_数据与复现.md', '不同入场时点回测.xlsx',
           '03_结果/组合汇总_全部时点.csv', '03_结果/逐笔_组合_全部时点.csv',
           '03_结果/基准99笔_不同时点收益明细.csv', '03_结果/信号级汇总_全部信号.csv']
hlines = []
for f in targets:
    if os.path.exists(f):
        h = hashlib.sha256(open(f, 'rb').read()).hexdigest()
        hlines.append(f'{h}  {f}')
open(f'{OUT}/关键文件sha256.txt', 'w', encoding='utf-8').write('\n'.join(hlines))

sys.exit(1 if n_fail else 0)
