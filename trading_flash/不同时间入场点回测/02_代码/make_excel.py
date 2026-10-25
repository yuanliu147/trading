# -*- coding: utf-8 -*-
"""把全部结果写入中文 Excel（中文表名、中文表头、百分比格式）。"""
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = '不同入场时点回测.xlsx'
S = pd.read_csv('03_结果/组合汇总_全部时点.csv')
A = pd.read_csv('03_结果/年度收益_全部时点.csv')
T = pd.read_csv('03_结果/逐笔_组合_全部时点.csv')
G = pd.read_csv('03_结果/信号级汇总_全部信号.csv')
D = pd.read_csv('03_结果/信号级明细_每日前二.csv')
O = pd.read_csv('03_结果/信号集合变化_按档位.csv')
X = pd.read_csv('03_结果/执行差异_对比基准.csv')
W = pd.read_csv('03_结果/基准99笔_不同时点收益明细.csv')
K = pd.read_csv('03_结果/跳过统计_全部时点.csv')
C = pd.read_csv('03_结果/当前信号_2026-10-23.csv')
CD = pd.read_csv('03_结果/当前信号_入场日期.csv')
J = pd.read_json('03_结果/指标.json', typ='series')

# ---------------- rename summary
S = S.rename(columns={
    'case_id': '编号', '类型': '类型', '入场方式': '入场方式',
    '第N个交易日': '第N个交易日', '说明': '说明',
    'cagr': '年化', 'dd': '最大回撤', 'calmar': 'Calmar',
    'sharpe': '夏普', 'final': '期末倍数', '期末权益_元': '期末权益(元)',
    '总利润_元': '总利润(元)', 'total_ret': '总收益',
    'n': '笔数', 'win': '胜率', 'avg': '平均单笔', 'median': '中位单笔',
    'best': '最好单笔', 'worst': '最差单笔', 'pf': '利润因子',
    'avg_win': '平均盈利', 'avg_loss': '平均亏损',
    'n_target': '止盈笔数', 'n_time': '到期笔数', 'n_eod': '期末笔数',
    'expo': '持仓时间占比', '基准年化': '基准年化',
    '年化差_vs基准': '年化差(vs基准)', '基准回撤': '基准回撤',
    '回撤差_vs基准': '回撤差(vs基准)', '是否基准': '是否基准',
    '首笔入场': '首笔入场', '末笔入场': '末笔入场'
})

PCT_COLS = ['年化', '最大回撤', '总收益', '胜率', '平均单笔', '中位单笔',
            '最好单笔', '最差单笔', '平均盈利', '平均亏损',
            '持仓时间占比', '基准年化', '年化差(vs基准)', '基准回撤',
            '回撤差(vs基准)']

# ---------------- annual rename
A = A.rename(columns={'case_id': '编号', '类型': '类型', '入场方式': '入场方式',
                      '第N个交易日': '第N个交易日', '说明': '说明'})

# ---------------- trades rename
T = T.rename(columns={
    'case_id': '编号', '类型': '类型', '入场方式': '入场方式',
    '第N个交易日': '第N个交易日', '说明': '说明',
    '信号日': '信号日', '入场日': '入场日', '退出日': '退出日',
    '代码': '代码', '信号排名': '信号排名',
    '250日回撤': '250日回撤(%)', '250日收益': '250日收益(%)',
    'RSI14': 'RSI14', '收盘位置': '收盘位置', '20日均额_亿': '20日均额(亿)',
    '入场价_后复权': '入场价(后复权)', '退出价_后复权': '退出价(后复权)',
    '退出原因': '退出原因', '单笔收益': '单笔收益(%)',
    '分配额_初始本金%': '分配额(初始本金%)',
    '仓位占当时权益%': '仓位占当时权益(%)',
    '盈亏贡献_初始本金%': '盈亏贡献(初始本金%)',
    '持有交易日': '持有交易日', '市场20日': '市场20日(%)'
})
T['退出原因'] = T['退出原因'].map({'target': '+20%止盈', 'time': '15日到期', 'eod': '期末未平'}).fillna(T['退出原因'])

# ---------------- signal summary rename
G = G.rename(columns={'case_id': '编号', '类型': '类型', '入场方式': '入场方式',
                      '第N个交易日': '第N个交易日', '说明': '说明'})
sig_pct = []
for c in G.columns:
    if c in ('编号', '类型', '入场方式', '第N个交易日', '说明', '信号总数'):
        continue
    if any(k in c for k in ['平均', '中位', '胜率', '涨幅', '追高', '差']):
        G[c] = G[c] * 100
        sig_pct.append(c)
G = G.rename(columns={c: c + '(%)' for c in sig_pct})

# ---------------- detail rename
D = D.rename(columns={'case_id': '编号', '入场方式': '入场方式',
                      '第N个交易日': '第N个交易日', '说明': '说明',
                      '信号日': '信号日', '代码': '代码', '信号排名': '信号排名',
                      '状态': '状态', '入场日': '入场日', '退出日': '退出日',
                      '单笔收益': '单笔收益(%)', '入场前涨幅': '入场前涨幅(%)',
                      '退出原因': '退出原因'})
D['退出原因'] = D['退出原因'].map({'target': '+20%止盈', 'time': '15日到期', 'eod': '期末未平'}).fillna(D['退出原因'])

# ---------------- overlap / diff rename
O = O.rename(columns={'第N个交易日': '第N个交易日', '信号日数': '信号日数',
                      '平均交集': '平均交集(只)', '平均并集': '平均并集(只)',
                      '平均Jaccard': '平均Jaccard', '前二平均重合只数': '前二平均重合(只)',
                      '前二完全重合比例': '前二完全重合比例', '前二存活率': '前二存活率'})
X = X.rename(columns={'case_id': '编号', '类型': '类型', '入场方式': '入场方式',
                      '第N个交易日': '第N个交易日', '交易数': '交易数',
                      '共同信号数': '共同信号数', '本档独有数': '本档独有数',
                      '基准独有数': '基准独有数', '交易集合Jaccard': '交易集合Jaccard',
                      '共同信号_本档平均': '共同信号_本档平均(%)',
                      '共同信号_基准平均': '共同信号_基准平均(%)',
                      '共同信号_差': '共同信号_差(%)',
                      '本档独有_平均': '本档独有_平均(%)',
                      '基准独有_平均': '基准独有_平均(%)',
                      '共同信号_本档贡献': '共同信号_本档贡献(%)',
                      '共同信号_基准贡献': '共同信号_基准贡献(%)',
                      '本档独有_贡献': '本档独有_贡献(%)',
                      '基准独有_贡献': '基准独有_贡献(%)'})
C = C.rename(columns={'250日收益': '250日收益(%)', '250日回撤': '250日回撤(%)'})

# ---------------- write
with pd.ExcelWriter(OUT, engine='openpyxl') as writer:
    # 说明
    notes = pd.DataFrame([
        ('标题', '主板“恐慌洗盘反转”：不同入场时点全量回测'),
        ('数据窗口', '2017-10-19 至 2026-10-23（面板从 2017-10-19 起，2017 为部分年度；2026 截至 10-23）'),
        ('信号（T 日收盘确认）', '主板；上市≥250日；20日均额≥1000万；剔除5%涨跌停样态；250日收益>0；250日回撤<-45%；RSI14<35；收盘位置>0.45；主板等权20日<-5%'),
        ('排序', '同日按250日回撤升序取前2；2仓各50%当前权益；无信号空仓'),
        ('入场档位', '信号后第 N 个交易日。N=1 次日；N=3 为母策略基准（2026-10-23 信号 → 2026-10-28 开盘）'),
        ('价格口径', '主口径：开盘价/收盘价；边界：当日最低价（乐观）/最高价（悲观）；另有 T 日收盘（前视参照）'),
        ('退出', '+20% 止盈（T+1 起，盘中触价）；持有15个交易日收盘卖；无普通止损；期末未平按最后收盘'),
        ('成本', '佣金万2.5（最低5元）+ 印花税千0.5（卖）+ 滑点千1（单边）；复权模拟按 0.175%/边'),
        ('主结果', 'O03（第3个交易日开盘）：年化 40.0%、最大回撤 -16.0%、99笔、胜率 71.7%、期末 20.75 倍、利润 197.5 万（10 万本金）'),
        ('全区间', '主档位（开盘/收盘，N=1..15）年化 8.0%–40.0%；第3日开盘是峰值，第2–4日开盘为 30.3%/40.0%/37.7%'),
        ('边界', '第3日最低价 52.3%、开盘 40.0%、收盘 27.6%、最高价 19.2%；同一天成交价可差 33pp'),
        ('信号点', '信号点在 T 日固定；延后入场不重算条件。若在入场日重新扫描，信号集合已大不相同：第3日平均 Jaccard 6.9%，前二存活率 13.5%'),
        ('信号级表', '全部=4195条信号逐条独立模拟；前二=每日回撤最深前2；基准99=O03组合实际成交的99条信号；收益均为单笔、扣0.175%/边'),
        ('基准99笔对照', '行=O03实际成交的99条信号；列=各档位下同一信号的独立单笔收益（%）；O03列应等于“基准组合收益”列'),
        ('信号集合变化', '同一信号日 T 与 T+N 的信号集合比较；Jaccard=交集/并集；前二存活率=T日前二在T+N仍有信号的比例；前二完全重合=T日与T+N前二完全相同的信号日比例'),
        ('执行差异', '每档 vs 基准O03：共同信号=两档都成交的(信号日,代码)；共同信号_差=本档平均-基准平均；本档独有/基准独有=因2仓占位不同而多/少做的交易'),
        ('CSV 单位', 'CSV 中组合汇总、信号级汇总、年度收益的收益/回撤/胜率是小数（0.40=40%）；逐笔、执行差异、基准99笔对照是百分数（9.05=9.05%）。Excel 已按列设为百分比格式或百分数表头'),
        ('重要限制', '参数是同一段历史里试出来的；40.0% 是峰值不是承诺；无盘中数据（只近似开盘/收盘/高低）；无退市股；未对全部档位做整手实际价复核'),
        ('不构成', '历史回测不是投资建议；先纸面，后小资金'),
    ], columns=['项目', '内容'])
    notes.to_excel(writer, sheet_name='说明', index=False)

    # 结论摘要：关键档位
    keep = ['C00', 'O01', 'O02', 'O03', 'O04', 'O05', 'O06', 'O07', 'O08',
            'O09', 'O10', 'O11', 'O12', 'O13', 'O14', 'O15',
            'C01', 'C02', 'C03', 'C04', 'C05', 'L03', 'H03']
    S[S['编号'].isin(keep)].to_excel(writer, sheet_name='结论摘要', index=False)
    S.to_excel(writer, sheet_name='组合汇总', index=False)
    A.to_excel(writer, sheet_name='年度收益', index=False)
    T.to_excel(writer, sheet_name='逐笔交易', index=False)
    W.to_excel(writer, sheet_name='基准99笔对照', index=False)
    G.to_excel(writer, sheet_name='信号级汇总', index=False)
    D.to_excel(writer, sheet_name='信号级明细_前二', index=False)
    O.to_excel(writer, sheet_name='信号集合变化', index=False)
    X.to_excel(writer, sheet_name='执行差异', index=False)
    K.to_excel(writer, sheet_name='跳过统计', index=False)

    # 当前信号 + 入场日期
    cur = C.copy()
    cur.to_excel(writer, sheet_name='当前信号', index=False, startrow=0)
    pd.DataFrame([[]]).to_excel(writer, sheet_name='当前信号', index=False, header=False, startrow=len(cur) + 2)
    CD.to_excel(writer, sheet_name='当前信号', index=False, startrow=len(cur) + 3)

    # 口径与参数
    params = pd.DataFrame([
        ('策略', '主板恐慌洗盘反转（母策略：trading_flash/2026-10-24_主板恐慌反转_年化30）'),
        ('起始资金', '100,000 元'),
        ('仓位数', '2'),
        ('单仓', '当前权益的 50%'),
        ('持有', '15 个交易日（含入场日）'),
        ('止盈', '+20%（T+1 起，盘中触价；开盘跳空高于止盈按开盘）'),
        ('止损', '无普通止损'),
        ('成本', '0.175%/边（复权模拟）'),
        ('一字涨停', '入场实际价 ≥ 前收×1.098 不买；被跳过不递补'),
        ('停牌', '入场日无行情不买'),
        ('复权', '后复权价计算收益；涨跌停/整手用不复权实际价'),
        ('样本', f"信号事件 {J['signal_events']:,} 条；信号日 {J['signal_days']} 天；标的 {J['signal_symbols']:,} 只"),
        ('当前信号', f"2026-10-23 共 {J['signal_events_2026_10_23']} 只；前二 603778.SH、605376.SH"),
        ('基准', 'O03：信号后第3个交易日开盘（母策略口径）'),
        ('主档位范围', f"年化 {J['main_cagr_min']*100:.1f}%–{J['main_cagr_max']*100:.1f}%，中位 {J['main_cagr_median']*100:.1f}%"),
        ('最差主档位', f"{J['worst_main']['case_id']} {J['worst_main']['label']}：{J['worst_main']['cagr']*100:.1f}% / 回撤 {J['worst_main']['dd']*100:.1f}%"),
        ('最好主档位', f"{J['best_main']['case_id']} {J['best_main']['label']}：{J['best_main']['cagr']*100:.1f}%"),
        ('数据限制', '本地库不含退市股；无盘中分时；行业分类缺失；2017 与 2026 为部分年度'),
    ], columns=['参数/项目', '值'])
    params.to_excel(writer, sheet_name='口径与参数', index=False)

# ---------------- format
from openpyxl import load_workbook
wb = load_workbook(OUT)

HEAD_FILL = PatternFill('solid', fgColor='1F4E79')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)


def style_sheet(ws, pct_cols=(), pct_units=(), num_cols=(), yuan_cols=(),
                multi_cols=(), widths=None, freeze='A2', row_height=None):
    # header
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if freeze:
        ws.freeze_panes = freeze
    headers = [c.value for c in ws[1]]
    for j, h in enumerate(headers, 1):
        if h is None:
            continue
        letter = get_column_letter(j)
        if h in pct_cols:
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                for c in row:
                    c.number_format = '0.0%'
        elif h in pct_units:
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                for c in row:
                    c.number_format = '0.00'
        elif h in yuan_cols:
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                for c in row:
                    c.number_format = '#,##0'
        elif h in multi_cols:
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                for c in row:
                    c.number_format = '0.00'
        elif h in num_cols:
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                for c in row:
                    c.number_format = '0'
        if widths and h in widths:
            ws.column_dimensions[letter].width = widths[h]
    if row_height:
        for r in range(1, ws.max_row + 1):
            ws.row_dimensions[r].height = row_height


style_sheet(wb['说明'], widths={'项目': 20, '内容': 120})
for r in range(2, wb['说明'].max_row + 1):
    wb['说明'].cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')

style_sheet(wb['结论摘要'], pct_cols=PCT_COLS,
            yuan_cols=('期末权益(元)', '总利润(元)'),
            multi_cols=('期末倍数', 'Calmar', '夏普', '利润因子'),
            num_cols=('第N个交易日', '笔数', '止盈笔数', '到期笔数', '期末笔数', '是否基准'),
            widths={'编号': 7, '类型': 6, '入场方式': 7, '第N个交易日': 8, '说明': 30,
                    '期末权益(元)': 14, '总利润(元)': 14})
style_sheet(wb['组合汇总'], pct_cols=PCT_COLS,
            yuan_cols=('期末权益(元)', '总利润(元)'),
            multi_cols=('期末倍数', 'Calmar', '夏普', '利润因子'),
            num_cols=('第N个交易日', '笔数', '止盈笔数', '到期笔数', '期末笔数', '是否基准'),
            widths={'编号': 7, '类型': 6, '入场方式': 7, '第N个交易日': 8, '说明': 30,
                    '期末权益(元)': 14, '总利润(元)': 14})
style_sheet(wb['年度收益'], pct_cols=tuple(str(y) for y in range(2017, 2027)),
            num_cols=('第N个交易日',), widths={'编号': 7, '说明': 30})
style_sheet(wb['逐笔交易'],
            pct_units=('250日回撤(%)', '250日收益(%)', '单笔收益(%)',
                       '分配额(初始本金%)', '仓位占当时权益(%)',
                       '盈亏贡献(初始本金%)', '市场20日(%)'),
            num_cols=('信号排名', '持有交易日'),
            multi_cols=('入场价(后复权)', '退出价(后复权)', 'RSI14', '收盘位置', '20日均额(亿)'),
            widths={'编号': 7, '说明': 28, '代码': 11, '信号日': 11, '入场日': 11, '退出日': 11})
style_sheet(wb['基准99笔对照'],
            pct_units=tuple(['基准组合收益'] + [c for c in W.columns if c.startswith(('O', 'C', 'L', 'H')) and c[1:].isdigit()]),
            multi_cols=('RSI14', '收盘位置'), num_cols=('信号排名',),
            widths={'信号日': 11, '代码': 11, '基准入场日': 11, '基准退出日': 11})
style_sheet(wb['信号级汇总'], pct_units=tuple(c + '(%)' for c in sig_pct),
            multi_cols=('全部_利润因子', '前二_利润因子', '基准99_利润因子'),
            num_cols=('第N个交易日', '信号总数') if '信号总数' in G.columns else ('第N个交易日',),
            widths={'编号': 7, '说明': 28})
style_sheet(wb['信号级明细_前二'], pct_units=('单笔收益(%)', '入场前涨幅(%)'),
            num_cols=('信号排名',), widths={'编号': 7, '说明': 26, '代码': 11, '信号日': 11, '入场日': 11, '退出日': 11})
style_sheet(wb['信号集合变化'], pct_cols=('平均Jaccard', '前二完全重合比例', '前二存活率'),
            multi_cols=('平均交集(只)', '平均并集(只)', '前二平均重合(只)'),
            num_cols=('第N个交易日', '信号日数'))
style_sheet(wb['执行差异'], pct_units=tuple(c for c in X.columns if '(%)' in c),
            multi_cols=('交易集合Jaccard',), num_cols=('交易数', '共同信号数', '本档独有数', '基准独有数'))
style_sheet(wb['跳过统计'], num_cols=('第N个交易日', '次数'))
style_sheet(wb['当前信号'], multi_cols=('收盘', 'RSI14', '收盘位置'),
            pct_units=('250日收益(%)', '250日回撤(%)'),
            num_cols=('排名',), widths={'代码': 12, '信号日': 11})
style_sheet(wb['口径与参数'], widths={'参数/项目': 18, '值': 120})
for r in range(2, wb['口径与参数'].max_row + 1):
    wb['口径与参数'].cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')

wb.save(OUT)
print('excel done', OUT)
