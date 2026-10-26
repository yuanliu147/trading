# -*- coding: utf-8 -*-
"""历史年份同态：交付校验。"""
from __future__ import annotations

import json
import math
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve()
OUT = HERE.parents[1]
RES = OUT / '03_结果'
FIG = OUT / '04_图表'
REPORT = OUT / '05_证据' / '校验.txt'

checks: list[tuple[bool, str]] = []


def check(cond: bool, msg: str):
    checks.append((bool(cond), msg))


def approx(a, b, tol=1e-4):
    return abs(float(a) - float(b)) <= tol


# ------------------------------------------------------------------ 输入
m = pd.read_parquet('/Users/admin/WorkSpace/trading/trading_hithink/初始研究/data/work/market_daily.parquet')
m['date'] = pd.to_datetime(m['date'])
m = m.sort_values('date').reset_index(drop=True)
cur = m.iloc[-1]
year = pd.read_csv(RES / '年度相似度.csv')
nn = pd.read_csv(RES / '最近邻_核心8.csv', parse_dates=['date'])
ev = pd.read_csv(RES / '同态事件_总表.csv')
st = pd.read_csv(RES / '同态期_策略起点.csv')
th = pd.read_csv(RES / '同态阈值汇总.csv')
rule = pd.read_csv(RES / '规则状态_事件.csv')
pred = pd.read_csv(RES / '距离预测力.csv')
state = pd.read_csv(RES / '当前状态.csv')
annual = pd.read_csv(OUT.parents[1] / '2026-10-24_主板恐慌反转_年化30/03_结果/最终_年度.csv',
                     parse_dates=['date'])
signal = pd.read_csv(OUT.parents[1] / '不同时间入场点回测/07_市场阶段适配/03_结果/全部信号日逐只_带阶段.csv',
                     parse_dates=['signal_date'])
starts = pd.read_csv(OUT.parents[1] / '不同时间入场点回测/07_市场阶段适配/03_结果/起点固定持有期.csv',
                     parse_dates=['signal_date'])

# 1 当前
check(cur['date'] == pd.Timestamp('2026-10-23'), f"当前日期 = 2026-10-23（实际 {cur['date'].date()}）")
check(approx(cur['equal_weight_index'], 1.743299, 5e-6), f"当前等权指数 = 1.743299（实际 {cur['equal_weight_index']:.6f}）")
check(approx(cur['market_volatility20'], 0.197754, 5e-5), '当前 20 日波动率 = 0.1978')
check(approx(cur['breadth20'], 0.189969, 5e-5), '当前宽度20 = 0.1900')
check(len(state) == 13, f'当前状态表 13 行（实际 {len(state)}）')

# 2 年度相似度
y18 = year.loc[year['年份'] == 2018].iloc[0]
y24 = year.loc[year['年份'] == 2024].iloc[0]
check('核心8最近日期' in year.columns, '年度表含核心8最近日期列')
check(y18['全状态13_排名'] == 1 and y24['全状态13_排名'] == 2, '全状态13：2018 #1，2024 #2')
check(y24['核心8_排名'] == 1 and y18['核心8_排名'] == 2, '核心8：2024 #1，2018 #2')
check(approx(y18['全状态13_距离'], 0.2986, 5e-4), f"2018 全状态距离 ≈0.299（实际 {y18['全状态13_距离']:.4f}）")
check(approx(y24['全状态13_距离'], 0.3221, 5e-4), f"2024 全状态距离 ≈0.322（实际 {y24['全状态13_距离']:.4f}）")
check(approx(y18['核心8_距离'], 0.1720, 5e-4), f"2018 核心8距离 ≈0.172（实际 {y18['核心8_距离']:.4f}）")
check(approx(y24['核心8_距离'], 0.1577, 5e-4), f"2024 核心8距离 ≈0.158（实际 {y24['核心8_距离']:.4f}）")
check(year['核心8_距离'].nsmallest(3).iloc[2] > 0.39, '核心8 第三名距离 > 0.39')
check(year['综合平均排名'].nsmallest(1).iloc[0] == 1.25, '2018 综合平均排名 1.25')

# 3 最近邻
first = nn.iloc[0]
check(first['date'] == pd.Timestamp('2024-06-14'), '核心8 最近邻 = 2024-06-14')
check(approx(first['d_核心8'], 0.1136, 5e-4), f"最近邻距离 ≈0.114（实际 {first['d_核心8']:.4f}）")
check((nn['date'] == pd.Timestamp('2018-05-30')).any(), '2018-05-30 在核心8前40')
check(len(nn) == 40, '核心8 最近邻表 40 行')

# 4 事件
e18 = ev.loc[ev['代表日'] == '2018-05-30'].iloc[0]
e24 = ev.loc[ev['代表日'] == '2024-06-14'].iloc[0]
check(approx(e18['市场60日'], -0.14763, 5e-4), '2018-05-30 市场60日 -14.8%')
check(approx(e18['市场120日'], -0.19475, 5e-4), '2018-05-30 市场120日 -19.5%')
check(approx(e24['市场60日'], -0.08115, 5e-4), '2024-06-14 市场60日 -8.1%')
check(approx(e24['市场120日'], 0.29725, 5e-4), '2024-06-14 市场120日 +29.7%')
check(e18['后续最低点日期'] == '2018-10-18' and approx(e18['后续最低点回撤'], -0.291987, 5e-4) and int(e18['后续最低点交易日']) == 94, '2018 后续最低点 2018-10-18 / -29.2% / 94日')
check(e24['后续最低点日期'] == '2024-09-18' and approx(e24['后续最低点回撤'], -0.099653, 5e-4) and int(e24['后续最低点交易日']) == 66, '2024 后续最低点 2024-09-18 / -10.0% / 66日')
check(len(ev) == 7, f'同态事件 7 个（实际 {len(ev)}）')
check((ev['主同态'] == '是').sum() == 2, '主同态 2 个（≤0.25）')
check(ev[['市场20日', '市场60日', '市场120日', '主同态']].notna().all().all(), '事件表关键列无缺失')

# 5 策略起点
s18 = st.loc[st['代表日'] == '2018-05-30'].iloc[0]
s24 = st.loc[st['代表日'] == '2024-06-14'].iloc[0]
check(approx(s18['下一可执行起点60日'], 0.098312, 5e-4), '2018 下一开始点60日 +9.8%')
check(approx(s18['下一可执行起点120日'], 0.160114, 5e-4), '2018 下一开始点120日 +16.0%')
check(approx(s24['下一可执行起点60日'], 0.025276, 5e-4), '2024 下一开始点60日 +2.5%')
check(approx(s24['下一可执行起点120日'], 0.113739, 5e-4), '2024 下一开始点120日 +11.4%')
check(approx(s18['区间起点120日回撤中位'], -0.131685, 5e-4), '2018 区间120日回撤中位 -13.2%')
check(approx(s24['区间起点120日回撤中位'], -0.124719, 5e-4), '2024 区间120日回撤中位 -12.5%')
check(s18['起点数'] == 3 and s24['起点数'] == 11, '2018/2024 起点数 3/11')

# 6 阈值 / 规则 / 预测力
check(th.loc[th['阈值'] == 0.25, '事件数'].iloc[0] == 2, '阈值0.25 事件数 2')
check(th.loc[th['阈值'] == 0.35, '事件数'].iloc[0] == 7, '阈值0.35 事件数 7')
check(len(rule) == 4, f'规则状态事件 4 个（实际 {len(rule)}）')
check(len(pred) == 3, '距离预测力 3 行')

# 7 年度收益 / 交易
a18 = annual.loc[annual['date'].dt.year == 2018, 'return'].iloc[0]
a24 = annual.loc[annual['date'].dt.year == 2024, 'return'].iloc[0]
check(approx(a18, 0.474056, 5e-4), '2018 策略年度 +47.4%')
check(approx(a24, 0.929837, 5e-4), '2024 策略年度 +93.0%')
for yy, n in [(2018, 68), (2024, 61)]:
    g = signal[(signal['signal_date'].dt.year == yy) & (signal['complete'] == True)]  # noqa: E712
    check(len(g) == n, f'{yy} 完整信号逐只 {n} 笔（实际 {len(g)}）')
    check(approx(g['ret'].mean(), 0.042583 if yy == 2018 else 0.064706, 5e-4),
          f'{yy} 单笔均值与 07 一致')
    check(approx((g['ret'] > 0).mean(), 0.632353 if yy == 2018 else 0.655738, 5e-4),
          f'{yy} 胜率与 07 一致')

check(len(starts) == 251, f'起点文件 251 行（实际 {len(starts)}）')
check(signal['complete'].sum() == 450, f'完整信号逐只 450 行（实际 {int(signal["complete"].sum())}）')

# 8 文件
for name in ['01_年度相似度.png', '02_同态路径.png', '03_距离与未来60日.png', '04_市场与策略.png']:
    p = FIG / name
    check(p.exists() and p.stat().st_size > 50_000, f'图表 {name} 存在且 >50KB')
xlsx = OUT / '历史年份同态.xlsx'
check(xlsx.exists(), 'Excel 存在')
wb = load_workbook(xlsx)
check(len(wb.sheetnames) == 8, f'Excel 8 个工作表（实际 {len(wb.sheetnames)}）')
check(wb.sheetnames == ['说明', '当前状态', '年度相似度', '同态事件', '策略起点', '规则状态', '距离预测力', '最近邻'],
      'Excel 工作表顺序正确')

html = OUT / '历史年份同态.html'
check(html.exists() and html.stat().st_size > 400_000, 'HTML 存在且 >400KB')
h = html.read_text(encoding='utf-8')
check(h.count('data:image/png;base64,') == 4, 'HTML 内嵌 4 张图')
check('2018-05-30' in h and '2024-06-14' in h and '0.299' in h and '0.322' in h,
      'HTML 含关键日期与更新后的距离')
check('0.308' not in h, 'HTML 不含旧的 0.308 距离')

for name in ['00_一页纸_历史年份同态.md', '01_完整报告_历史年份同态.md', '03_数据与复现.md', 'README.md']:
    check((OUT / name).exists(), f'{name} 存在')

# 9 关键数字一致性
check(approx(state.loc[state['指标'] == '250日涨幅', '当前值'].iloc[0], -0.05151, 5e-5),
      '当前状态 250日涨幅 -5.2%')
check(approx(state.loc[state['指标'] == '20日波动率', '当前值'].iloc[0], 0.197754, 5e-5),
      '当前状态 20日波动率 0.198')
check(approx(state.loc[state['指标'] == '跌停占比', '当前值'].iloc[0], 0.002507, 5e-6),
      '当前状态 跌停占比 0.25%')

# ------------------------------------------------------------------ 输出
lines = ['历史年份同态 · 交付校验', f'日期：{cur["date"].date()}', '']
passed = sum(1 for ok, _ in checks if ok)
lines.append(f'通过：{passed}/{len(checks)}')
lines.append('')
for ok, msg in checks:
    lines.append(('✓' if ok else '✗') + ' ' + msg)
lines.append('')
lines.append('数据：沪深主板等权 market_daily.parquet；策略 07 起点固定持有期.csv；年度 最终_年度.csv。')
lines.append('注意：相似度为描述性口径，标准化含全样本；同态日高度重叠，不是独立样本。')
REPORT.write_text('\n'.join(lines), encoding='utf-8')

print('\n'.join(lines))
if passed != len(checks):
    raise SystemExit(1)
