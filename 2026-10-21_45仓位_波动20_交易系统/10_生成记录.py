# -*- coding: utf-8 -*-
"""生成 45% + 波动20 独立系统的逐笔记录、年度统计、Excel 和权益图。"""
import os
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

plt.rcParams['font.sans-serif'] = ['PingFang SC', 'Arial Unicode MS', 'Heiti SC']
plt.rcParams['axes.unicode_minus'] = False

SRC = '/Users/admin/WorkSpace/trading/2026-10-20_涨幅趋势_全维度定制'
OUT = '/Users/admin/WorkSpace/trading/2026-10-21_45仓位_波动20_交易系统'

t = pd.read_csv(os.path.join(SRC, '03_结果/回撤优化_逐笔_收益优先_45_波动20.csv'),
                parse_dates=['entry_date', 'exit_date'])
e = pd.read_csv(os.path.join(SRC, '03_结果/回撤优化_权益_收益优先_45_波动20.csv'),
                parse_dates=['date']).set_index('date')
cand = pd.read_csv(os.path.join(SRC, '03_结果/回撤优化_候选总表.csv'))
row = cand[cand['方案'] == '收益优先_45_波动20'].iloc[0]

# ---------- 逐笔增强 ----------
t = t.sort_values(['entry_date', 'code']).reset_index(drop=True)
t['序号'] = np.arange(1, len(t) + 1)
t['买入金额'] = t['entry'] * t['shares']
t['卖出金额'] = t['exit'] * t['shares']
t['毛盈亏'] = (t['exit'] - t['entry']) * t['shares']
t['费用'] = t['毛盈亏'] - t['pnl']
t['账户贡献'] = t['weight'] * t['ret']
t['年份'] = t['exit_date'].dt.year
t['盈亏'] = np.where(t['pnl'] >= 0, '盈利', '亏损')

trade = t.rename(columns={
    'code': '代码', 'entry_date': '买入日', 'exit_date': '卖出日', 'entry': '买入价',
    'exit': '卖出价', 'ret': '收益率', 'reason': '退出原因', 'hold': '持有天数',
    'shares': '股数', 'pnl': '净盈亏', 'weight': '买入仓位',
})
cols = ['序号', '代码', '买入日', '卖出日', '持有天数', '买入价', '卖出价', '股数',
        '买入金额', '卖出金额', '毛盈亏', '费用', '净盈亏', '收益率', '买入仓位',
        '账户贡献', '退出原因', '盈亏', '年份']
trade = trade[cols]
trade.to_csv(os.path.join(OUT, '逐笔交易记录.csv'), index=False, encoding='utf-8-sig')

# ---------- 年度统计 ----------
def pf(x):
    pos = x[x > 0].sum(); neg = -x[x < 0].sum()
    return pos / neg if neg > 0 else np.nan

yearly = t.groupby('年份').agg(
    交易笔数=('pnl', 'size'), 净盈亏=('pnl', 'sum'), 胜率=('ret', lambda x: (x > 0).mean()),
    平均收益率=('ret', 'mean'), 平均账户贡献=('账户贡献', 'mean'),
    最大单笔盈利=('pnl', 'max'), 最大单笔亏损=('pnl', 'min'),
).reset_index()
yearly['盈亏比'] = t.groupby('年份')['pnl'].apply(pf).values
yearly.to_csv(os.path.join(OUT, '年度统计.csv'), index=False, encoding='utf-8-sig')

# ---------- 权益增强 ----------
eq = e.copy()
eq['回撤'] = eq['equity'] / eq['equity'].cummax() - 1
eq['最低价回撤'] = eq['equity_low'] / eq['equity'].cummax() - 1
eq_out = eq.reset_index().rename(columns={
    'date': '日期', 'equity': '权益', 'cash': '现金', 'npos': '持仓数',
    'equity_low': '最低价权益', 'ret': '日收益', 'exposure': '总仓位',
})
eq_out.to_csv(os.path.join(OUT, '权益曲线.csv'), index=False, encoding='utf-8-sig')

# ---------- 汇总 ----------
final = e.equity.iloc[-1]
cagr = (final / 20000) ** (252 / len(e)) - 1
maxdd = (e.equity / e.equity.cummax() - 1).min()
maxdd_low = (e.equity_low / e.equity.cummax() - 1).min()
win = (t.ret > 0).mean()
pf_all = t.pnl[t.pnl > 0].sum() / -t.pnl[t.pnl < 0].sum()
worst = t['账户贡献'].min()
best = t['账户贡献'].max()
top5 = t['账户贡献'].nlargest(5).sum() / np.log(final / 20000)
summary = pd.DataFrame([
    ('数据区间', '2018-01-01 至 2026-10-13'),
    ('本金', '20,000 元'),
    ('期末资金', f'{final:,.0f} 元'),
    ('全期连续年化', f'{cagr*100:.1f}%'),
    ('分段拼接年化', '17.3%'),
    ('训练 / 检验年化', f'{row["训练年化"]*100:.1f}% / {row["检验年化"]*100:.1f}%'),
    ('最大回撤（收盘）', f'{maxdd*100:.1f}%'),
    ('最大回撤（最低价）', f'{maxdd_low*100:.1f}%'),
    ('交易笔数', f'{len(t)} 笔'),
    ('胜率', f'{win*100:.1f}%'),
    ('盈亏比', f'{pf_all:.2f}'),
    ('平均持有', f'{t.hold.mean():.1f} 个交易日'),
    ('平均总仓位', f'{e.exposure.mean()*100:.0f}%'),
    ('最差单笔账户亏损', f'{worst*100:.1f}%'),
    ('最好单笔账户盈利', f'{best*100:.1f}%'),
    ('前五笔账户贡献占比', f'{top5*100:.0f}%'),
    ('去掉前五笔年化', f'{row["去前五笔年化"]*100:.1f}%'),
    ('成本 0.35% 年化/回撤', f'{row["成本0.35%年化"]*100:.1f}% / {row["成本0.35%回撤"]*100:.1f}%'),
], columns=['项目', '数值'])
summary.to_csv(os.path.join(OUT, '系统统计.csv'), index=False, encoding='utf-8-sig')

# ---------- 图表 ----------
fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(13, 8.6), sharex=True,
                               gridspec_kw={'height_ratios': [3, 1.15]})
fig.subplots_adjust(top=0.90, bottom=0.10, left=0.08, right=0.98, hspace=0.08)
ax1.plot(e.index, e.equity / 10000, color='#1F77B4', lw=1.8, label='账户权益')
ax1.fill_between(e.index, 2, e.equity / 10000, color='#1F77B4', alpha=.08)
ax1.axhline(2, color='black', lw=.6, alpha=.3)
ax1.set_ylabel('账户权益（万元）')
fig.suptitle('45% + 波动20：2万元账户权益与回撤（2018–2026.10）', fontsize=15, y=0.97)
ax1.text(.01, .88, f'期末 {final/10000:.1f} 万 ｜ 年化 {cagr*100:.1f}% ｜ 最大回撤 {maxdd*100:.1f}% ｜ '
         f'胜率 {win*100:.1f}% ｜ 盈亏比 {pf_all:.2f}', transform=ax1.transAxes,
         fontsize=9, va='top', color='#333333', bbox=dict(facecolor='white', alpha=.85, edgecolor='none'))
ax1.grid(alpha=.22); ax1.legend(frameon=False, fontsize=9, loc='upper left')
ax2.plot(e.index, eq['回撤'], color='#D62728', lw=1.1, label='收盘权益回撤')
ax2.plot(e.index, eq['最低价回撤'], color='#FF9896', lw=.8, alpha=.8, label='最低价回撤')
ax2.axhline(0, color='black', lw=.6, alpha=.3)
ax2.yaxis.set_major_formatter(FuncFormatter(lambda y, pos: f'{y:.0%}'))
ax2.set_ylabel('回撤'); ax2.grid(alpha=.22)
ax2.legend(frameon=False, fontsize=8, ncol=2, loc='lower left')
fig.text(0.5, 0.025, '口径：全期连续资金；2021年末清仓再起算17.3%，2022年新账户18.5%',
         ha='center', fontsize=8.5, color='#555555')
plt.savefig(os.path.join(OUT, '权益与回撤.png'), dpi=200)
plt.close()

# ---------- Excel ----------
xls = os.path.join(OUT, '逐笔交易记录.xlsx')
with pd.ExcelWriter(xls, engine='openpyxl') as w:
    summary.to_excel(w, sheet_name='系统统计', index=False)
    trade.to_excel(w, sheet_name='逐笔交易', index=False)
    yearly.to_excel(w, sheet_name='年度统计', index=False)
    rules = pd.DataFrame([
        ('市场闸门', '全A收盘在20日均线以上的比例 > 50%；不满足只平不开。'),
        ('趋势结构', '收盘 > 60日均线；20日均线 > 60日均线；60日均线高于20日前。'),
        ('趋势年龄', '5–30 个交易日。'),
        ('波动', '14日平均真实波幅 / 收盘 <= 4.4%。'),
        ('情绪', '过去20日涨停次数 = 0。'),
        ('板块', '行业宽度 >= 60%；行业指数20日涨幅 > 5%。'),
        ('位置', '距60日高点回撤 > -8%。'),
        ('可买性', '价格>=2元；20日成交额>=3000万元；次日开盘买；高开>5%、一字涨停、停牌不买。'),
        ('排序', '同日候选按20日成交额从小到大排序，取前2只。'),
        ('单票上限', '45% × 波动系数；波动系数=min(1, 20%/全A等权20日已实现波动)，下限0.3。'),
        ('股数', '目标金额 ÷ 买入价，向下取整到100股；不足100股跳过；股价>90元通常买不了。'),
        ('现金', '两只合计<=90%；现金不足减手；不融资。'),
        ('退出', '持有20个交易日收盘卖完；盘中-20%灾难止损；一字跌停不卖，顺延。'),
        ('不做', '不加仓、不摊平、不当天买回、不设窄止损、不设均线/移动止盈、不分批。'),
        ('成本', '单边0.15%或最低5元；实盘按0.25%–0.35%做预算。'),
    ], columns=['条款', '内容'])
    rules.to_excel(w, sheet_name='规则摘要', index=False)
    note = pd.DataFrame([
        ('来源', '2026-10-20_涨幅趋势_全维度定制 / 05_25%回测审计与回撤优化.md'),
        ('数据', '全A后复权日线，2016-10-13至2026-10-13；含102只退市/缺失股，共5675只。'),
        ('口径', '信号t日收盘，t+1开盘成交；整手、最低佣金、停牌/退市处理已计入。'),
        ('重要提醒', '全期连续年化21.6%；分段拼接17.3%，2022年新账户18.5%。不要把21.6%当未来预期。'),
        ('收益集中', '前5笔贡献约83%；去掉前5笔年化8.3%。'),
        ('尾部风险', '平稳块自举5%最坏回撤约-43.6%；历史-26.4%不是最坏情况。'),
        ('样本外', '2022–2026参与过规则选择，不是干净样本外；小市值退市股仍未补全。'),
        ('用途', '研究与纸面交易记录，不构成投资建议。'),
    ], columns=['项目', '说明'])
    note.to_excel(w, sheet_name='说明', index=False)

# 格式化
from openpyxl import load_workbook
wb = load_workbook(xls)
head_fill = PatternFill('solid', fgColor='1F4E79')
head_font = Font(bold=True, color='FFFFFF', size=10)
thin = Side(style='thin', color='D9D9D9')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for ws in wb.worksheets:
    for c in ws[1]:
        c.fill = head_fill; c.font = head_font; c.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        maxlen = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[letter].width = min(max(10, maxlen * 1.8), 42)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical='center', wrap_text=True)
# 数字格式
ws = wb['逐笔交易']
for r in range(2, ws.max_row + 1):
    for col, fmt in [('F', '0.00'), ('G', '0.00'), ('H', '#,##0'), ('I', '#,##0.00'),
                     ('J', '#,##0.00'), ('K', '#,##0.00'), ('L', '#,##0.00'), ('M', '#,##0.00'),
                     ('N', '0.00%'), ('O', '0.00%'), ('P', '0.00%')]:
        ws[f'{col}{r}'].number_format = fmt
ws = wb['年度统计']
for r in range(2, ws.max_row + 1):
    for col, fmt in [('B', '#,##0'), ('C', '#,##0.00'), ('D', '0.0%'), ('E', '0.00%'),
                     ('F', '0.00%'), ('G', '#,##0.00'), ('H', '#,##0.00'), ('I', '0.00')]:
        ws[f'{col}{r}'].number_format = fmt
wb.save(xls)

print('saved', OUT)
print(summary.to_string(index=False))
print('\n年度统计')
print(yearly.round(3).to_string(index=False))
